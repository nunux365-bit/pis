"""Unit tests for O2C_OHC (no live DB / OpenAI)."""

from __future__ import annotations

from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest

from app.agents.o2c_ohc.json_validation import validate_contract_payload
from app.agents.o2c_ohc.pdf_probe import infer_ingestion_class
from app.agents.o2c_ohc.billing_profile import (
    BILLING_PROFILE_GENERIC,
    BILLING_PROFILE_TACO,
    BILLING_PROFILE_TCS,
    infer_billing_profile,
)
from app.agents.o2c_ohc.o2c_utils import _working_days_inclusive
from app.agents.o2c_ohc.pivot_slicer_attendance import (
    _cell_matches_allowed,
    _selected_values_for_slicer,
)
from app.agents.o2c_ohc.attendance_summary import parse_ohc_summary_workbook
from app.agents.o2c_ohc.ohc_roll_workbook_parser import (
    OFF_ROLL_CFG,
    UNIFIED_OHC_ATTENDANCE_SHEET,
    _parse_sheet,
    parse_ohc_roll_workbook,
    try_parse_ohc_roll_workbook,
    workbook_has_roll_tabs,
    workbook_has_unified_ohc_attendance_tab,
)
from app.agents.o2c_ohc.duplicate_day_cell_merge import (
    collapse_duplicate_day_cells,
    normalize_day_cell_attendance_value,
)
from app.agents.o2c_ohc.roll_to_summary_export import (
    _attendance_bucket,
    _fill_for_attendance_value,
)
from app.agents.o2c_ohc.folder_scanner import PdfCandidate, _file_sha256
from app.agents.o2c_ohc.llm_extract import _parse_payload_json
from app.agents.o2c_ohc.llm_extract import (
    CONTRACT_PROMPT_PROFILE_GENERIC,
    CONTRACT_PROMPT_PROFILE_TACO,
    base_instructions_for_prompt_profile,
)
from app.agents.o2c_ohc.pipeline import infer_contract_prompt_profile, process_one_contract_pdf


def test_validate_contract_payload_empty() -> None:
    errs = validate_contract_payload({})
    assert errs


def test_working_days_inclusive() -> None:
    assert _working_days_inclusive(date(2025, 3, 3), date(2025, 3, 7)) == 5  # Mon–Fri


def test_parse_ohc_roll_workbook_minimal(tmp_path: Path) -> None:
    from openpyxl import Workbook

    on_name = "OHC Attendance On Roll"
    off_name = "OHC Attendance Off Roll"
    wb = Workbook()
    w1 = wb.active
    w1.title = on_name
    w1.cell(2, 14, date(2026, 2, 1))
    w1.cell(3, 1, "Site-Alpha")
    w1.cell(3, 2, "Test Nurse")
    w1.cell(3, 3, "EMP001")
    w1.cell(3, 6, "Nurse")
    w1.cell(3, 14, "Present ( G )")
    w1.cell(3, 46, 20)

    w2 = wb.create_sheet(off_name)
    w2.cell(2, 15, date(2026, 2, 1))
    w2.cell(3, 1, "Site-Beta")
    w2.cell(3, 2, "Off Doc")
    w2.cell(3, 3, "HLD-1")
    w2.cell(3, 6, "Doctor")
    w2.cell(3, 15, "Week Off")

    p = tmp_path / "roll.xlsx"
    wb.save(p)

    assert workbook_has_roll_tabs(p)
    ds = parse_ohc_roll_workbook(p)
    assert ds.meta["total_ohc_sites"] == 2
    m = ds.to_client_site_attendance_map()
    assert "Site-Alpha" in m and "Site-Beta" in m
    # Present/absent come from date cells only, not trailing summary columns.
    assert m["Site-Alpha"][0]["present_days"] == 1
    assert m["Site-Beta"][0]["present_days"] == 0
    rollup = ds.mis_rollup_by_site_role()
    assert "Nurse" in rollup["Site-Alpha"]["by_role"]

    auto = parse_ohc_summary_workbook(p, engine="auto")
    assert set(auto.keys()) >= {"Site-Alpha", "Site-Beta"}


def test_parse_unified_ohc_roll_workbook_minimal(tmp_path: Path) -> None:
    """Single ``OHC Attendance`` tab + string date headers (no dual On/Off tabs)."""
    from openpyxl import Workbook

    unified = UNIFIED_OHC_ATTENDANCE_SHEET
    wb = Workbook()
    ws = wb.active
    ws.title = unified
    ws.cell(2, 15, "1-Apr-26")
    ws.cell(2, 16, "2-Apr-26")
    ws.cell(3, 1, "Site-Unified")
    ws.cell(3, 2, "Doc One")
    ws.cell(3, 3, "EMP-U1")
    ws.cell(3, 6, "Doctor")
    ws.cell(3, 15, "Present ( G )")
    ws.cell(3, 16, "Absent")

    p = tmp_path / "unified_roll.xlsx"
    wb.save(p)

    assert not workbook_has_roll_tabs(p)
    assert workbook_has_unified_ohc_attendance_tab(p)
    ds = try_parse_ohc_roll_workbook(p)
    assert ds is not None
    assert ds.meta.get("parser") == "ohc_roll_workbook_parser_unified_v1"
    m = ds.to_client_site_attendance_map()
    assert "Site-Unified" in m
    assert m["Site-Unified"][0]["present_days"] == 1
    assert m["Site-Unified"][0]["absent_days"] == 1
    assert m["Site-Unified"][0]["roll_type"] == "OHC Attendance"

    auto = parse_ohc_summary_workbook(p, engine="auto")
    assert "Site-Unified" in auto


def test_parse_sheet_duplicate_calendar_day_first_column_blank() -> None:
    """Wide templates repeat the same calendar day; left slot blank must not inherit right slot."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = OFF_ROLL_CFG["sheet"]
    ws.cell(2, 15, date(2026, 2, 1))
    ws.cell(2, 16, date(2026, 2, 1))
    for r, name, c15, c16 in (
        (3, "Alice", None, "Present ( G )"),
        (4, "Bob", "Week Off", "Present ( G )"),
    ):
        ws.cell(r, 1, "Site A")
        ws.cell(r, 2, name)
        ws.cell(r, 3, f"E{r}")
        for c in range(4, 15):
            ws.cell(r, c, "x")
        ws.cell(r, 15, c15)
        ws.cell(r, 16, c16)

    out = _parse_sheet(wb, OFF_ROLL_CFG, "Off Roll")
    by_name = {s["name"]: s for s in out["Site A"]}
    assert "2026-02-01" not in by_name["Alice"].get("attendance", {})
    assert by_name["Bob"]["attendance"]["2026-02-01"] == "Week Off"


def test_parse_ohc_summary_workbook(tmp_path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Emp Id", "Name", "Client Site", "Roll", "Role", "Present"])
    ws.append(["E1", "Alice", "ACME - Plant 1", "OnRoll", "NURSE_GNM", 22])
    ws.append(["E2", "Bob", "ACME - Plant 1", "OffRoll", "NURSE_GNM", 20])
    ws.append(["E3", "Cara", "BETA - WH", "OnRoll", "MO_MBBS", 21])
    p = tmp_path / "ohc.xlsx"
    wb.save(p)

    m = parse_ohc_summary_workbook(p, engine="flat")
    assert set(m.keys()) == {"ACME - Plant 1", "BETA - WH"}
    assert len(m["ACME - Plant 1"]) == 2
    assert m["ACME - Plant 1"][0]["present_days"] == Decimal(22)


def test_cell_matches_allowed() -> None:
    assert _cell_matches_allowed("Foo", {"Foo"})
    assert _cell_matches_allowed("foo", {"Foo"})


def test_selected_values_for_slicer_positional() -> None:
    cache_fields = [("Month", ["Jan", "Feb"]), ("Site", ["A", "B"])]
    slicer = {"sourceName": "Month", "pivotFieldIndex": None, "items": [(None, True), (None, False)]}
    name, vals, all_sel = _selected_values_for_slicer(cache_fields, slicer)
    assert name == "Month"
    assert all_sel is False
    assert vals == {"Jan"}


def test_collapse_duplicate_day_cells_first_column_wins() -> None:
    assert collapse_duplicate_day_cells(["Week Off", "Present ( G )"]) == "Week Off"
    assert collapse_duplicate_day_cells(["Present ( G )", "Week off"]) == "Present ( G )"


def test_collapse_duplicate_day_cells_absent_when_first() -> None:
    assert collapse_duplicate_day_cells(["Absent", "Present ( G )"]) == "Absent"


def test_collapse_duplicate_day_cells_preserves_first_blank() -> None:
    assert collapse_duplicate_day_cells(["", "Present ( G )"]) is None
    assert collapse_duplicate_day_cells([None, "Absent"]) is None


def test_normalize_non_text_day_cell_counts_as_blank() -> None:
    assert normalize_day_cell_attendance_value(42) is None
    assert normalize_day_cell_attendance_value(3.14) is None
    assert collapse_duplicate_day_cells([42, "Present ( G )"]) is None


def test_attendance_bucket_blank_and_unknown_not_present() -> None:
    assert _attendance_bucket(None) == "blank"
    assert _attendance_bucket("") == "blank"
    assert _attendance_bucket("   ") == "blank"
    assert _attendance_bucket("#REF!") == "other"
    assert _attendance_bucket("Present ( G )") == "present"


def test_fill_only_explicit_statuses() -> None:
    assert _fill_for_attendance_value("Present ( G )") is not None
    assert _fill_for_attendance_value("Absent") is not None
    assert _fill_for_attendance_value("Week off") is not None
    assert _fill_for_attendance_value(None) is None
    assert _fill_for_attendance_value("") is None
    assert _fill_for_attendance_value("???") is None


def test_tcs_taco_top_level_folders_are_not_auto_skipped_for_ingest() -> None:
    """Regression: ``tcs/`` and ``taco/`` paths use profile rules but must still run full ingest."""
    import app.agents.o2c_ohc.pipeline as pipeline_mod

    assert not hasattr(pipeline_mod, "skip_ingest_for_tcs_taco_top_folder")


def test_infer_contract_prompt_profile_folders() -> None:
    assert infer_contract_prompt_profile("a.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("tcs/foo.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("client_tcs/foo.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("tcs-2025/foo.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("generic/x.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("other/x.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert infer_contract_prompt_profile("taco/x.pdf") == CONTRACT_PROMPT_PROFILE_TACO
    assert infer_contract_prompt_profile("TACO/x.pdf") == CONTRACT_PROMPT_PROFILE_TACO
    assert infer_contract_prompt_profile("my_taco/x.pdf") == CONTRACT_PROMPT_PROFILE_TACO
    assert infer_contract_prompt_profile("taco_imports/x.pdf") == CONTRACT_PROMPT_PROFILE_TACO
    assert infer_contract_prompt_profile("montacola/x.pdf") == CONTRACT_PROMPT_PROFILE_GENERIC


def test_infer_billing_profile() -> None:
    assert infer_billing_profile("a.pdf", CONTRACT_PROMPT_PROFILE_GENERIC) == BILLING_PROFILE_GENERIC
    assert infer_billing_profile("tcs/foo.pdf", CONTRACT_PROMPT_PROFILE_GENERIC) == BILLING_PROFILE_TCS
    assert infer_billing_profile("client_tcs/foo.pdf", CONTRACT_PROMPT_PROFILE_GENERIC) == BILLING_PROFILE_TCS
    assert infer_billing_profile("taco/x.pdf", CONTRACT_PROMPT_PROFILE_TACO) == BILLING_PROFILE_TACO
    assert infer_billing_profile("tcs/x.pdf", CONTRACT_PROMPT_PROFILE_TACO) == BILLING_PROFILE_TACO


def test_base_instructions_taco_has_columnar_map_generic_does_not() -> None:
    taco = base_instructions_for_prompt_profile(CONTRACT_PROMPT_PROFILE_TACO)
    gen = base_instructions_for_prompt_profile(CONTRACT_PROMPT_PROFILE_GENERIC)
    assert "Rate tables for MIS (staffing / OHC) — columnar map" in taco
    assert "columnar map" not in gen.lower()
    assert "TCS / general IT OHC" in gen
    assert "How to read the table" in taco
    assert "MERGED CELLS" in taco
    assert "Col D: FMO fee" in taco
    assert "Col F: MO fee" in taco
    assert "N_FMO" in taco and "N_MO" in taco
    assert "Internal 3-pass" in taco
    assert "3a — Col E shared-MO" in taco
    assert "N\u22121" in taco  # unicode minus in “N−1” anti-pattern
    assert "Exception A" in taco
    assert "MO → `rate_lines` emission" in taco
    assert "**N_MO** MO lines" in taco
    assert "MO line count" in taco
    assert "Col D vs Col F" in taco
    assert "Chakan cluster — first merged MO band" in taco
    assert "3a → 3b → 3c" in taco
    assert "IPD Hinjewadi" in taco and "OHC_ADMIN_INVOICE_PCT" in taco
    assert "HELD IN ABEYANCE" in taco
    assert "Self-validation" in taco
    assert "contract_payload JSON" in taco
    assert "contracted_quantity" in taco and "payment_due_days" in taco


def test_infer_ingestion_class_text_pdf(tmp_path: Path) -> None:
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Hello contract " * 50)
    p = tmp_path / "t.pdf"
    doc.save(p)
    doc.close()
    assert infer_ingestion_class(p) in ("text_native", "mixed")


def test_process_one_contract_pdf_retries_llm_invalid_json(monkeypatch, tmp_path: Path) -> None:
    pdf = tmp_path / "c.pdf"
    pdf.write_bytes(b"%PDF-1.7 fake")
    c = PdfCandidate(
        absolute_path=pdf,
        relative_path="x/c.pdf",
        sha256="abc",
        mtime=datetime(2026, 1, 1, tzinfo=UTC),
    )

    calls = {"n": 0}

    async def _extract(*_a, **_k):
        calls["n"] += 1
        if calls["n"] == 1:
            raise ValueError("llm_invalid_json: Unterminated string")
        return (
            "# T\n\nok",
            {
                "extraction_metadata": {"overall_confidence": 0.9, "needs_human_review": False},
                "client": {"legal_name": "Acme", "short_name": "Acme"},
                "sites": [{"site_key": "acme_mumbai", "canonical_name": "x", "display_name": "x", "service_category": "ohc"}],
                "contract": {
                    "title": "t",
                    "contract_kind": "sow",
                    "effective_from": "2026-01-01",
                    "termination_notice_days": 30,
                },
                "parties": [{"party_role": "service_provider", "legal_name": "Tata 1MG"}, {"party_role": "primary_client", "legal_name": "Acme"}],
                "payment_terms": {"payment_due_days": 30, "gst_rate": 18.0, "tds_applicable": False},
                "rate_lines": [{
                    "site_key": "acme_mumbai",
                    "billing_model": "fixed_monthly",
                    "role_code": "MO_MBBS",
                    "description": "MO",
                    "attendance_required": False,
                    "schedule_type": "none",
                    "schedule_config": {},
                    "billing_rules": {},
                    "billing_rule_text": "x",
                }],
            },
            {"id": "r1"},
        )

    async def _ingest_vec(*_a, **_k):
        return {"status": "ok"}

    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.extract_markdown_and_payload_async", _extract)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.pdf_page_count", lambda *_a, **_k: 1)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.infer_ingestion_class", lambda *_a, **_k: "text_native")
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.validate_contract_payload", lambda *_a, **_k: [])
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline._ingest_contract_db_and_index_vector_async", _ingest_vec)
    failed_logs = {"n": 0}

    async def _log_fail(*_a, **_k):
        failed_logs["n"] += 1

    cleared = {"n": 0}

    async def _clear_fail(**_k):
        cleared["n"] += 1

    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.log_failed_contract_parsing_async", _log_fail)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.clear_failed_contract_parsing_async", _clear_fail)

    out = process_one_contract_pdf(c, contracts_root=tmp_path, max_repair_attempts=2)
    assert out["status"] == "ok"
    assert out["attempts"] == 2
    assert out.get("contract_prompt_profile") == CONTRACT_PROMPT_PROFILE_GENERIC
    assert calls["n"] == 2
    assert failed_logs["n"] == 0
    assert cleared["n"] == 1


def test_process_one_contract_pdf_computes_sha256_when_candidate_empty(monkeypatch, tmp_path: Path) -> None:
    pdf = tmp_path / "c.pdf"
    pdf.write_bytes(b"%PDF-1.7 fake contract bytes")
    expected_sha = _file_sha256(pdf)
    c = PdfCandidate(
        absolute_path=pdf,
        relative_path="x/c.pdf",
        sha256="",
        mtime=datetime(2026, 1, 1, tzinfo=UTC),
    )

    async def _extract(*_a, **_k):
        return (
            "# T\n\nok",
            {
                "extraction_metadata": {"overall_confidence": 0.9, "needs_human_review": False},
                "client": {"legal_name": "Acme", "short_name": "Acme"},
                "sites": [{"site_key": "acme_mumbai", "canonical_name": "x", "display_name": "x", "service_category": "ohc"}],
                "contract": {
                    "title": "t",
                    "contract_kind": "sow",
                    "effective_from": "2026-01-01",
                    "termination_notice_days": 30,
                },
                "parties": [{"party_role": "service_provider", "legal_name": "Tata 1MG"}, {"party_role": "primary_client", "legal_name": "Acme"}],
                "payment_terms": {"payment_due_days": 30, "gst_rate": 18.0, "tds_applicable": False},
                "rate_lines": [{
                    "site_key": "acme_mumbai",
                    "billing_model": "fixed_monthly",
                    "role_code": "MO_MBBS",
                    "description": "MO",
                    "attendance_required": False,
                    "schedule_type": "none",
                    "schedule_config": {},
                    "billing_rules": {},
                    "billing_rule_text": "x",
                }],
            },
            {"id": "r1"},
        )

    captured: dict[str, str] = {}

    async def _ingest_vec(*_a, file_sha256: str, **_k):
        captured["file_sha256"] = file_sha256
        return {"status": "ok"}

    async def _clear_fail(**_k):
        return None

    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.extract_markdown_and_payload_async", _extract)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.pdf_page_count", lambda *_a, **_k: 1)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.infer_ingestion_class", lambda *_a, **_k: "text_native")
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.validate_contract_payload", lambda *_a, **_k: [])
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline._ingest_contract_db_and_index_vector_async", _ingest_vec)
    monkeypatch.setattr("app.agents.o2c_ohc.pipeline.clear_failed_contract_parsing_async", _clear_fail)

    out = process_one_contract_pdf(c, contracts_root=tmp_path, max_repair_attempts=1)
    assert out["status"] == "ok"
    assert captured["file_sha256"] == expected_sha


def test_parse_payload_json_requires_strict_json_object() -> None:
    good = '{"contract": {}, "client": {}, "sites": [], "parties": [], "payment_terms": {}, "rate_lines": [], "extraction_metadata": {}}'
    parsed = _parse_payload_json(good)
    assert isinstance(parsed, dict)

    bad = "```json\\n{'contract': {}}\\n```"
    with pytest.raises(ValueError):
        _parse_payload_json(bad)
