"""Manpower → ``clinical_role_hint`` index (Employee ID → Specialist)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.agents.o2c_ohc.manpower_clinical_hint import (
    append_manpower_period_churn_rows,
    build_clinical_role_hint,
    build_manpower_clinical_hint_index,
    build_manpower_employee_index_for_site,
    enrich_amap_with_manpower_hints,
)
from app.agents.o2c_ohc.manpower_clinical_hint import ManpowerSiteRow
from app.config.settings import settings


def _manpower_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(["Name", "OHC Name", "Employee ID", "Specialist", "Role"])
    ws.append(["Alice", "Site-A", "E-100", "Counsellor", "Doctor"])
    ws.append(["Bob", "Site-A", "E-200", "Radiology", "Doctor"])
    ws.append(["Dup", "Site-A", "E-300", "First", "Nurse"])
    ws.append(["Dup2", "Site-A", "E-300", "Second", "Nurse"])
    wb.save(path)


def test_build_manpower_index_maps_employee_id_to_specialist(tmp_path: Path) -> None:
    p = tmp_path / "wb.xlsx"
    _manpower_xlsx(p)
    idx = build_manpower_clinical_hint_index(p)
    assert "Counsellor" in idx["E-100"]
    assert "Radiology" in idx["E-200"]
    assert "Second" in idx["E-300"]  # last site row wins for legacy id-only index


def test_enrich_amap_sets_clinical_role_hint_always_present(tmp_path: Path) -> None:
    p = tmp_path / "wb.xlsx"
    _manpower_xlsx(p)
    amap = {
        "Site-A": [
            {"employee_external_id": "E-100", "role_code": "Doctor"},
            {"employee_external_id": "E-999", "role_code": "Nurse"},
        ],
    }
    enrich_amap_with_manpower_hints(amap, p)
    assert amap["Site-A"][0]["clinical_role_hint"] == "attendance role: Doctor; manpower specialist: Counsellor"
    assert amap["Site-A"][1]["clinical_role_hint"] is None


def test_enrich_without_manpower_sheet_sets_null_hints(tmp_path: Path) -> None:
    p = tmp_path / "no_mp.xlsx"
    wb = Workbook()
    wb.active.append(["x"])
    wb.save(p)
    amap = {"S": [{"employee_external_id": "X", "role_code": "Doctor"}]}
    enrich_amap_with_manpower_hints(amap, p)
    assert amap["S"][0]["clinical_role_hint"] is None


def test_disabled_skips_read(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    p = tmp_path / "wb.xlsx"
    _manpower_xlsx(p)
    monkeypatch.setattr(settings, "o2c_manpower_clinical_hint_enabled", False)
    assert build_manpower_clinical_hint_index(p) == {}
    amap = {"S": [{"employee_external_id": "E-100", "role_code": "Doctor"}]}
    enrich_amap_with_manpower_hints(amap, p)
    assert amap["S"][0].get("clinical_role_hint") is None


def test_mis_summary_att_slim_includes_clinical_role_hint_key(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.agents.o2c_ohc.mis_summary_llm import _mis_summary_openai_chat_inputs

    monkeypatch.setattr(settings, "openai_api_key", "sk-test-dummy")

    tup = _mis_summary_openai_chat_inputs(
        client_site_key="Site",
        period_start=date(2026, 4, 1),
        period_end=date(2026, 4, 30),
        working_days=22,
        attendance_records=[
            {
                "employee_external_id": "E1",
                "employee_name": "A",
                "role_code": "Doctor",
                "clinical_role_hint": "Counsellor",
                "present_days": 20,
                "absent_days": None,
            }
        ],
        rate_lines=[{"id": "00000000-0000-0000-0000-000000000001", "billing_model": "fixed_monthly"}],
        past_corrections=None,
        billing_profile="generic",
    )
    assert tup is not None
    _, _, user = tup
    assert "clinical_role_hint" in user
    assert "Counsellor" in user
    assert "attendance_roll_role" not in user
    assert "manpower_designation" not in user


def test_detailed_json_includes_clinical_role_hint_when_set() -> None:
    from app.agents.o2c_ohc.mis_details import build_detailed_json_from_records

    dj = build_detailed_json_from_records(
        "Site-X",
        [
            {
                "employee_name": "Ann",
                "employee_external_id": "E1",
                "role_code": "Doctor",
                "clinical_role_hint": "Counsellor",
                "roll_type": "On Roll",
            }
        ],
        period_start=date(2026, 4, 1),
        period_end=date(2026, 4, 2),
    )
    assert dj["employees"][0]["clinical_role_hint"] == "Counsellor"


def _manpower_designation_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(
        [
            "Name",
            "OHC Name",
            "Employee ID",
            "Designation (as per PO/KAM)",
            "Role",
            "Date of joining",
            "Last working day",
        ]
    )
    ws.append(
        [
            "Dr Amit",
            "Site-A",
            "129851",
            "General Physician",
            "Doctor",
            "2026-04-16",
            None,
        ]
    )
    ws.append(
        [
            "Dr Dipak",
            "Site-A",
            "HLD-D-40",
            "General Physician",
            "Doctor",
            "2024-10-22",
            "2026-04-15",
        ]
    )
    wb.save(path)


def test_manpower_client_column_maps_site(tmp_path: Path) -> None:
    """Production Manpower tab uses ``Client`` instead of ``OHC Name`` (no Specialist column)."""
    p = tmp_path / "client_col.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(
        [
            "Name",
            "Client",
            "Employee ID",
            "Designation (as per PO/KAM)",
            "Contact No.",
            "Role",
        ]
    )
    ws.append(["Alice", "TACO CD-Pune-Chakan-phase II", "E-100", "MO", "999", "Doctor"])
    wb.save(p)
    idx = build_manpower_employee_index_for_site(p, "TACO CD-Pune-Chakan-phase II")
    assert "E-100" in idx
    assert idx["E-100"].designation == "MO"
    amap = {"TACO CD-Pune-Chakan-phase II": [{"employee_external_id": "E-100", "role_code": "Doctor"}]}
    enrich_amap_with_manpower_hints(amap, p)
    assert "manpower designation: MO" in (amap["TACO CD-Pune-Chakan-phase II"][0]["clinical_role_hint"] or "")


def test_designation_fallback_when_no_specialist_column(tmp_path: Path) -> None:
    p = tmp_path / "des.xlsx"
    _manpower_designation_xlsx(p)
    idx = build_manpower_employee_index_for_site(p, "Site-A")
    assert idx["129851"].designation == "General Physician"
    amap = {"Site-A": [{"employee_external_id": "129851", "role_code": "Doctor"}]}
    enrich_amap_with_manpower_hints(amap, p)
    row = amap["Site-A"][0]
    assert row["clinical_role_hint"] == "attendance role: Doctor; manpower designation: General Physician"
    assert row["role_code"] == "Doctor"
    assert "manpower_designation" not in row
    assert row["doj"] == "2026-04-16"


def test_build_clinical_role_hint_roll_and_designation() -> None:
    mp = ManpowerSiteRow(
        employee_id="129851",
        ohc_name="Site-A",
        designation="General Physician",
        specialist=None,
        mp_role="Doctor",
        doj_iso="2026-04-16",
        lwd_iso=None,
        employee_name="Dr Amit",
    )
    hint = build_clinical_role_hint(roll_role="Doctor", mp=mp)
    assert hint == "attendance role: Doctor; manpower designation: General Physician"


def test_multi_site_same_eid_uses_site_scoped_dates(tmp_path: Path) -> None:
    """Same employee id on two OHCs must not leak DOJ/LWD across sites."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(
        [
            "Name",
            "OHC Name",
            "Employee ID",
            "Designation (as per PO/KAM)",
            "Role",
            "Date of joining",
            "Last working day",
        ]
    )
    ws.append(["A SiteA", "Site-A", "SHARED-1", "GP A", "Doctor", "2026-04-16", None])
    ws.append(["A SiteB", "Site-B", "SHARED-1", "GP B", "Doctor", "2026-05-01", None])
    p = tmp_path / "multi.xlsx"
    wb.save(p)
    amap = {
        "Site-A": [{"employee_external_id": "SHARED-1", "role_code": "Doctor"}],
        "Site-B": [{"employee_external_id": "SHARED-1", "role_code": "Doctor"}],
    }
    enrich_amap_with_manpower_hints(amap, p)
    assert amap["Site-A"][0]["doj"] == "2026-04-16"
    assert amap["Site-B"][0]["doj"] == "2026-05-01"


def test_duplicate_eid_same_site_prefers_row_with_dates(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(
        ["Name", "OHC Name", "Employee ID", "Designation", "Role", "Date of joining", "Last working day"]
    )
    ws.append(["Sparse", "Site-A", "E-1", None, "Doctor", None, None])
    ws.append(["Rich", "Site-A", "E-1", "General Physician", "Doctor", "2026-04-16", "2026-04-15"])
    p = tmp_path / "dup.xlsx"
    wb.save(p)
    idx = build_manpower_employee_index_for_site(p, "Site-A")
    assert idx["E-1"].doj_iso == "2026-04-16"
    assert idx["E-1"].lwd_iso == "2026-04-15"


def test_site_scoped_clinical_hint_index(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manpower"
    ws.append(["Name", "OHC Name", "Employee ID", "Specialist", "Role"])
    ws.append(["A", "Site-A", "E-1", "Alpha", "Doctor"])
    ws.append(["B", "Site-B", "E-1", "Beta", "Doctor"])
    p = tmp_path / "scoped.xlsx"
    wb.save(p)
    assert "Alpha" in build_manpower_clinical_hint_index(p, client_site_key="Site-A")["E-1"]
    assert "Beta" in build_manpower_clinical_hint_index(p, client_site_key="Site-B")["E-1"]


def test_adds_manpower_only_joiner_for_billing_period(tmp_path: Path) -> None:
    p = tmp_path / "des.xlsx"
    _manpower_designation_xlsx(p)
    roll = [
        {
            "employee_external_id": "HLD-D-40",
            "role_code": "Doctor",
            "lwd": "2026-04-15",
        }
    ]
    enrich_amap_with_manpower_hints({"Site-A": roll}, p)
    assert {r["employee_external_id"] for r in roll} == {"HLD-D-40"}

    ps, pe = date(2026, 4, 1), date(2026, 4, 30)
    with_churn = append_manpower_period_churn_rows(roll, p, "Site-A", period_start=ps, period_end=pe)
    eids = {r["employee_external_id"] for r in with_churn}
    assert eids == {"HLD-D-40", "129851"}
    joiner = next(r for r in with_churn if r["employee_external_id"] == "129851")
    assert joiner.get("_from_manpower_only") is True
    assert joiner["doj"] == "2026-04-16"


def test_detailed_json_omits_clinical_when_null() -> None:
    from app.agents.o2c_ohc.mis_details import build_detailed_json_from_records

    dj = build_detailed_json_from_records(
        "Site-X",
        [
            {
                "employee_name": "Ann",
                "employee_external_id": "E1",
                "role_code": "Doctor",
                "clinical_role_hint": None,
                "roll_type": "On Roll",
            }
        ],
        period_start=date(2026, 4, 1),
        period_end=date(2026, 4, 2),
    )
    assert "clinical_role_hint" not in dj["employees"][0]
