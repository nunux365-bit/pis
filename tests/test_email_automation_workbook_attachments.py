"""Receivables attachment resolution (.xlsx, .xlsb, and .zip)."""

from __future__ import annotations

import os
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.email_automation.engine import excel_reader as xr
from app.email_automation.engine.classifier import InboundClassifier
from app.email_automation.gmail_sa import AttachmentStub, FetchedMessage
from app.email_automation.pipeline.workbook_attachments import (
    WorkbookResolveError,
    convert_xlsb_to_xlsx,
    normalize_receivable_workbook,
    resolve_downloaded_workbook,
)
from app.email_automation.workflow_packs import REGISTRY
from app.services import receivable_dashboard as rd

RECEIVABLE_XLSB_FIXTURE = Path(
    os.environ.get(
        "RECEIVABLE_XLSB_FIXTURE",
        "/Users/pankaj/Downloads/Receivable-08-07.2026.xlsb",
    )
)
RECEIVABLE_XLSX_FIXTURE = Path(
    os.environ.get(
        "RECEIVABLE_XLSX_FIXTURE",
        "/Users/pankaj/Downloads/Receivable-08-07.2026.xlsx",
    )
)


def test_resolve_passthrough_xlsx(tmp_path: Path) -> None:
    p = tmp_path / "Receivables-May.xlsx"
    p.write_bytes(b"x")
    assert resolve_downloaded_workbook(p) == p.resolve()


def test_resolve_passthrough_xlsb_converts(tmp_path: Path) -> None:
    if not RECEIVABLE_XLSB_FIXTURE.is_file():
        pytest.skip(f"fixture not found: {RECEIVABLE_XLSB_FIXTURE}")

    src = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    src.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())
    out = resolve_downloaded_workbook(src)
    assert out.suffix.lower() == ".xlsx"
    assert out.name == "Receivable-08-07.2026.xlsx"
    assert out.exists()
    assert src.exists()
    wb = load_workbook(out, read_only=True, data_only=True)
    try:
        assert "Receivable as on 08.07.2026" in wb.sheetnames
    finally:
        wb.close()


def test_convert_xlsb_to_xlsx_round_trip_openpyxl(tmp_path: Path) -> None:
    if not RECEIVABLE_XLSB_FIXTURE.is_file():
        pytest.skip(f"fixture not found: {RECEIVABLE_XLSB_FIXTURE}")

    src = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    src.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())
    out = convert_xlsb_to_xlsx(src)
    assert out.exists()
    wb = load_workbook(out, read_only=True, data_only=True)
    try:
        assert len(wb.sheetnames) >= 1
        ws = wb["Receivable as on 08.07.2026"]
        rows = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))
        assert rows[0][2] == "Code"
    finally:
        wb.close()


def test_resolve_zip_single_workbook(tmp_path: Path) -> None:
    xlsx = tmp_path / "Receivables-May.xlsx"
    xlsx.write_bytes(b"fake-xlsx")
    zpath = tmp_path / "Receivables-May.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        zf.write(xlsx, arcname="Receivables-May.xlsx")
    xlsx.unlink()
    out = resolve_downloaded_workbook(zpath)
    assert out.name == "Receivables-May.xlsx"
    assert out.exists()


def test_resolve_zip_single_xlsb_workbook(tmp_path: Path) -> None:
    if not RECEIVABLE_XLSB_FIXTURE.is_file():
        pytest.skip(f"fixture not found: {RECEIVABLE_XLSB_FIXTURE}")

    xlsb = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    xlsb.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())
    zpath = tmp_path / "Receivables-May.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        zf.write(xlsb, arcname=xlsb.name)
    xlsb.unlink()
    out = resolve_downloaded_workbook(zpath)
    assert out.suffix.lower() == ".xlsx"
    assert out.name == "Receivable-08-07.2026.xlsx"
    assert out.exists()


def test_resolve_zip_empty_raises(tmp_path: Path) -> None:
    zpath = tmp_path / "Receivables.zip"
    with zipfile.ZipFile(zpath, "w"):
        pass
    with pytest.raises(WorkbookResolveError, match="empty"):
        resolve_downloaded_workbook(zpath)


def test_resolve_zip_no_receivable_workbook_raises(tmp_path: Path) -> None:
    zpath = tmp_path / "Receivables.zip"
    with zipfile.ZipFile(zpath, "w") as zf:
        zf.writestr("report.xlsx", b"x")
    with pytest.raises(WorkbookResolveError, match="no receivables"):
        resolve_downloaded_workbook(zpath)


def test_normalize_receivable_workbook_rejects_bad_name(tmp_path: Path) -> None:
    p = tmp_path / "report.xlsb"
    p.write_bytes(b"x")
    with pytest.raises(WorkbookResolveError, match="does not match"):
        normalize_receivable_workbook(p)


def test_classifier_accepts_receivable_zip() -> None:
    pack = REGISTRY["PAYMENT_REMINDER_WEEKLY"]
    rule = pack.classifier_rule
    clf = InboundClassifier([rule])
    att = AttachmentStub(
        "Receivables 27-May-26.zip",
        "application/zip",
        50_000,
        "att-zip",
    )
    msg = FetchedMessage(
        id="m1",
        thread_id=None,
        sender="Bhawna Gandhi <bhawna.gandhi@1mg.com>",
        subject="Receivable/ Overdue as on 27th May'26",
        received_at_ms=None,
        headers={},
        attachments=(att,),
    )
    res = clf.classify(msg)
    assert res is not None
    assert res.attachment.filename.endswith(".zip")


def test_classifier_accepts_receivable_xlsb() -> None:
    pack = REGISTRY["PAYMENT_REMINDER_WEEKLY"]
    clf = InboundClassifier([pack.classifier_rule])
    att = AttachmentStub(
        "Receivable-08-07.2026.xlsb",
        "application/vnd.ms-excel.sheet.binary.macroEnabled.12",
        4_000_000,
        "att-xlsb",
    )
    msg = FetchedMessage(
        id="m2",
        thread_id=None,
        sender="Bhawna Gandhi <bhawna.gandhi@1mg.com>",
        subject="Receivable as on 08 Jul 2026",
        received_at_ms=None,
        headers={},
        attachments=(att,),
    )
    res = clf.classify(msg)
    assert res is not None
    assert res.attachment.filename.endswith(".xlsb")


@pytest.mark.slow
def test_receivable_dashboard_builds_from_converted_xlsb(tmp_path: Path) -> None:
    if not RECEIVABLE_XLSB_FIXTURE.is_file():
        pytest.skip(f"fixture not found: {RECEIVABLE_XLSB_FIXTURE}")

    src = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    src.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())
    xlsx = resolve_downloaded_workbook(src)
    payload = rd.build_receivable_payload(xlsx, {})
    assert payload is not None
    assert payload["meta"]["parser_version"] == rd.PARSER_VERSION
    assert payload["kpi_lakh"]
    assert payload["all_parties"]


def test_ingest_regex_accepts_xlsb_filename() -> None:
    from app.email_automation.pipeline.workbook_attachments import (
        RECEIVABLE_WORKBOOK_INGEST_RE,
    )

    assert RECEIVABLE_WORKBOOK_INGEST_RE.search("Receivable-08-07.2026.xlsb")
    assert RECEIVABLE_WORKBOOK_INGEST_RE.search("Receivables-May.xlsx")
    assert not RECEIVABLE_WORKBOOK_INGEST_RE.search("report.csv")


@pytest.mark.slow
def test_excel_reader_reads_converted_xlsb_receivable_tab(tmp_path: Path) -> None:
    if not RECEIVABLE_XLSB_FIXTURE.is_file():
        pytest.skip(f"fixture not found: {RECEIVABLE_XLSB_FIXTURE}")

    src = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    src.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())
    xlsx = resolve_downloaded_workbook(src)
    data = xr.read_sheet(
        xlsx,
        "Receivable as on",
        header_hints=(
            "code",
            "business unit",
            "net due",
            "0-1 months",
            "receivables",
        ),
    )
    assert data.sheet_name.startswith("Receivable as on")
    assert len(data.rows) > 100
    assert "code#0" in data.canonical_keys


@pytest.mark.slow
def test_xlsb_conversion_matches_native_xlsx_dashboard_and_reader(
    tmp_path: Path,
) -> None:
    """Golden parity: same as-on date exported as .xlsx vs .xlsb must parse identically."""

    if not RECEIVABLE_XLSB_FIXTURE.is_file() or not RECEIVABLE_XLSX_FIXTURE.is_file():
        pytest.skip(
            "paired fixtures not found — set RECEIVABLE_XLSB_FIXTURE and "
            "RECEIVABLE_XLSX_FIXTURE"
        )

    native_xlsx = tmp_path / RECEIVABLE_XLSX_FIXTURE.name
    native_xlsx.write_bytes(RECEIVABLE_XLSX_FIXTURE.read_bytes())
    xlsb = tmp_path / RECEIVABLE_XLSB_FIXTURE.name
    xlsb.write_bytes(RECEIVABLE_XLSB_FIXTURE.read_bytes())

    converted = resolve_downloaded_workbook(xlsb)
    native_dash = rd.build_receivable_payload(native_xlsx, {})
    conv_dash = rd.build_receivable_payload(converted, {})
    assert native_dash is not None and conv_dash is not None

    for section in (
        "kpi_lakh",
        "unbilled_lakh",
        "business_units",
        "all_parties",
        "top_parties",
        "collection_grids",
    ):
        assert native_dash[section] == conv_dash[section], section

    hints = ("code", "business unit", "net due", "0-1 months", "receivables")
    native_sheet = xr.read_sheet(native_xlsx, "Receivable as on", header_hints=hints)
    conv_sheet = xr.read_sheet(converted, "Receivable as on", header_hints=hints)
    assert native_sheet.canonical_keys == conv_sheet.canonical_keys
    assert len(native_sheet.rows) == len(conv_sheet.rows)

    def _by_code(rows: tuple[dict[str, object], ...]) -> dict[str, dict[str, object]]:
        out: dict[str, dict[str, object]] = {}
        for row in rows:
            code = row.get("code#0")
            if code:
                out[str(code)] = row
        return out

    native_by_code = _by_code(native_sheet.rows)
    conv_by_code = _by_code(conv_sheet.rows)
    assert native_by_code.keys() == conv_by_code.keys()
    for code in native_by_code:
        assert native_by_code[code] == conv_by_code[code], code
