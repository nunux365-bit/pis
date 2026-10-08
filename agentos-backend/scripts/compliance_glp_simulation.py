#!/usr/bin/env python3
"""
GLP compliance simulation: local recordings → ffmpeg normalize → Deepgram → OpenAI rubric
→ build Rollup row (same code path as production Sheets append) → diff vs reference xlsx.

Setup (from ``agentos-backend/``):

  export PYTHONPATH=.
  # Keys: use agentos-backend/.env (OPENAI_*, COMPLIANCE_DEEPGRAM_API_KEY) or export in shell.
  python scripts/compliance_glp_simulation.py \\
    --recordings-dir \"$HOME/Downloads/drive-download-20260430T104911Z-3-001\" \\
    --rollup-xlsx \"$HOME/Downloads/glp/glp1_consult_rollup.xlsx\"

  # Same transcript as golden ``*.deepgram.json`` / ``*.txt`` (isolates LLM vs live STT):
  python scripts/compliance_glp_simulation.py ... \\
    --reference-json-dir \"/path/to/_transcripts-deepgram\"

Do not commit API keys. Process env overrides .env when set.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import math
import sys
from pathlib import Path
from typing import Any

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.agents.compliance_call.adapters.normalize_ffmpeg import normalize_audio_file  # noqa: E402
from app.agents.compliance_call.adapters.rubric_openai import run_rubric_eval  # noqa: E402
from app.agents.compliance_call.adapters.transcribe_deepgram import (  # noqa: E402
    canonical_transcript_from_deepgram_response,
    transcribe_wav_to_result,
)
from app.agents.compliance_call.doctor_slug import doctor_name_from_filename, doctor_slug_from_filename  # noqa: E402
from app.agents.compliance_call.glp_scoring import deepgram_summary_for_persist, structural_from_deepgram  # noqa: E402
from app.agents.compliance_call.glp_sheet_rows import (  # noqa: E402
    _fmt_ts_utc,
    build_glp_rollup_row,
    glp_rollup_headers,
)
from app.config.settings import settings  # noqa: E402

log = logging.getLogger("compliance_glp_simulation")

_MEDIA_SUFFIXES = {".mpeg", ".mp3", ".m4a", ".wav", ".mp4", ".webm", ".ogg", ".mpga", ".aac", ".flac"}


def _list_recordings(root: Path) -> list[Path]:
    out: list[Path] = []
    for p in root.iterdir():
        if not p.is_file():
            continue
        if p.suffix.lower() in _MEDIA_SUFFIXES:
            out.append(p)
    return sorted(out, key=lambda x: (doctor_slug_from_filename(x.name), x.name.lower()))


def _deepgram_json_for_media(media_path: Path, ref_dir: Path) -> Path:
    """``Copy of Dr X..mpeg`` → ``Copy of Dr X..deepgram.json`` (single ``.mpeg`` suffix strip)."""
    stem = media_path.name.rsplit(".", 1)[0]
    p = ref_dir / f"{stem}.deepgram.json"
    if not p.is_file():
        raise FileNotFoundError(f"Reference Deepgram JSON not found: {p}")
    return p


def _transcript_and_summary_from_reference_json(jpath: Path) -> tuple[str, dict[str, Any]]:
    payload = json.loads(jpath.read_text(encoding="utf-8"))
    text = canonical_transcript_from_deepgram_response(payload)
    dg = deepgram_summary_for_persist(
        payload,
        structural_from_deepgram(payload, diarize=settings.compliance_deepgram_diarize),
    )
    return text, dg


def _read_reference_rollup(xlsx: Path) -> tuple[list[str], list[list[Any]]]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if "Rollup" not in wb.sheetnames:
        raise RuntimeError(f"Expected 'Rollup' sheet in {xlsx}, got {wb.sheetnames}")
    ws = wb["Rollup"]
    header_row = next(ws.iter_rows(min_row=4, max_row=4, values_only=True))
    headers = [("" if c is None else str(c).strip()) for c in header_row][:49]
    canon = list(glp_rollup_headers())
    if headers != canon:
        log.warning(
            "Reference header row 4 differs from glp_rollup_headers() (comparing by column index). "
            "First diff at: %s",
            next((i for i, (a, b) in enumerate(zip(headers, canon)) if a != b), None),
        )
    data: list[list[Any]] = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        vals = list(row)
        while len(vals) < 49:
            vals.append(None)
        vals = vals[:49]
        if all(v is None or (isinstance(v, str) and not str(v).strip()) for v in vals):
            break
        data.append(vals)
    return canon, data


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isfinite(v) and abs(v - round(v)) < 1e-9:
            return str(int(round(v)))
        return str(v).rstrip("0").rstrip(".") if "." in str(v) else str(v)
    return str(v).strip()


def _cells_close(a: str, b: str) -> bool:
    if a == b:
        return True
    try:
        fa, fb = float(a), float(b)
        return math.isclose(fa, fb, rel_tol=0.02, abs_tol=0.25)
    except ValueError:
        return False


def _compare_rows(
    expected: list[Any],
    actual: list[str],
    *,
    skip_cols: set[int],
    headers: list[str],
) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    for i, (e, a) in enumerate(zip(expected, actual)):
        if i in skip_cols:
            continue
        es, astr = _cell_str(e), _cell_str(a)
        if _cells_close(es, astr) or es == astr:
            continue
        diffs.append(
            {
                "col": i,
                "header": headers[i] if i < len(headers) else "?",
                "expected": es[:500],
                "actual": astr[:500],
            }
        )
    return diffs


def _norm_doc_name(s: str) -> str:
    return " ".join((s or "").strip().lower().split())


def _pair_recordings_to_expected_rows(
    recordings: list[Path],
    expected_rows: list[list[Any]],
) -> list[tuple[Path, list[Any]]]:
    """Match each xlsx row to a file by **Doctor** (col 1), in sheet order; consume files per doctor by filename."""
    from collections import defaultdict, deque

    buckets: dict[str, deque[Path]] = defaultdict(deque)
    for p in sorted(recordings, key=lambda x: x.name.lower()):
        key = _norm_doc_name(doctor_name_from_filename(p.name))
        buckets[key].append(p)

    out: list[tuple[Path, list[Any]]] = []
    for row in expected_rows:
        exp_name = _norm_doc_name(str(row[1] if len(row) > 1 else ""))
        if not exp_name:
            raise RuntimeError("Reference row missing Doctor (column B)")
        q = buckets.get(exp_name)
        if not q:
            raise RuntimeError(
                f"No recording left for reference doctor {row[1]!r} (normalized={exp_name!r}). "
                f"Check filenames vs sheet order."
            )
        out.append((q.popleft(), row))

    leftover = {k: list(v) for k, v in buckets.items() if v}
    if leftover:
        raise RuntimeError(f"Unmatched recordings after pairing: {leftover}")
    return out


def _rollup_from_pipeline(
    *,
    doctor_name: str,
    doctor_slug: str,
    eval_doc: dict[str, Any],
    deepgram_summary: dict[str, Any],
) -> list[str]:
    dg = deepgram_summary if isinstance(deepgram_summary, dict) else {}
    struct = dg.get("structural") if isinstance(dg.get("structural"), dict) else {}
    dg_ts = _fmt_ts_utc(str(dg.get("deepgram_created") or ""))
    dp = eval_doc.get("domain_pcts") if isinstance(eval_doc.get("domain_pcts"), dict) else {}
    patient_summary = str(eval_doc.get("patient_summary") or "")
    comments = str(eval_doc.get("comments") or "")
    scored_at = str(eval_doc.get("scored_at") or "")
    grade_label = str(eval_doc.get("grade_label") or "")
    status_text = f"{eval_doc.get('grade') or ''} ({grade_label})".strip() if grade_label else str(
        eval_doc.get("grade") or ""
    )
    return build_glp_rollup_row(
        serial_no="",
        doctor_name=doctor_name or doctor_slug,
        patient_summary=patient_summary,
        date_scored=scored_at,
        call_timestamp_utc=dg_ts,
        structural=struct,
        domain_pcts=dp,
        composite_pct=eval_doc.get("composite_pct"),
        grade=str(eval_doc.get("grade") or ""),
        grade_label=grade_label,
        status_text=status_text,
        eval_doc=eval_doc,
        comments=comments,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Simulate GLP pipeline vs reference Rollup xlsx.")
    parser.add_argument(
        "--recordings-dir",
        type=Path,
        required=True,
        help="Folder with audio/video files (e.g. Drive export).",
    )
    parser.add_argument(
        "--rollup-xlsx",
        type=Path,
        required=True,
        help="Reference ``glp1_consult_rollup.xlsx`` (Rollup sheet, header row 4, data from row 5).",
    )
    parser.add_argument(
        "--rubric-version",
        default="",
        help="Passed to run_rubric_eval; empty uses settings.compliance_rubric_version.",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Compare all columns including Patient / Date / timestamps / Comments (default skips LLM-volatile cols).",
    )
    parser.add_argument(
        "--reference-json-dir",
        type=Path,
        default=None,
        help="If set, load transcript + structural from ``<stem>.deepgram.json`` here (skips ffmpeg + Deepgram API).",
    )
    parser.add_argument("-q", "--quiet", action="store_true")
    parser.add_argument(
        "--json-out",
        type=Path,
        default=None,
        help="Write full per-file summary JSON to this path.",
    )
    args = parser.parse_args()

    logging.basicConfig(level=logging.WARNING if args.quiet else logging.INFO)

    if not (settings.openai_api_key or "").strip():
        log.error("OPENAI_API_KEY is not set (.env or environment).")
        return 1
    ref_json = args.reference_json_dir.expanduser().resolve() if args.reference_json_dir else None
    if not ref_json and not (settings.compliance_deepgram_api_key or "").strip():
        log.error("COMPLIANCE_DEEPGRAM_API_KEY is not set (or pass --reference-json-dir to skip Deepgram).")
        return 1

    rec_dir = args.recordings_dir.expanduser().resolve()
    xlsx = args.rollup_xlsx.expanduser().resolve()
    if not rec_dir.is_dir():
        log.error("Recordings dir not found: %s", rec_dir)
        return 1
    if not xlsx.is_file():
        log.error("Rollup xlsx not found: %s", xlsx)
        return 1

    headers, expected_rows = _read_reference_rollup(xlsx)
    recordings = _list_recordings(rec_dir)

    if len(recordings) != len(expected_rows):
        log.error(
            "Row count mismatch: %d recordings vs %d data rows in xlsx (after header row 4). "
            "Use same folder / sheet as the reference.",
            len(recordings),
            len(expected_rows),
        )
        return 1

    # Default: skip columns that are expected to differ between runs (LLM + clock).
    skip_cols: set[int] = set()
    if not args.strict:
        skip_cols.update(
            {
                0,  # conversation_id (we emit empty)
                2,  # Patient summary (LLM)
                3,  # Date Scored (wall clock)
                4,  # Call Timestamp — STT metadata may differ slightly
                5,  # Duration (minor float)
                6,  # Total turns (segmentation)
                7,  # Avg turn
                8,  # Longest turn
                9,  # Speech activity %
                10,  # Dr:Pt ratio (structural)
                19,  # Status — reference often stores label only; pipeline uses "Grade (label)"
                48,  # Comments (LLM)
            }
        )

    results: list[dict[str, Any]] = []
    all_ok = True

    try:
        paired = _pair_recordings_to_expected_rows(recordings, expected_rows)
    except RuntimeError as e:
        log.error("%s", e)
        return 1

    for path, exp_row in paired:
        slug = doctor_slug_from_filename(path.name)
        dname = doctor_name_from_filename(path.name)
        log.info("Processing %s (doctor=%s)", path.name, dname)

        wav: Path | None = None
        try:
            if ref_json:
                jpath = _deepgram_json_for_media(path, ref_json)
                transcript, dg_summary = _transcript_and_summary_from_reference_json(jpath)
                log.info("  using reference JSON %s", jpath.name)
            else:
                wav = normalize_audio_file(path, loudnorm=settings.compliance_ffmpeg_loudnorm)
                tr = asyncio.run(transcribe_wav_to_result(wav))
                transcript = str(tr.get("transcript_text") or "")
                dg_summary = tr.get("deepgram_summary") if isinstance(tr.get("deepgram_summary"), dict) else {}

            rv = (args.rubric_version or "").strip() or None
            eval_doc = asyncio.run(
                run_rubric_eval(
                    transcript_text=transcript,
                    doctor_slug=slug,
                    doctor_name=dname,
                    rubric_version=rv,
                    deepgram_summary=dg_summary,
                )
            )
            actual_row = _rollup_from_pipeline(
                doctor_name=dname,
                doctor_slug=slug,
                eval_doc=eval_doc,
                deepgram_summary=dg_summary,
            )
        except Exception as e:
            log.exception("Failed on %s", path.name)
            results.append(
                {
                    "file": path.name,
                    "ok": False,
                    "error": f"{type(e).__name__}: {e}"[:2000],
                }
            )
            all_ok = False
            continue
        finally:
            if wav is not None:
                wav.unlink(missing_ok=True)

        diffs = _compare_rows(exp_row, actual_row, skip_cols=skip_cols, headers=headers)
        file_ok = len(diffs) == 0
        if not file_ok:
            all_ok = False
        results.append(
            {
                "file": path.name,
                "ok": file_ok,
                "doctor": dname,
                "diffs": diffs,
                "diff_count": len(diffs),
                "eval": {
                    "composite_pct": eval_doc.get("composite_pct"),
                    "grade": eval_doc.get("grade"),
                    "grade_label": eval_doc.get("grade_label"),
                    "model": eval_doc.get("model"),
                    "patient_summary": str(eval_doc.get("patient_summary") or "").strip(),
                    "comments": str(eval_doc.get("comments") or "").strip(),
                },
            }
        )
        if file_ok:
            log.info("  MATCH (within rules): %s", path.name)
        else:
            log.warning("  MISMATCH: %s — %d column diffs (showing up to 8)", path.name, len(diffs))
            for d in diffs[:8]:
                log.warning('    col %d [%s]: expected=%r actual=%r', d["col"], d["header"], d["expected"], d["actual"])

    summary = {
        "recordings": len(recordings),
        "pairing": "by_reference_doctor_column",
        "reference_json_dir": str(ref_json) if ref_json else None,
        "strict_mode": args.strict,
        "skipped_col_indices": sorted(skip_cols),
        "all_match": all_ok,
        "run_config": {
            "compliance_deepgram_model": (settings.compliance_deepgram_model or "").strip(),
            "compliance_deepgram_diarize": settings.compliance_deepgram_diarize,
            "compliance_openai_model": (settings.compliance_openai_model or "").strip()
            or (settings.openai_chat_model or "").strip(),
            "max_output_tokens": int(
                settings.compliance_openai_max_output_tokens or settings.o2c_llm_max_output_tokens
            ),
        },
        "results": results,
    }
    if args.json_out:
        args.json_out.write_text(json.dumps(summary, indent=2), encoding="utf-8")
        log.info("Wrote %s", args.json_out)

    print(json.dumps({"all_match": all_ok, "files": len(results)}, indent=2))
    return 0 if all_ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
