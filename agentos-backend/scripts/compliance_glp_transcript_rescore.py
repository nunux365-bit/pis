#!/usr/bin/env python3
"""
GLP compliance re-score on **already-transcribed** calls.

Unlike ``scripts/compliance_glp_simulation.py`` this does **not** need recordings
(no ffmpeg / Deepgram) and does **not** need a reference Rollup xlsx. It reads the
transcripts already stored on completed ``compliance_call`` ``workflow_runs``
(``output_data.transcript_text`` / ``normalized_transcript_text`` + ``deepgram_summary``),
re-runs the OpenAI rubric (same code path as production), builds the Rollup row
(``build_glp_rollup_row``) and — optionally — writes those rows out as an xlsx.

Setup (from ``agentos-backend/``)::

  export PYTHONPATH=.
  # Keys/DB: use agentos-backend/.env (OPENAI_*, DATABASE_URL) or export in shell.

  # Re-score everything transcribed, write the rollup sheet:
  python scripts/compliance_glp_transcript_rescore.py \\
    --rollup-out "$HOME/Downloads/glp/rescored_rollup.xlsx"

  # Only calls created between two dates (WorkflowRun.created_at, inclusive):
  python scripts/compliance_glp_transcript_rescore.py \\
    --created-after 2026-07-01 --created-before 2026-07-27 \\
    --rollup-out "$HOME/Downloads/glp/rescored_rollup.xlsx"

  # Skip the LLM entirely and just rebuild the sheet from the stored eval:
  python scripts/compliance_glp_transcript_rescore.py \\
    --reuse-stored-eval --rollup-out "$HOME/Downloads/glp/from_stored.xlsx"

Filters (all optional, combinable): ``--created-after/before``, ``--conversation-id``,
``--workflow-run-id``, ``--doctor``, ``--source-type``, ``--limit``.

Do not commit API keys. Process env overrides .env when set.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import os
import re
import sys
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from sqlalchemy import inspect as sa_inspect, select  # noqa: E402

from app.agents.compliance_call.adapters.rubric_openai import run_rubric_eval  # noqa: E402
from app.agents.compliance_call.adapters.sheets_append import sheet_serial_no_for_state  # noqa: E402
from app.agents.compliance_call.glp_scoring import load_glp_rubric  # noqa: E402
from app.agents.compliance_call.glp_sheet_rows import (  # noqa: E402
    _fmt_ts_utc,
    build_glp_detailed_row,
    build_glp_rollup_row,
    glp_detailed_headers,
    glp_rollup_headers,
)
from app.config.settings import settings  # noqa: E402
from app.db.models import ComplianceCallRun, WorkflowRun, WorkflowRunStatus  # noqa: E402
from app.db.session import AsyncSessionLocal, engine  # noqa: E402
from app.services.compliance_call_projection import (  # noqa: E402
    compliance_call_run_from_workflow_run,
)

import uuid  # noqa: E402

log = logging.getLogger("compliance_glp_transcript_rescore")

CANONICAL_COMPLIANCE_CALL_WORKFLOW_KEY = "compliance_call"

_TRUTHY = {"1", "true", "yes", "on"}


def _maybe_inject_system_trust(*, enabled: bool) -> None:
    """Verify TLS via the OS trust store (macOS Keychain) instead of certifi/OpenSSL.

    Local-only: behind a corporate MITM proxy (e.g. Zscaler) certifi does not trust the
    re-signing CA and OpenSSL 3.x rejects it. Staging/production egress directly, so this
    is **off by default** — enable with --use-system-trust or RESCORE_USE_SYSTEM_TRUST=1.
    """
    if not enabled:
        return
    try:
        import truststore

        truststore.inject_into_ssl()
        log.info("system trust store injected (truststore) — TLS verified via OS keychain")
    except ImportError:
        log.warning(
            "--use-system-trust set but 'truststore' is not installed; "
            "run `uv pip install truststore`. Continuing with default (certifi) trust."
        )


# ---------------------------------------------------------------------------
# CLI date parsing
# ---------------------------------------------------------------------------
def _parse_dt(s: str, *, end_of_day: bool) -> datetime:
    """Accept ``YYYY-MM-DD`` or full ISO. Naive inputs are treated as UTC.

    A bare date as ``--created-after`` means 00:00:00; as ``--created-before`` it
    means 23:59:59.999999 so the whole day is inclusive.
    """
    raw = (s or "").strip()
    try:
        dt = datetime.fromisoformat(raw)
    except ValueError:
        d = date.fromisoformat(raw)
        dt = datetime.combine(d, time.max if end_of_day else time.min)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


# ---------------------------------------------------------------------------
# Load candidate runs
# ---------------------------------------------------------------------------
def _input_str(inp: dict[str, Any], key: str) -> str:
    v = inp.get(key)
    return str(v).strip() if v is not None else ""


async def _load_runs(args: argparse.Namespace) -> list[WorkflowRun]:
    async with AsyncSessionLocal() as db:
        q = select(WorkflowRun).where(
            WorkflowRun.workflow_key == CANONICAL_COMPLIANCE_CALL_WORKFLOW_KEY
        )
        if not args.any_status:
            q = q.where(WorkflowRun.status == WorkflowRunStatus.COMPLETED.value)
        if args.created_after:
            q = q.where(WorkflowRun.created_at >= _parse_dt(args.created_after, end_of_day=False))
        if args.created_before:
            q = q.where(WorkflowRun.created_at <= _parse_dt(args.created_before, end_of_day=True))
        if args.workflow_run_id:
            q = q.where(WorkflowRun.id.in_([w.strip() for w in args.workflow_run_id]))
        q = q.order_by(WorkflowRun.created_at.asc())
        rows = (await db.execute(q)).scalars().all()

    # JSONB / substring filters applied in Python (robust across dialects).
    wanted_conv = {str(c).strip() for c in (args.conversation_id or [])}
    doc_needle = (args.doctor or "").strip().lower()
    out: list[WorkflowRun] = []
    for r in rows:
        inp = r.input_data if isinstance(r.input_data, dict) else {}
        source = _input_str(inp, "source") or (
            "mysql_call" if r.mysql_call_id is not None else "gdrive"
        )
        if args.source_type != "all":
            want = "mysql_call" if args.source_type == "mysql" else "gdrive"
            if source != want:
                continue
        if wanted_conv:
            conv = _input_str(inp, "mysql_second_opinion_conversation_id")
            if conv not in wanted_conv:
                continue
        if doc_needle and doc_needle not in _input_str(inp, "doctor_name").lower():
            continue
        out.append(r)
        if args.limit and len(out) >= args.limit:
            break
    return out


# ---------------------------------------------------------------------------
# Extract stored transcript
# ---------------------------------------------------------------------------
def _stored_transcript(od: dict[str, Any], source: str) -> str:
    """Grading text the way the graph fed the rubric: normalized first, else canonical.

    (The utterance-level ``grading_transcript`` is not persisted, so ``auto`` falls
    back to canonical when no normalized text is stored.)
    """
    norm = str(od.get("normalized_transcript_text") or "").strip()
    canon = str(od.get("transcript_text") or "").strip()
    if source == "normalized":
        return norm
    if source == "canonical":
        return canon
    return norm or canon


# ---------------------------------------------------------------------------
# Build the Rollup / Detailed row (same shape as compliance_glp_simulation.py)
# ---------------------------------------------------------------------------
def _rollup_common_fields(
    *, eval_doc: dict[str, Any], deepgram_summary: dict[str, Any]
) -> dict[str, Any]:
    dg = deepgram_summary if isinstance(deepgram_summary, dict) else {}
    struct = dg.get("structural") if isinstance(dg.get("structural"), dict) else {}
    dp = eval_doc.get("domain_pcts") if isinstance(eval_doc.get("domain_pcts"), dict) else {}
    grade_label = str(eval_doc.get("grade_label") or "")
    status_text = (
        f"{eval_doc.get('grade') or ''} ({grade_label})".strip()
        if grade_label
        else str(eval_doc.get("grade") or "")
    )
    return {
        "structural": struct,
        "domain_pcts": dp,
        "composite_pct": eval_doc.get("composite_pct"),
        "grade": str(eval_doc.get("grade") or ""),
        "grade_label": grade_label,
        "status_text": status_text,
        "patient_summary": str(eval_doc.get("patient_summary") or ""),
        "date_scored": str(eval_doc.get("scored_at") or ""),
        "call_timestamp_utc": _fmt_ts_utc(str(dg.get("deepgram_created") or "")),
        "comments": str(eval_doc.get("comments") or ""),
    }


def _rollup_row(
    *, serial_no: str, doctor_name: str, eval_doc: dict[str, Any], deepgram_summary: dict[str, Any]
) -> list[str]:
    f = _rollup_common_fields(eval_doc=eval_doc, deepgram_summary=deepgram_summary)
    return build_glp_rollup_row(
        serial_no=serial_no,
        doctor_name=doctor_name,
        patient_summary=f["patient_summary"],
        date_scored=f["date_scored"],
        call_timestamp_utc=f["call_timestamp_utc"],
        structural=f["structural"],
        domain_pcts=f["domain_pcts"],
        composite_pct=f["composite_pct"],
        grade=f["grade"],
        grade_label=f["grade_label"],
        status_text=f["status_text"],
        eval_doc=eval_doc,
        comments=f["comments"],
    )


def _detailed_row(
    *,
    rubric: dict[str, Any],
    serial_no: str,
    doctor_name: str,
    eval_doc: dict[str, Any],
    deepgram_summary: dict[str, Any],
) -> list[str]:
    f = _rollup_common_fields(eval_doc=eval_doc, deepgram_summary=deepgram_summary)
    return build_glp_detailed_row(
        rubric=rubric,
        serial_no=serial_no,
        doctor_name=doctor_name,
        patient_summary=f["patient_summary"],
        date_scored=f["date_scored"],
        call_timestamp_utc=f["call_timestamp_utc"],
        structural=f["structural"],
        domain_pcts=f["domain_pcts"],
        composite_pct=f["composite_pct"],
        grade=f["grade"],
        grade_label=f["grade_label"],
        status_text=f["status_text"],
        eval_doc=eval_doc,
        comments=f["comments"],
    )


# Excel: cells cap at 32,767 chars; control chars (except \t \n \r) are illegal in xlsx.
_XLSX_MAX_CELL = 32_767
_XLSX_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xlsx_safe(s: str) -> str:
    return _XLSX_ILLEGAL.sub("", str(s or ""))[:_XLSX_MAX_CELL]


# ---------------------------------------------------------------------------
# Write the rollup workbook (Rollup sheet: header row 4, data row 5 — round-trips
# with compliance_glp_simulation.py --rollup-xlsx / glp_validate_reference_workbook.py).
# An optional Transcript column is appended AFTER the fixed 49 (Rollup) / 93 (Detailed)
# columns so the canonical layout still round-trips (readers take only the first 49/93).
# ---------------------------------------------------------------------------
def _write_rollup_xlsx(
    out_path: Path,
    *,
    rollup_rows: list[list[str]],
    detailed_rows: list[list[str]] | None,
    rubric: dict[str, Any],
    transcripts: list[str] | None = None,
) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wsr = wb.active
    wsr.title = "Rollup"
    headers = list(glp_rollup_headers())
    wsr.cell(row=1, column=1, value="GLP-1 Consult Rollup (re-scored from stored transcripts)")
    for c, h in enumerate(headers, start=1):
        wsr.cell(row=4, column=c, value=h)
    if transcripts is not None:
        wsr.cell(row=4, column=len(headers) + 1, value="Transcript")
    for i, row in enumerate(rollup_rows, start=5):
        for c, val in enumerate(row, start=1):
            wsr.cell(row=i, column=c, value=val)
        if transcripts is not None:
            wsr.cell(row=i, column=len(headers) + 1, value=_xlsx_safe(transcripts[i - 5]))

    if detailed_rows is not None:
        wsd = wb.create_sheet("Detailed")
        dheaders = list(glp_detailed_headers(rubric))
        wsd.cell(row=1, column=1, value="GLP-1 Consult Detailed (re-scored from stored transcripts)")
        for c, h in enumerate(dheaders, start=1):
            wsd.cell(row=4, column=c, value=h)
        if transcripts is not None:
            wsd.cell(row=4, column=len(dheaders) + 1, value="Transcript")
        for i, row in enumerate(detailed_rows, start=5):
            for c, val in enumerate(row, start=1):
                wsd.cell(row=i, column=c, value=val)
            if transcripts is not None:
                wsd.cell(row=i, column=len(dheaders) + 1, value=_xlsx_safe(transcripts[i - 5]))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------------------
# Persist re-scored result: update workflow_runs.output_data.eval, then
# DELETE + INSERT the compliance_call_runs projection row (what the dashboard reads).
# ---------------------------------------------------------------------------
def _dec_str(v: Any) -> Any:
    return str(v) if v is not None else None


async def _persist_rescored(
    wid: uuid.UUID,
    *,
    eval_doc: dict[str, Any],
    reuse_stored_eval: bool,
    commit: bool,
) -> dict[str, Any]:
    """Write the new eval back and rebuild the projection row (delete-then-insert).

    ``commit=False`` computes the change and rolls back (dry run) so you can see
    the before/after without touching production data.
    """
    async with AsyncSessionLocal() as db:
        r = await db.get(WorkflowRun, wid)
        if r is None or r.workflow_key != CANONICAL_COMPLIANCE_CALL_WORKFLOW_KEY:
            return {"persisted": False, "reason": "run_not_found"}

        prev_out = dict(r.output_data) if isinstance(r.output_data, dict) else {}
        prev_eval = prev_out.get("eval") if isinstance(prev_out.get("eval"), dict) else {}
        prev_marker = {
            "composite_pct": prev_eval.get("composite_pct"),
            "grade": prev_eval.get("grade"),
            "grade_label": prev_eval.get("grade_label"),
        }

        # Only overwrite output_data.eval when we produced a fresh one. When reusing
        # the stored eval we just rebuild the projection row from what is already there.
        if not reuse_stored_eval:
            new_out = dict(prev_out)
            new_out["eval"] = eval_doc
            new_out["rescored_at"] = datetime.now(timezone.utc).isoformat()
            new_out["rescored_prev"] = prev_marker
            r.output_data = new_out  # reassign so SQLAlchemy tracks the JSONB change

        existing = (
            await db.execute(
                select(ComplianceCallRun).where(ComplianceCallRun.workflow_run_id == wid)
            )
        ).scalar_one_or_none()
        projection_deleted = existing is not None
        if existing is not None:
            await db.delete(existing)

        # Flush the pending output_data UPDATE + projection DELETE together (the DELETE must land
        # before the re-INSERT to avoid the unique(workflow_run_id) collision). The flush fires
        # updated_at's onupdate=func.now(), which expires the attribute; refresh eagerly so building
        # the projection payload below doesn't trigger a lazy (sync) reload → MissingGreenlet.
        await db.flush()
        insp = sa_inspect(r)
        if insp.expired or insp.expired_attributes:
            await db.refresh(r)

        payload = compliance_call_run_from_workflow_run(r)
        payload.pop("workflow_run_id", None)
        db.add(ComplianceCallRun(id=uuid.uuid4(), workflow_run_id=wid, **payload))

        if commit:
            await db.commit()
        else:
            await db.rollback()

        return {
            "persisted": bool(commit),
            "dry_run": not commit,
            "projection_deleted": projection_deleted,
            "output_data_eval_updated": not reuse_stored_eval,
            "prev": {
                "composite_pct": _dec_str(prev_marker.get("composite_pct")),
                "grade": prev_marker.get("grade"),
                "grade_label": prev_marker.get("grade_label"),
            },
            "new": {
                "composite_pct": _dec_str(payload.get("composite_pct")),
                "grade": payload.get("grade"),
                "grade_label": payload.get("grade_label"),
            },
        }


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> int:
    parser = argparse.ArgumentParser(
        description="Re-score GLP compliance from stored call transcripts (no recordings, no reference xlsx)."
    )
    parser.add_argument("--created-after", default=None, help="Only runs with created_at >= this date/ISO (inclusive).")
    parser.add_argument("--created-before", default=None, help="Only runs with created_at <= this date/ISO (inclusive; bare date = end of day).")
    parser.add_argument("--conversation-id", action="append", default=None, help="Filter by MySQL second_opinion_conversation_id (repeatable).")
    parser.add_argument("--workflow-run-id", action="append", default=None, help="Filter by workflow_runs.id (repeatable).")
    parser.add_argument("--doctor", default=None, help="Case-insensitive substring match on stored doctor_name.")
    parser.add_argument("--source-type", choices=["all", "mysql", "gdrive"], default="all")
    parser.add_argument("--any-status", action="store_true", help="Include non-completed runs (default: COMPLETED only).")
    parser.add_argument("--limit", type=int, default=0, help="Max runs to process (0 = no limit).")
    parser.add_argument("--reuse-stored-eval", action="store_true", help="Skip OpenAI; build rows from the eval stored on the run.")
    parser.add_argument("--transcript-source", choices=["auto", "canonical", "normalized"], default="auto", help="Which stored transcript to re-score (default auto: normalized else canonical).")
    parser.add_argument("--rubric-version", default="", help="Passed to run_rubric_eval; empty uses settings default.")
    parser.add_argument("--rollup-out", type=Path, default=None, help="Write the rollup workbook here (optional).")
    parser.add_argument("--detailed", action="store_true", help="Also add a 93-col Detailed sheet to --rollup-out.")
    parser.add_argument(
        "--with-transcript",
        action="store_true",
        help="Append a 'Transcript' column (the text that was scored) after the fixed columns in --rollup-out.",
    )
    parser.add_argument(
        "--persist",
        action="store_true",
        help="Write results back to the DB: update workflow_runs.output_data.eval, then DELETE+INSERT the "
        "compliance_call_runs projection row (what the dashboard reads). Dry-run unless --yes is given.",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help="Actually commit --persist writes to the database (irreversible). Without it, --persist only previews.",
    )
    parser.add_argument("--json-out", type=Path, default=None, help="Write per-run summary JSON here.")
    parser.add_argument(
        "--use-system-trust",
        action="store_true",
        help="LOCAL ONLY: verify TLS via the OS trust store (truststore) for corporate MITM proxies "
        "like Zscaler. Off by default; also enabled by RESCORE_USE_SYSTEM_TRUST=1. Not needed in staging/prod.",
    )
    parser.add_argument("-q", "--quiet", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.WARNING if args.quiet else logging.INFO)

    _maybe_inject_system_trust(
        enabled=args.use_system_trust
        or (os.environ.get("RESCORE_USE_SYSTEM_TRUST", "").strip().lower() in _TRUTHY)
    )

    if not args.reuse_stored_eval and not (settings.openai_api_key or "").strip():
        log.error("OPENAI_API_KEY is not set (.env or environment). Use --reuse-stored-eval to skip the LLM.")
        return 1

    # Single event loop for the whole run: the module-level async engine binds its pool to
    # one loop, so nested asyncio.run() calls would orphan asyncpg connections at teardown.
    return asyncio.run(_amain(args))


async def _amain(args: argparse.Namespace) -> int:
    try:
        return await _process(args)
    finally:
        await engine.dispose()


async def _process(args: argparse.Namespace) -> int:
    runs = await _load_runs(args)
    if not runs:
        log.error("No matching compliance_call runs found for the given filters.")
        return 1
    log.info("Loaded %d candidate run(s).", len(runs))

    rubric = load_glp_rubric() if args.detailed else {}
    rv = (args.rubric_version or "").strip() or None

    persist_commit = bool(args.persist and args.yes)
    if args.persist:
        if persist_commit:
            log.warning(
                "PERSIST + --yes: will UPDATE workflow_runs.output_data.eval and DELETE+INSERT "
                "compliance_call_runs for %d run(s). This is irreversible.",
                len(runs),
            )
        else:
            log.warning("PERSIST dry-run (no --yes): computing DB changes and rolling back — nothing is written.")

    rollup_rows: list[list[str]] = []
    detailed_rows: list[list[str]] | None = [] if args.detailed else None
    row_transcripts: list[str] = []  # parallel to rollup_rows; used only when --with-transcript
    results: list[dict[str, Any]] = []
    persisted_count = 0
    all_ok = True

    for r in runs:
        inp = r.input_data if isinstance(r.input_data, dict) else {}
        od = r.output_data if isinstance(r.output_data, dict) else {}
        source = _input_str(inp, "source") or ("mysql_call" if r.mysql_call_id is not None else "gdrive")
        dname = _input_str(inp, "doctor_name")
        slug = _input_str(inp, "doctor_slug")
        conv_raw = inp.get("mysql_second_opinion_conversation_id")
        serial_no = sheet_serial_no_for_state(
            ingest_source=source, mysql_second_opinion_conversation_id=conv_raw
        )
        dg_summary = od.get("deepgram_summary") if isinstance(od.get("deepgram_summary"), dict) else {}
        transcript = _stored_transcript(od, args.transcript_source)

        log.info("Run %s doctor=%s conv=%s transcript_chars=%d", r.id, dname or slug, serial_no or "-", len(transcript))

        try:
            if args.reuse_stored_eval:
                eval_doc = od.get("eval") if isinstance(od.get("eval"), dict) else {}
                if not eval_doc:
                    raise RuntimeError("no stored eval on run (drop --reuse-stored-eval to re-run the rubric)")
            else:
                if not transcript:
                    raise RuntimeError("no stored transcript on run (nothing to re-score)")
                eval_doc = await run_rubric_eval(
                    transcript_text=transcript,
                    doctor_slug=slug,
                    doctor_name=dname,
                    rubric_version=rv,
                    deepgram_summary=dg_summary,
                )
            rollup_rows.append(
                _rollup_row(
                    serial_no=serial_no,
                    doctor_name=dname or slug,
                    eval_doc=eval_doc,
                    deepgram_summary=dg_summary,
                )
            )
            if detailed_rows is not None:
                detailed_rows.append(
                    _detailed_row(
                        rubric=rubric,
                        serial_no=serial_no,
                        doctor_name=dname or slug,
                        eval_doc=eval_doc,
                        deepgram_summary=dg_summary,
                    )
                )
            row_transcripts.append(transcript)  # keep aligned with rollup_rows
        except Exception as e:
            log.exception("Failed on run %s", r.id)
            all_ok = False
            results.append({"workflow_run_id": str(r.id), "ok": False, "error": f"{type(e).__name__}: {e}"[:2000]})
            continue

        persist_info: dict[str, Any] | None = None
        if args.persist:
            try:
                persist_info = await _persist_rescored(
                    r.id,
                    eval_doc=eval_doc,
                    reuse_stored_eval=bool(args.reuse_stored_eval),
                    commit=persist_commit,
                )
                if persist_info.get("persisted"):
                    persisted_count += 1
                log.info(
                    "  persist run=%s %s: %s → %s",
                    r.id,
                    "COMMITTED" if persist_info.get("persisted") else "dry-run",
                    persist_info.get("prev"),
                    persist_info.get("new"),
                )
            except Exception as e:
                log.exception("Persist failed on run %s", r.id)
                all_ok = False
                persist_info = {"persisted": False, "error": f"{type(e).__name__}: {e}"[:2000]}

        results.append(
            {
                "workflow_run_id": str(r.id),
                "ok": True,
                "conversation_id": serial_no,
                "doctor": dname or slug,
                "source": source,
                "created_at": r.created_at.isoformat() if r.created_at else None,
                "reused_stored_eval": bool(args.reuse_stored_eval),
                "eval": {
                    "composite_pct": eval_doc.get("composite_pct"),
                    "grade": eval_doc.get("grade"),
                    "grade_label": eval_doc.get("grade_label"),
                    "model": eval_doc.get("model"),
                },
                "persist": persist_info,
            }
        )

    if args.rollup_out and rollup_rows:
        out_path = args.rollup_out.expanduser().resolve()
        _write_rollup_xlsx(
            out_path,
            rollup_rows=rollup_rows,
            detailed_rows=detailed_rows,
            rubric=rubric,
            transcripts=row_transcripts if args.with_transcript else None,
        )
        log.info(
            "Wrote rollup workbook: %s (%d row(s)%s)",
            out_path,
            len(rollup_rows),
            ", +Transcript col" if args.with_transcript else "",
        )

    summary = {
        "runs_processed": len(results),
        "rows_built": len(rollup_rows),
        "all_ok": all_ok,
        "persist": {"enabled": bool(args.persist), "committed": persist_commit, "rows_written": persisted_count},
        "reuse_stored_eval": bool(args.reuse_stored_eval),
        "transcript_source": args.transcript_source,
        "filters": {
            "created_after": args.created_after,
            "created_before": args.created_before,
            "conversation_id": args.conversation_id,
            "workflow_run_id": args.workflow_run_id,
            "doctor": args.doctor,
            "source_type": args.source_type,
            "any_status": args.any_status,
            "limit": args.limit or None,
        },
        "rollup_out": str(args.rollup_out.expanduser().resolve()) if (args.rollup_out and rollup_rows) else None,
        "results": results,
    }
    if args.json_out:
        args.json_out.write_text(json.dumps(summary, indent=2), encoding="utf-8")
        log.info("Wrote %s", args.json_out)

    print(
        json.dumps(
            {
                "all_ok": all_ok,
                "runs": len(results),
                "rows": len(rollup_rows),
                "persist_committed": persist_commit,
                "persisted_rows": persisted_count,
            },
            indent=2,
        )
    )
    return 0 if all_ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
