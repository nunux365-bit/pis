#!/usr/bin/env python3
"""
Validate GLP reference artifacts (README §2 / §3) without calling Deepgram or OpenAI.

Checks (data + deterministic code only):
1. ``manifest.json`` order vs ``glp1_consult_rollup.xlsx`` / ``Detailed`` row count.
2. For each call: ``*.txt`` == canonical string from ``*.deepgram.json`` (same as
   ``canonical_transcript_from_deepgram_response``).
3. Structural columns (Rollup cols 5–10) from JSON vs Excel (matches README: Dr:Pt = NA when no diarize).
4. Rollup item columns (28) → recomputed domain % + composite via ``apply_deterministic_scores``
   vs Excel (§13 tolerance).
5. ``Detailed`` sheet: 93 columns; first 20 headers match Rollup; per-item Status columns match Rollup.

Usage (from ``agentos-backend/``)::

  PYTHONPATH=. python scripts/glp_validate_reference_workbook.py \\
    --transcripts-dir /path/to/_transcripts-deepgram \\
    --rollup-xlsx /path/to/glp1_consult_rollup.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from openpyxl import load_workbook  # noqa: E402

from app.agents.compliance_call.adapters.transcribe_deepgram import (  # noqa: E402
    canonical_transcript_from_deepgram_response,
)
from app.agents.compliance_call.glp_scoring import (  # noqa: E402
    apply_deterministic_scores,
    load_glp_rubric,
    structural_from_deepgram,
)
from app.agents.compliance_call.glp_sheet_rows import (  # noqa: E402
    ROLLUP_SCORED_ITEM_IDS,
    glp_rollup_headers,
)


def _load_manifest(transcripts_dir: Path) -> list[dict[str, Any]]:
    mpath = transcripts_dir / "manifest.json"
    if not mpath.is_file():
        raise FileNotFoundError(mpath)
    return json.loads(mpath.read_text(encoding="utf-8"))["files"]


def _stem_from_manifest_json(rel_json: str) -> str:
    return Path(rel_json).name.replace(".deepgram.json", "")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--transcripts-dir", type=Path, required=True)
    ap.add_argument("--rollup-xlsx", type=Path, required=True)
    ap.add_argument("--rubric-json", type=Path, default=None)
    args = ap.parse_args()

    td = args.transcripts_dir.expanduser().resolve()
    xlsx = args.rollup_xlsx.expanduser().resolve()
    rubric = load_glp_rubric(args.rubric_json) if args.rubric_json else load_glp_rubric()

    files = _load_manifest(td)
    stems = [_stem_from_manifest_json(f["json"]) for f in files]

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    wsr, wsd = wb["Rollup"], wb["Detailed"]
    hr = list(next(wsr.iter_rows(min_row=4, max_row=4, values_only=True)))
    hd = list(next(wsd.iter_rows(min_row=4, max_row=4, values_only=True)))
    r4r = [str(x).strip() if x is not None else "" for x in hr]
    r4d = [str(x).strip() if x is not None else "" for x in hd]

    issues: list[str] = []
    if list(glp_rollup_headers()) != r4r:
        issues.append("Rollup row 4 headers differ from glp_rollup_headers()")
    if len(r4d) != 93:
        issues.append(f"Detailed row 4: expected 93 columns, got {len(r4d)}")
    if r4r[:20] != r4d[:20]:
        issues.append("Rollup/Detailed first 20 header cells differ")

    data_rows = list(wsr.iter_rows(min_row=5, max_row=100, values_only=True))
    n_excel = 0
    for row in data_rows:
        vals = list(row)[:49]
        if all(v is None or (isinstance(v, str) and not str(v).strip()) for v in vals):
            break
        n_excel += 1

    if n_excel != len(stems):
        issues.append(f"manifest files ({len(stems)}) vs excel data rows ({n_excel})")

    col_by = {r4r[i]: i for i in range(len(r4r)) if r4r[i]}
    start_items = r4r.index("Safety: MTC/MEN2")

    for i, stem in enumerate(stems):
        jf = td / f"{stem}.deepgram.json"
        tf = td / f"{stem}.txt"
        if not jf.is_file():
            issues.append(f"missing {jf.name}")
            continue
        payload = json.loads(jf.read_text(encoding="utf-8"))
        canon = canonical_transcript_from_deepgram_response(payload)
        if tf.is_file():
            ref_txt = tf.read_text(encoding="utf-8").strip()
            if canon != ref_txt:
                issues.append(f"transcript mismatch stem={stem!r} len {len(canon)} vs {len(ref_txt)}")

        st = structural_from_deepgram(payload, diarize=False)
        ri = 5 + i
        xl = list(next(wsr.iter_rows(min_row=ri, max_row=ri, values_only=True)))
        pairs = [
            ("Duration (min)", st["duration_minutes"], xl[col_by["Duration (min)"]]),
            ("Total Turns", st["total_speaking_turns"], xl[col_by["Total Turns"]]),
            ("Avg Turn (s)", st["avg_turn_length_seconds"], xl[col_by["Avg Turn (s)"]]),
            ("Longest Turn (s)", st["longest_turn_seconds"], xl[col_by["Longest Turn (s)"]]),
            ("Speech Activity (%)", st["speech_activity_pct"], xl[col_by["Speech Activity (%)"]]),
            ("Dr:Pt Ratio", st["doctor_to_patient_ratio"], xl[col_by["Dr:Pt Ratio"]]),
        ]
        for name, a, b in pairs:
            if isinstance(b, float) and isinstance(a, (int, float)):
                if round(float(a), 4) != round(float(b), 4):
                    issues.append(f"{stem} structural {name}: py={a} xlsx={b}")
            else:
                if str(a).strip() != str(b).strip():
                    issues.append(f"{stem} structural {name}: py={a!r} xlsx={b!r}")

        status_by_id: dict[str, str] = {}
        for j, iid in enumerate(ROLLUP_SCORED_ITEM_IDS):
            v = xl[start_items + j]
            status_by_id[iid] = str(v or "").strip().upper() or "NO"
        det = apply_deterministic_scores(rubric, status_by_id)
        for dk, h in [
            ("safety_screening", "Safety (%)"),
            ("clinical_assessment", "Clinical (%)"),
            ("psychosocial_lifestyle", "Psychosoc (%)"),
            ("drug_education_informed_consent", "Drug Ed (%)"),
            ("nutritional_counseling", "Nutrition (%)"),
            ("followup_planning_monitoring", "Follow-up (%)"),
        ]:
            ex = float(xl[col_by[h]])
            py = float(det["domain_pcts"].get(dk) or 0.0)
            if abs(ex - py) > 0.05:
                issues.append(f"row {ri} domain {dk}: excel={ex} recomputed={py}")
        ex_c = float(xl[col_by["Composite (%)"]])
        if abs(ex_c - float(det["composite_pct"])) > 0.11:
            issues.append(f"row {ri} composite: excel={ex_c} recomputed={det['composite_pct']}")

        # Detailed vs Rollup item statuses
        dr = list(next(wsd.iter_rows(min_row=ri, max_row=ri, values_only=True)))
        for j in range(28):
            if xl[start_items + j] != dr[20 + 2 * j]:
                issues.append(f"row {ri} item j={j} rollup vs detailed status mismatch")

    out = {"ok": len(issues) == 0, "issue_count": len(issues), "issues": issues}
    print(json.dumps(out, indent=2))
    return 0 if not issues else 1


if __name__ == "__main__":
    raise SystemExit(main())
