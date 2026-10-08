"""Document parsing for PDF, DOCX, TXT, and MD files."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    """Result of document parsing."""
    filename: str
    content: str
    page_count: int | None = None
    metadata: dict | None = None


def parse_document(file_path: str | Path) -> ParsedDocument:
    """
    Parse a document and extract text content.

    Supports: PDF, DOCX, TXT, MD

    Args:
        file_path: Path to the document

    Returns:
        ParsedDocument with extracted content

    Raises:
        ValueError: If file type is not supported
        FileNotFoundError: If file doesn't exist
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()
    filename = path.name

    if suffix == ".pdf":
        return _parse_pdf(path, filename)
    elif suffix == ".docx":
        return _parse_docx(path, filename)
    elif suffix in (".txt", ".md", ".markdown"):
        return _parse_text(path, filename)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _parse_pdf(path: Path, filename: str) -> ParsedDocument:
    """Parse PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("PyMuPDF (fitz) is required for PDF parsing. Install with: pip install pymupdf")

    doc = fitz.open(path)
    pages = []

    for page in doc:
        text = page.get_text("text")
        if text.strip():
            pages.append(text)

    content = "\n\n".join(pages)
    page_count = len(doc)
    doc.close()

    log.debug("Parsed PDF %s: %d pages, %d chars", filename, page_count, len(content))

    return ParsedDocument(
        filename=filename,
        content=content,
        page_count=page_count,
        metadata={"source_type": "pdf"},
    )


def _parse_docx(path: Path, filename: str) -> ParsedDocument:
    """Parse DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")

    doc = Document(path)
    paragraphs = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    content = "\n\n".join(paragraphs)

    log.debug("Parsed DOCX %s: %d paragraphs, %d chars", filename, len(paragraphs), len(content))

    return ParsedDocument(
        filename=filename,
        content=content,
        metadata={"source_type": "docx"},
    )


def _parse_text(path: Path, filename: str) -> ParsedDocument:
    """Parse plain text or markdown files."""
    content = path.read_text(encoding="utf-8")

    log.debug("Parsed text %s: %d chars", filename, len(content))

    source_type = "markdown" if path.suffix.lower() in (".md", ".markdown") else "text"

    return ParsedDocument(
        filename=filename,
        content=content,
        metadata={"source_type": source_type},
    )
