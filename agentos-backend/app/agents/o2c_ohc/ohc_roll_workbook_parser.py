"""
Parse OHC attendance workbooks using the same layout as ``generate_ohc_json.py`` (Downloads/agentos).

Supports:

- **Dual tabs:** ``OHC Attendance On Roll`` + ``OHC Attendance Off Roll`` (legacy).
- **Unified tab:** ``OHC Attendance`` alone — same wide-row shape as Off Roll (meta through col 14,
  dates from col 15, trailing summary metrics). Used when dual tabs are absent.

Date headers may be ``datetime`` / ``date`` cells or strings such as ``1-Apr-26``.

Produces an in-memory :class:`OHCAttendanceDataset` and can convert to
``ClientSiteAttendanceMap`` for attendance → MIS pipelines.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from app.agents.o2c_ohc.attendance_summary import AttendanceRecord, ClientSiteAttendanceMap
from app.agents.o2c_ohc.duplicate_day_cell_merge import collapse_duplicate_day_cells

# Mirrors generate_ohc_json.py ON_ROLL_CFG / OFF_ROLL_CFG
ON_ROLL_SHEET = "OHC Attendance On Roll"
OFF_ROLL_SHEET = "OHC Attendance Off Roll"
# Single-sheet layout (On+Off merged); column map matches production unified export (Apr 2026 style).
UNIFIED_OHC_ATTENDANCE_SHEET = "OHC Attendance"

ON_ROLL_CFG: dict[str, Any] = {
    "sheet": ON_ROLL_SHEET,
    "meta_cols": {
        "ohc_name": 1,
        "name": 2,
        "employee_id": 3,
        "contact_no": 4,
        "state": 5,
        "role": 6,
        "doj": 7,
        "status": 8,
        "lwd": 9,
        "payroll_type": 10,
        "frequency": 11,
        "kam": 12,
        "cluster_leads": 13,
    },
    "date_start_col": 14,
    "summary_cols": {
        "target_present_days": 45,
        "actual_present_days": 46,
        "week_off": 47,
        "half_day": 48,
        "leave": 49,
        "holiday": 50,
        "locum_duty": 51,
        "absent": 52,
        "total_present_days": 53,
        "ot": 54,
    },
}

OFF_ROLL_CFG: dict[str, Any] = {
    "sheet": OFF_ROLL_SHEET,
    "meta_cols": {
        "ohc_name": 1,
        "name": 2,
        "employee_id": 3,
        "contact_no": 4,
        "state": 5,
        "role": 6,
        "doj": 7,
        "status": 8,
        "lwd": 9,
        "payroll_type": 10,
        "frequency": 11,
        "kam": 12,
        "vendor_name": 13,
        "ops_manager": 14,
    },
    "date_start_col": 15,
    "summary_cols": {
        "target_present_days": 46,
        "actual_present_days": 47,
        "week_off": 48,
        "half_day": 49,
        "leave": 50,
        "holiday": 51,
        "locum_duty": 52,
        "absent_days": 53,
        "present_days_last_month": 54,
        "present_days_current_month": 55,
        "total_present_days": 56,
        "ot": 57,
    },
}

UNIFIED_ROLL_CFG: dict[str, Any] = {
    "sheet": UNIFIED_OHC_ATTENDANCE_SHEET,
    "meta_cols": {
        "ohc_name": 1,
        "name": 2,
        "employee_id": 3,
        "contact_no": 4,
        "state": 5,
        "role": 6,
        "doj": 7,
        "status": 8,
        "lwd": 9,
        "payroll_type": 10,
        "frequency": 11,
        "kam": 12,
        "vendor_name": 13,
        "cluster_leads": 14,
    },
    "date_start_col": 15,
    "summary_cols": {
        "target_present_days": 46,
        "actual_present_days": 47,
        "week_off": 48,
        "half_day": 49,
        "holiday": 50,
        "locum_duty": 51,
        "attendance_not_marked": 52,
        "absent": 53,
        "ot": 54,
    },
}

ATTENDANCE_TAXONOMY: dict[str, Any] = {
    "present_variants": [
        {"pattern": "Present ( G )", "shift": "G", "counts_as": "present"},
        {"pattern": "Present ( M )", "shift": "M", "counts_as": "present"},
        {"pattern": "Present ( E )", "shift": "E", "counts_as": "present"},
        {"pattern": "Present ( N )", "shift": "N", "counts_as": "present"},
        {"pattern": "Present", "shift": None, "counts_as": "present"},
    ],
    "locum_variants": [
        {"pattern": "Locum Duty ( G ) / <name>", "shift": "G", "counts_as": "locum_duty"},
        {"pattern": "Locum Duty ( M ) / <name>", "shift": "M", "counts_as": "locum_duty"},
        {"pattern": "Locum Duty ( E ) / <name>", "shift": "E", "counts_as": "locum_duty"},
        {"pattern": "Locum Duty ( N ) / <name>", "shift": "N", "counts_as": "locum_duty"},
    ],
    "non_present": [
        {"value": "Week Off", "counts_as": "week_off"},
        {"value": "Leave", "counts_as": "leave"},
        {"value": "Maternity Leave", "counts_as": "leave"},
        {"value": "Holiday", "counts_as": "week_off"},
        {"value": "Absent", "counts_as": "absent"},
        {"value": "Left", "counts_as": "not_applicable"},
    ],
}


def _cell_str(ws: Any, row: int, col: int) -> str:
    v = ws.cell(row, col).value
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _cell_num(ws: Any, row: int, col: int) -> Any:
    v = ws.cell(row, col).value
    if v is None or (isinstance(v, str) and not str(v).strip()):
        return ""
    try:
        n = float(v)
        return int(n) if n == int(n) else n
    except (ValueError, TypeError):
        return str(v).strip()


def _iso_date_header_value(v: Any) -> str | None:
    """Normalize header row cell to ``YYYY-MM-DD`` for date columns; else ``None``."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            d = from_excel(v)
            if isinstance(d, datetime):
                return d.strftime("%Y-%m-%d")
            if isinstance(d, date):
                return d.isoformat()
        except Exception:
            pass
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        for fmt in ("%d-%b-%y", "%d-%b-%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                parsed = datetime.strptime(s, fmt)
                return parsed.strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def _detect_dates(ws: Any, header_row: int, start_col: int, max_cols: int = 120) -> list[tuple[int, str]]:
    """
    Contiguous date columns from ``start_col`` (row ``header_row``).

    Accepts ``date`` / ``datetime`` / Excel serial / common string headers (e.g. ``1-Apr-26``).
    Stops at the first non-empty cell that is not a parseable date (e.g. summary headers).
    """
    dates: list[tuple[int, str]] = []
    for c in range(start_col, start_col + max_cols):
        v = ws.cell(header_row, c).value
        iso = _iso_date_header_value(v)
        if iso:
            dates.append((c, iso))
            continue
        if v is None or (isinstance(v, str) and not str(v).strip()):
            if dates:
                break
            continue
        break
    return dates


def _parse_sheet(wb: Any, cfg: dict[str, Any], source_label: str) -> dict[str, list[dict[str, Any]]]:
    from collections import defaultdict

    if cfg["sheet"] not in wb.sheetnames:
        return {}
    ws = wb[cfg["sheet"]]
    header_row = 2
    data_start = 3
    date_cols = _detect_dates(ws, header_row, cfg["date_start_col"])
    if not date_cols:
        return {}

    # Same calendar day can appear in multiple columns (e.g. wide template + duplicate block).
    # Merge left-to-right with blank preserved — matches ``roll_to_summary_export``.
    cols_by_date: dict[str, list[int]] = defaultdict(list)
    for col, date_str in date_cols:
        cols_by_date[date_str].append(col)

    ohc_staff: dict[str, list[dict[str, Any]]] = defaultdict(list)
    max_row = ws.max_row or data_start

    for row in range(data_start, max_row + 1):
        ohc = _cell_str(ws, row, cfg["meta_cols"]["ohc_name"])
        name = _cell_str(ws, row, cfg["meta_cols"]["name"])
        if not ohc or not name:
            continue
        meta = {k: _cell_str(ws, row, c) for k, c in cfg["meta_cols"].items()}
        meta["source"] = source_label
        # Back-compat: keep non-blank attendance exactly as before.
        attendance: dict[str, str] = {}
        # New: keep all detected dates (including blanks as None) for date-wise metrics.
        attendance_all_days: dict[str, str | None] = {}
        for date_str in sorted(cols_by_date.keys()):
            raw_vals = [ws.cell(row, c).value for c in cols_by_date[date_str]]
            merged = collapse_duplicate_day_cells(raw_vals)
            attendance_all_days[date_str] = merged
            if merged:
                attendance[date_str] = merged
        summary = {k: _cell_num(ws, row, c) for k, c in cfg["summary_cols"].items()}
        meta["attendance"] = attendance
        meta["attendance_all_days"] = attendance_all_days
        meta["summary"] = summary
        ohc_staff[ohc].append(meta)

    return dict(ohc_staff)


def workbook_has_roll_tabs(path: Path | str) -> bool:
    p = Path(path).expanduser().resolve()
    wb = load_workbook(p, read_only=True)
    try:
        names = set(wb.sheetnames)
        return ON_ROLL_SHEET in names and OFF_ROLL_SHEET in names
    finally:
        wb.close()


def workbook_has_unified_ohc_attendance_tab(path: Path | str) -> bool:
    """True when a single ``OHC Attendance`` wide tab is present without dual roll tabs."""
    p = Path(path).expanduser().resolve()
    wb = load_workbook(p, read_only=True)
    try:
        names = set(wb.sheetnames)
        if ON_ROLL_SHEET in names and OFF_ROLL_SHEET in names:
            return False
        return UNIFIED_OHC_ATTENDANCE_SHEET in names
    finally:
        wb.close()


def _bucket_status(v: Any) -> str:
    """
    Bucket a per-day attendance status (date-by-date only; never use Excel summary columns).
    - absent, week_off, leave, blank are explicit buckets.
    - Leave matches substring ``leave`` (case-insensitive), except when clearly ``present``.
    - Blank is empty / None only.
    - Holiday is treated as week_off.
    - ``Left`` is ``not_applicable`` (see ``ATTENDANCE_TAXONOMY``); it does not add to present_days.
    - Everything else (Present, Locum, OT, unknown non-empty text) counts as **present**.
    """
    if v is None:
        return "blank"
    s = str(v).strip()
    if not s:
        return "blank"
    sl = s.strip().lower()
    sl_flat = re.sub(r"[\s_-]+", "", sl)
    if "halfday" in sl_flat and "present" not in sl:
        return "half_day"
    if "absent" in sl and "present" not in sl:
        return "absent"
    if "week off" in sl or sl.replace(" ", "") == "weekoff" or "holiday" in sl:
        return "week_off"
    if "leave" in sl and "present" not in sl:
        return "leave"
    if sl == "left":
        return "not_applicable"
    return "present"


def _daywise_attendance_counts(att_all: dict[str, Any]) -> dict[str, Decimal]:
    """
    Per-person metrics from the full date axis in ``attendance_all_days`` only.
    present_days = all days that are not absent, week_off, leave, blank, or not_applicable.
    expected_days = total_days - week_off_days (computed from the same loop).
    """
    counts = {
        "present_days": Decimal("0"),
        "absent_days": Decimal("0"),
        "week_off_days": Decimal("0"),
        "leave_days": Decimal("0"),
        "blank_days": Decimal("0"),
    }
    for v in att_all.values():
        b = _bucket_status(v)
        if b == "absent":
            counts["absent_days"] += Decimal("1")
        elif b == "half_day":
            counts["present_days"] += Decimal("0.5")
            counts["absent_days"] += Decimal("0.5")
        elif b == "week_off":
            counts["week_off_days"] += Decimal("1")
        elif b == "leave":
            counts["leave_days"] += Decimal("1")
        elif b == "blank":
            counts["blank_days"] += Decimal("1")
        elif b == "not_applicable":
            pass
        else:
            counts["present_days"] += Decimal("1")
    return counts


def _dates_inclusive(d0: date, d1: date) -> list[date]:
    if d1 < d0:
        return [d0]
    out: list[date] = []
    cur = d0
    while cur <= d1:
        out.append(cur)
        cur = cur + timedelta(days=1)
    return out


def _staff_to_attendance_record(site_key: str, staff: dict[str, Any]) -> AttendanceRecord:
    summary = staff.get("summary") or {}
    att_all = staff.get("attendance_all_days") or staff.get("attendance") or {}
    counts = _daywise_attendance_counts(att_all)
    total_days = len(att_all)
    expected_days = max(0, Decimal(total_days) - counts["week_off_days"])
    emp_id = (staff.get("employee_id") or "").strip() or (staff.get("name") or "").strip()
    return {
        "client_site_key": site_key,
        "employee_external_id": emp_id[:120],
        "employee_name": (staff.get("name") or "").strip() or None,
        "role_code": (staff.get("role") or "").strip() or None,
        "lwd": (staff.get("lwd") or "").strip() or None,
        "doj": (staff.get("doj") or "").strip() or None,
        # Always from date-by-date cells; never from trailing Excel summary columns.
        "present_days": counts["present_days"],
        "absent_days": counts["absent_days"],
        "total_days": total_days,
        "expected_days": expected_days,
        "week_off_days": float(counts["week_off_days"]),
        "leave_days": float(counts["leave_days"]),
        "blank_days": float(counts["blank_days"]),
        "period_filter": "",
        "roll_filter": "",
        "roll_type": (staff.get("source") or "").strip() or None,
        # Day-wise statuses (date->status/None) to support OHC Detailed without reparsing workbook.
        "ohc_attendance_daywise": {k: (v if v not in ("",) else None) for k, v in att_all.items()},
        "ohc_summary_metrics": {k: v for k, v in summary.items() if v != ""},
        "ohc_attendance_days": len(att_all),
    }


@dataclass
class OHCSiteBlock:
    """One client-site (OHC name) with staff rows as produced by the roll tabs."""

    site_key: str
    staff_count: int
    roles: dict[str, int]
    staff: list[dict[str, Any]]


@dataclass
class OHCAttendanceDataset:
    """Full parse: suitable for MIS export and invoice bridge."""

    meta: dict[str, Any]
    taxonomy: dict[str, Any]
    sites: dict[str, OHCSiteBlock]

    def to_client_site_attendance_map(self) -> ClientSiteAttendanceMap:
        out: ClientSiteAttendanceMap = {}
        for key, block in self.sites.items():
            out[key] = [_staff_to_attendance_record(key, s) for s in block.staff]
        return out

    def site_period_attendance_metrics(
        self,
        *,
        site_key: str,
        period_start: date,
        period_end: date,
    ) -> list[dict[str, Any]]:
        """
        Date-by-date counts for MIS inputs (do not use Summary-sheet totals).

        Returns per-person dicts with:
        - employee_external_id, employee_name, role_code, roll_type
        - total_days, expected_days (= total_days - week_off_days)
        - present_days, absent_days, blank_days, week_off_days, leave_days
        """
        block = self.sites.get(site_key)
        if not block:
            return []
        days = _dates_inclusive(period_start, period_end)
        total_days = len(days)
        out: list[dict[str, Any]] = []
        for s in block.staff:
            att_all: dict[str, Any] = s.get("attendance_all_days") or s.get("attendance") or {}
            counts = {
                "present_days": Decimal("0"),
                "absent_days": Decimal("0"),
                "week_off_days": Decimal("0"),
                "leave_days": Decimal("0"),
                "blank_days": Decimal("0"),
            }
            for d in days:
                ds = d.isoformat()
                b = _bucket_status(att_all.get(ds))
                if b == "absent":
                    counts["absent_days"] += Decimal("1")
                elif b == "half_day":
                    counts["present_days"] += Decimal("0.5")
                    counts["absent_days"] += Decimal("0.5")
                elif b == "week_off":
                    counts["week_off_days"] += Decimal("1")
                elif b == "leave":
                    counts["leave_days"] += Decimal("1")
                elif b == "blank":
                    counts["blank_days"] += Decimal("1")
                elif b == "not_applicable":
                    pass
                else:
                    counts["present_days"] += Decimal("1")
            expected_days = max(Decimal("0"), Decimal(total_days) - counts["week_off_days"])
            emp_id = (s.get("employee_id") or "").strip() or (s.get("name") or "").strip() or f"row_{len(out)+1}"
            out.append(
                {
                    "employee_external_id": emp_id[:200],
                    "employee_name": (s.get("name") or "").strip() or None,
                    "role_code": (s.get("role") or "").strip() or None,
                    "roll_type": (s.get("source") or "").strip() or None,
                    "total_days": total_days,
                    "expected_days": expected_days,
                    **counts,
                }
            )
        return out

    def site_period_detailed_json(
        self,
        *,
        site_key: str,
        period_start: date,
        period_end: date,
    ) -> dict[str, Any]:
        """Detailed grid for OHC Detailed sheet using the parsed roll attendance."""
        block = self.sites.get(site_key)
        if not block:
            return {"ok": True, "site": site_key, "employees": [], "dates": [], "grid": {}}
        days = _dates_inclusive(period_start, period_end)
        dates = [d.isoformat() for d in days]
        employees: list[dict[str, Any]] = []
        grid: dict[str, dict[str, str | None]] = {}
        # Make employee name keys unique to avoid collisions.
        seen: dict[str, int] = {}
        for s in block.staff:
            name = (s.get("name") or "").strip() or (s.get("employee_id") or "").strip() or "employee"
            role = (s.get("role") or "").strip()
            src = (s.get("source") or "").strip()
            k = name
            n = seen.get(k, 0) + 1
            seen[k] = n
            if n > 1:
                k = f"{name} ({n})"
            employees.append({"name": k, "role": role, "roll_source": src})
            att_all: dict[str, Any] = s.get("attendance_all_days") or s.get("attendance") or {}
            grid[k] = {
                ds: (att_all.get(ds) if att_all.get(ds) not in ("", None) else None) for ds in dates
            }
        return {
            "ok": True,
            "site": site_key,
            "period": {"start": period_start.isoformat(), "end": period_end.isoformat()},
            "employees": employees,
            "dates": dates,
            "grid": grid,
        }

    def mis_rollup_by_site_role(self) -> dict[str, Any]:
        """Aggregates for MIS-style reporting (client-site × role)."""
        rollup: dict[str, Any] = {}
        for site_key, block in self.sites.items():
            by_role: dict[str, dict[str, Decimal]] = {}
            for s in block.staff:
                role = (s.get("role") or "").strip() or "_unknown"
                rec = _staff_to_attendance_record(site_key, s)
                bucket = by_role.setdefault(
                    role,
                    {
                        "headcount": Decimal(0),
                        "present_days_sum": Decimal(0),
                        "absent_days_sum": Decimal(0),
                    },
                )
                bucket["headcount"] += Decimal(1)
                bucket["present_days_sum"] += _decimal(rec.get("present_days"))
                bucket["absent_days_sum"] += _decimal(rec.get("absent_days"))
            rollup[site_key] = {
                "staff_count": block.staff_count,
                "by_role": {k: {x: float(y) for x, y in v.items()} for k, v in by_role.items()},
            }
        return rollup

    def to_json_dict(self) -> dict[str, Any]:
        """Shape close to ``generate_ohc_json.py`` output (for export / debug)."""
        ohc_sites: dict[str, Any] = {}
        for k, block in self.sites.items():
            ohc_sites[k] = {
                "staff_count": block.staff_count,
                "roles": block.roles,
                "staff": block.staff,
            }
        return {
            "_meta": self.meta,
            "attendance_taxonomy": self.taxonomy,
            "ohc_sites": ohc_sites,
        }


def _decimal(v: Any) -> Decimal:
    if v is None:
        return Decimal(0)
    if isinstance(v, Decimal):
        return v
    try:
        return Decimal(str(v))
    except Exception:
        return Decimal(0)


def _build_ohc_attendance_dataset(
    p: Path,
    combined: dict[str, list[dict[str, Any]]],
    *,
    parser_meta: str,
) -> OHCAttendanceDataset:
    from collections import defaultdict

    if not combined:
        raise ValueError("No OHC attendance rows parsed.")

    all_dates: set[str] = set()
    for lst in combined.values():
        for s in lst:
            all_dates.update((s.get("attendance") or {}).keys())
    sorted_dates = sorted(all_dates)
    billing_period = sorted_dates[0][:7] if sorted_dates else "unknown"

    unique_roles = sorted(
        {
            (s.get("role") or "").strip()
            for staff_rows in combined.values()
            for s in staff_rows
            if (s.get("role") or "").strip()
        }
    )
    role_order = {role: idx for idx, role in enumerate(unique_roles)}
    sites: dict[str, OHCSiteBlock] = {}
    for ohc in sorted(combined.keys()):
        staff = combined[ohc]
        staff.sort(
            key=lambda s: (role_order.get((s.get("role") or "").strip(), 99), s.get("name", ""))
        )
        roles: dict[str, int] = {}
        for s in staff:
            r = (s.get("role") or "").strip() or ""
            roles[r] = roles.get(r, 0) + 1
        sites[ohc] = OHCSiteBlock(
            site_key=ohc,
            staff_count=len(staff),
            roles=dict(sorted({k: v for k, v in roles.items()}.items())),
            staff=staff,
        )

    meta = {
        "source_file": p.name,
        "billing_period": billing_period,
        "total_ohc_sites": len(sites),
        "total_staff": sum(b.staff_count for b in sites.values()),
        "date_range": (
            {"from": sorted_dates[0], "to": sorted_dates[-1]} if sorted_dates else {}
        ),
        "parser": parser_meta,
    }
    return OHCAttendanceDataset(meta=meta, taxonomy=ATTENDANCE_TAXONOMY, sites=sites)


def parse_ohc_roll_workbook(path: Path | str) -> OHCAttendanceDataset:
    """
    Parse On Roll + Off Roll sheets. Raises ``ValueError`` if date columns cannot be found.
    """
    p = Path(path).expanduser().resolve()
    if not p.is_file():
        raise FileNotFoundError(str(p))
    wb = load_workbook(p, data_only=True)
    try:
        if ON_ROLL_SHEET not in wb.sheetnames or OFF_ROLL_SHEET not in wb.sheetnames:
            raise ValueError(
                f"Workbook missing roll tabs: need {ON_ROLL_SHEET!r} and {OFF_ROLL_SHEET!r}"
            )
        on_roll = _parse_sheet(wb, ON_ROLL_CFG, "On Roll")
        off_roll = _parse_sheet(wb, OFF_ROLL_CFG, "Off Roll")
        if not on_roll and not off_roll:
            raise ValueError("No date columns / no data parsed from On Roll or Off Roll tabs.")
    finally:
        wb.close()

    from collections import defaultdict

    combined: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for ohc, lst in on_roll.items():
        combined[ohc].extend(lst)
    for ohc, lst in off_roll.items():
        combined[ohc].extend(lst)

    return _build_ohc_attendance_dataset(
        p, dict(combined), parser_meta="ohc_roll_workbook_parser_v1"
    )


def parse_unified_ohc_roll_workbook(path: Path | str) -> OHCAttendanceDataset:
    """
    Parse a single **OHC Attendance** sheet (unified on+off layout). Raises ``ValueError`` if
    no date columns or no data rows.
    """
    p = Path(path).expanduser().resolve()
    if not p.is_file():
        raise FileNotFoundError(str(p))
    wb = load_workbook(p, data_only=True)
    try:
        if UNIFIED_OHC_ATTENDANCE_SHEET not in wb.sheetnames:
            raise ValueError(f"Workbook missing unified tab {UNIFIED_OHC_ATTENDANCE_SHEET!r}")
        unified = _parse_sheet(wb, UNIFIED_ROLL_CFG, "OHC Attendance")
        if not unified:
            raise ValueError("No date columns / no data parsed from unified OHC Attendance tab.")
    finally:
        wb.close()

    return _build_ohc_attendance_dataset(
        p, dict(unified), parser_meta="ohc_roll_workbook_parser_unified_v1"
    )


def try_parse_ohc_roll_workbook(path: Path | str) -> OHCAttendanceDataset | None:
    """Dual roll tabs first; else unified ``OHC Attendance`` tab."""
    p = Path(path).expanduser().resolve()
    if workbook_has_roll_tabs(p):
        try:
            return parse_ohc_roll_workbook(p)
        except (ValueError, OSError):
            return None
    if workbook_has_unified_ohc_attendance_tab(p):
        try:
            return parse_unified_ohc_roll_workbook(p)
        except (ValueError, OSError):
            return None
    return None
