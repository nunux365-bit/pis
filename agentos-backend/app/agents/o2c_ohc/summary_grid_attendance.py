"""
Google Sheets–style OHC Summary: data validation on Summary!A1 (period), B1 (client-site), E1 (roll).

**Intended parse policy** (product default):
1. Set **first** (A1) and **third** (E1) dropdowns to **All**, recalculate, then read **Manpower!AB**
   for the client–site list.
2. For **each** client-site, set **B1**, recalculate, parse the grid.
3. **Row 2** = role label per column (stored as ``role_code`` / ``summary_row2_role``); **row 3** =
   employee name; **column A** from row 4 = dates; body cells = attendance status.

**openpyxl** cannot recalculate dynamic arrays — use ``parse_summary_formula_grid_snapshot`` only for
the **last saved** B1 (and optionally attach metadata that period/roll were normalized for your
pipeline). Full multi-site requires **xlwings** + Microsoft Excel (``try_attendance_all_sites_xlwings``).

Contract ↔ attendance **role** and **client_site** string matching is intentionally out of scope here.
"""

from __future__ import annotations

import logging
import re
import zipfile
from collections.abc import Callable
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from app.agents.o2c_ohc.attendance_summary import AttendanceRecord, ClientSiteAttendanceMap
from app.config.settings import settings

log = logging.getLogger(__name__)


def _local(tag: str) -> str:
    return tag.split("}", 1)[-1] if tag.startswith("{") else tag


def all_period_label() -> str:
    return (settings.o2c_attendance_all_period_label or "All").strip()


def all_roll_label() -> str:
    return (settings.o2c_attendance_all_roll_label or "All").strip()


def summary_sheet_name() -> str:
    return (settings.o2c_attendance_summary_sheet or "Summary").strip()


def detect_summary_dropdown_workbook(xlsx_path: Path) -> bool:
    """True if Summary sheet has list validations on A1, B1, and E1 (OHC pattern)."""
    path = xlsx_path.expanduser().resolve()
    with zipfile.ZipFile(path, "r") as zf:
        try:
            raw = zf.read("xl/worksheets/sheet1.xml")
        except KeyError:
            return False
    root = ET.fromstring(raw)
    found: set[str] = set()
    for dv in root.iter():
        if _local(dv.tag) != "dataValidation":
            continue
        if dv.get("type") != "list":
            continue
        sq = (dv.get("sqref") or "").replace("$", "").upper()
        if sq in ("A1", "B1", "E1"):
            found.add(sq)
    return found >= {"A1", "B1", "E1"}


def iter_client_sites_from_manpower(xlsx_path: Path, *, max_row: int = 5000) -> list[str]:
    """Unique non-empty strings from Manpower column AB (B1 list source in the sample file).

    **Note:** With A1/E1 = All, this list can differ from a file saved under other filters.
    Prefer refreshing sites **after** forcing All in Excel (see xlwings path).
    """
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        if "Manpower" not in wb.sheetnames:
            return []
        ws = wb["Manpower"]
        seen: set[str] = set()
        out: list[str] = []
        for row in ws.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=28,
            max_col=28,
            values_only=True,
        ):
            v = row[0] if row else None
            if v is None:
                continue
            s = str(v).strip()
            if not s or s in seen:
                continue
            seen.add(s)
            out.append(s)
        return out
    finally:
        wb.close()


def _count_present_like(status: Any) -> bool:
    if status is None:
        return False
    t = str(status).lower()
    return "present" in t or "locum" in t


def _looks_like_employee_name_cell(value: Any) -> bool:
    """
    Google/Excel exports often add extra columns (codes, phones, states) with row 3 filled.
    Keep columns that look like a person name, not metadata.
    """
    if value is None:
        return False
    s = str(value).strip()
    if len(s) < 3:
        return False
    if re.fullmatch(r"[\d\s\-+().]+", s):
        return False
    if s.upper().startswith("HLD-"):
        return False
    if not re.search(r"[A-Za-z]", s):
        return False
    low = s.lower()
    if low in frozenset({"maharashtra", "gujarat", "karnataka", "all"}):
        return False
    # Bare job titles as row-3 "name" are usually header spill, not people
    if low in frozenset({"nurse", "doctor", "driver"}):
        return False
    return True


def _cell_at_matrix(data: list[Any], r: int, c: int) -> Any:
    """1-based Excel row/col into matrix from xlwings range (0-based list)."""
    ri, ci = r - 1, c - 1
    if ri < 0 or ci < 0 or ri >= len(data):
        return None
    row = data[ri]
    if row is None or not isinstance(row, (list, tuple)) or ci >= len(row):
        return None
    return row[ci]


def parse_summary_grid_with_cell_getter(
    cell: Callable[[int, int], Any],
    *,
    site_key: str,
    period_filter: str,
    roll_filter: str,
    role_row: int = 2,
    name_row: int = 3,
    data_start_row: int = 4,
    max_row: int = 400,
    max_col: int = 60,
    filter_non_employee_columns: bool = True,
) -> list[AttendanceRecord]:
    """
    Parse Summary attendance using **row ``role_row``** as role and **row ``name_row``** as employee.

    ``role_code`` / ``summary_row2_role`` hold the same value (row 2 in the standard layout).
    """
    records: list[AttendanceRecord] = []
    for col in range(2, max_col + 1):
        emp = cell(name_row, col)
        if emp is None or not str(emp).strip():
            continue
        if filter_non_employee_columns and not _looks_like_employee_name_cell(emp):
            continue
        role_raw = cell(role_row, col)
        role_s = str(role_raw).strip() if role_raw is not None else None
        present_days = Decimal(0)
        for row in range(data_start_row, max_row + 1):
            d = cell(row, 1)
            if d is None:
                continue
            status = cell(row, col)
            if _count_present_like(status):
                present_days += Decimal(1)
        records.append(
            {
                "client_site_key": site_key,
                "employee_external_id": str(emp).strip()[:120],
                "employee_name": str(emp).strip(),
                # Row 2 in OHC Summary = role label (map to contract role_code later).
                "role_code": role_s,
                "summary_row2_role": role_s,
                "present_days": present_days,
                "absent_days": None,
                "period_filter": period_filter,
                "roll_filter": roll_filter,
                "roll_type": roll_filter,
            }
        )
    return records


def _records_from_summary_ws(
    ws: Any,
    site_key: str,
    period_filter: str,
    roll_filter: str,
    *,
    max_row: int | None = None,
    max_col: int | None = None,
) -> list[AttendanceRecord]:
    mr = max_row if max_row is not None else min(ws.max_row or 400, 500)
    mc = max_col if max_col is not None else min(ws.max_column or 60, 80)

    def cell(r: int, c: int) -> Any:
        return ws.cell(r, c).value

    return parse_summary_grid_with_cell_getter(
        cell,
        site_key=site_key,
        period_filter=period_filter,
        roll_filter=roll_filter,
        max_row=mr,
        max_col=mc,
    )


def parse_summary_formula_grid_snapshot(
    xlsx_path: Path,
    *,
    sheet_name: str | None = None,
    normalize_period_roll_to_all_in_output: bool = False,
) -> ClientSiteAttendanceMap | None:
    """
    Last saved calculated grid for **current** Summary!B1 (openpyxl cannot recalc).

    If ``normalize_period_roll_to_all_in_output`` is True, each record's ``period_filter`` /
    ``roll_filter`` / ``roll_type`` are set to the configured **All** labels (for downstream keys)
    even when A1/E1 in the file differ — **the numeric grid is still whatever was last saved**.
    """
    path = xlsx_path.expanduser().resolve()
    if not detect_summary_dropdown_workbook(path):
        return None

    sh = sheet_name or summary_sheet_name()
    wb = load_workbook(path, data_only=True)
    try:
        if sh not in wb.sheetnames:
            return None
        ws = wb[sh]
        site = ws["B1"].value
        if site is None or not str(site).strip():
            return None
        site_key = str(site).strip()
        a1 = ws["A1"].value
        e1 = ws["E1"].value
        p_eff = all_period_label() if normalize_period_roll_to_all_in_output else (
            str(a1).strip() if a1 is not None else all_period_label()
        )
        r_eff = all_roll_label() if normalize_period_roll_to_all_in_output else (
            str(e1).strip() if e1 is not None else all_roll_label()
        )
        recs = _records_from_summary_ws(ws, site_key, p_eff, r_eff)
        return {site_key: recs}
    finally:
        wb.close()


def _unique_ab_from_xlwings(manpower_sheet: Any, *, max_row: int = 4000) -> list[str]:
    vals = manpower_sheet.range((1, 28), (max_row, 28)).value
    if vals is None:
        return []
    flat: list[Any]
    if not isinstance(vals, (list, tuple)):
        flat = [vals]
    else:
        flat = []
        for row in vals:
            if isinstance(row, (list, tuple)):
                flat.extend(row)
            else:
                flat.append(row)
    seen: set[str] = set()
    out: list[str] = []
    for v in flat:
        if v is None:
            continue
        s = str(v).strip()
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
    return out


def try_attendance_all_sites_xlwings(
    xlsx_path: Path,
    *,
    sheet_name: str | None = None,
    max_sites: int | None = None,
    manpower_max_row: int = 4000,
    grid_max_row: int = 400,
    grid_max_col: int = 60,
) -> ClientSiteAttendanceMap | None:
    """
    Production path for Google-Sheet-style Summary:

    1. Set **A1** and **E1** to **All** (labels from settings).
    2. Recalculate, then collect unique **Manpower!AB** sites.
    3. For each site, set **B1**, recalculate, parse grid (row 2 = role, row 3 = name).

    Requires ``pip install xlwings`` and Microsoft Excel.
    """
    try:
        import xlwings as xw
    except ImportError:
        return None

    sh = sheet_name or summary_sheet_name()
    path = str(xlsx_path.expanduser().resolve())
    pall = all_period_label()
    rall = all_roll_label()

    app = xw.App(visible=False)
    combined: dict[str, list[AttendanceRecord]] = {}
    wb = None
    try:
        wb = app.books.open(path)
        summary = wb.sheets[sh]
        manpower = wb.sheets["Manpower"]

        summary.range("A1").value = pall
        summary.range("E1").value = rall
        wb.app.calculate()

        sites = _unique_ab_from_xlwings(manpower, max_row=manpower_max_row)
        if max_sites is not None:
            sites = sites[: max(0, max_sites)]

        for site in sites:
            summary.range("B1").value = site
            wb.app.calculate()
            data = summary.range((1, 1), (grid_max_row, grid_max_col)).value
            if not data:
                combined[str(site).strip()] = []
                continue
            if not isinstance(data[0], (list, tuple)):
                data = [data]

            def cell(r: int, c: int) -> Any:
                return _cell_at_matrix(data, r, c)

            recs = parse_summary_grid_with_cell_getter(
                cell,
                site_key=str(site).strip(),
                period_filter=pall,
                roll_filter=rall,
                max_row=grid_max_row,
                max_col=grid_max_col,
            )
            combined[str(site).strip()] = recs
    except Exception as e:
        log.warning("xlwings multi-site attendance failed: %s", e)
        combined = {}
    finally:
        try:
            if wb is not None:
                wb.close()
        except Exception:
            pass
        try:
            app.quit()
        except Exception:
            pass

    return combined if combined else None


def parse_google_sheet_attendance(
    xlsx_path: Path,
    *,
    prefer_xlwings: bool = True,
    max_sites_xlwings: int | None = None,
) -> tuple[ClientSiteAttendanceMap, dict[str, Any]]:
    """
    Preferred entry point: full **All / All / iterate B1** via xlwings when possible; else snapshot.

    Returns ``(map, meta)`` where ``meta`` includes ``mode`` and ``warnings``.
    """
    path = xlsx_path.expanduser().resolve()
    meta: dict[str, Any] = {"mode": None, "warnings": []}

    if prefer_xlwings:
        m = try_attendance_all_sites_xlwings(path, max_sites=max_sites_xlwings)
        if m is not None:
            meta["mode"] = "xlwings_force_all_all_iterate_client_site"
            meta["site_count"] = len(m)
            return m, meta

    snap = parse_summary_formula_grid_snapshot(
        path,
        normalize_period_roll_to_all_in_output=True,
    )
    if snap is not None:
        meta["mode"] = "openpyxl_snapshot_single_client_site"
        meta["warnings"].append(
            "Only Summary!B1 as last saved; A1/E1 normalized to All in record metadata only. "
            "Install xlwings + Excel for full multi-site with forced All/All."
        )
        return snap, meta

    meta["mode"] = "none"
    meta["warnings"].append("Workbook does not match OHC Summary dropdown layout.")
    return {}, meta
