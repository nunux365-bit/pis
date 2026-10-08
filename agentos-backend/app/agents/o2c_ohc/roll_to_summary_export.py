"""
Build a Summary-style grid from **OHC Attendance On Roll** + **OHC Attendance Off Roll** wide tabs.

Each output worksheet is one client–site (OHC Name). Layout mirrors Summary: row 1 A1/B1/E1,
row 2 = role, row 3 = ``Date`` + employee names, column A from row 4 = calendar dates,
body = collapsed status text. A **footer** block lists per-employee day counts (present, absent,
week off, leave, locum, **blank**, other). Blanks are not colored and not counted as present.

Duplicate date columns (two Excel columns per calendar day) are merged with a fixed priority
so totals align with the live Summary tab (validated against TACO sample).
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.agents.o2c_ohc.duplicate_day_cell_merge import collapse_duplicate_day_cells
from app.agents.o2c_ohc.summary_grid_attendance import (
    all_period_label,
    all_roll_label,
    _looks_like_employee_name_cell,
)

ON_ROLL_SHEET = "OHC Attendance On Roll"
OFF_ROLL_SHEET = "OHC Attendance Off Roll"

# Fixed columns (1-based): A=OHC Name, B=Name, F=Role (both roll sheets in sample file)
COL_SITE = 1
COL_NAME = 2
COL_ROLE = 6
HEADER_DATE_ROW = 2
DATA_START_ROW = 3

# Semantic fills (original workbook uses conditional formatting — not stored in cell XML).
# Palette aligned with common Google Sheets / Excel attendance templates.
_FILL_FILTER_ROW = PatternFill(patternType="solid", fgColor="5B9BD5")
_FILL_ROLE_ROW = PatternFill(patternType="solid", fgColor="BDD7EE")
_FILL_NAME_ROW = PatternFill(patternType="solid", fgColor="DEEBF7")
_FILL_DATE_COL = PatternFill(patternType="solid", fgColor="F2F2F2")
_FILL_PRESENT = PatternFill(patternType="solid", fgColor="C6EFCE")
_FILL_ABSENT = PatternFill(patternType="solid", fgColor="FFC7CE")
_FILL_WEEK_OFF = PatternFill(patternType="solid", fgColor="D9D9D9")
_FILL_LEAVE = PatternFill(patternType="solid", fgColor="FFEB9C")
_FILL_LOCUM = PatternFill(patternType="solid", fgColor="9BC2E6")
_FILL_SUMMARY_BAND = PatternFill(patternType="solid", fgColor="E2E2E2")

_FONT_HEADER = Font(bold=True, color="FFFFFF")
_FONT_GRID = Font(size=10)
_FONT_SUMMARY_LABEL = Font(bold=True, size=10)
_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _attendance_bucket(v: Any) -> str:
    """
    Classify a cell for totals. **Blank is blank** — never treated as present.
    Unknown non-empty text is ``other`` (not present).
    """
    if v is None:
        return "blank"
    s = str(v).strip()
    if not s:
        return "blank"
    sl = s.lower().strip()
    if "absent" in sl and "present" not in sl:
        return "absent"
    # Leave should match `%leave%` (case-insensitive)
    if "leave" in sl and "present" not in sl:
        return "leave"
    if "locum" in sl:
        return "locum"
    if "week off" in sl or sl.replace(" ", "") == "weekoff":
        return "week_off"
    if "present" in sl:
        return "present"
    return "other"


def _fill_for_attendance_value(v: Any) -> PatternFill | None:
    """Only color **explicit** statuses. Blanks and unknown text stay uncolored (no implied present)."""
    b = _attendance_bucket(v)
    if b == "present":
        return _FILL_PRESENT
    if b == "absent":
        return _FILL_ABSENT
    if b == "week_off":
        return _FILL_WEEK_OFF
    if b == "leave":
        return _FILL_LEAVE
    if b == "locum":
        return _FILL_LOCUM
    return None


def _excel_safe_sheet_title(site: str, used: set[str]) -> str:
    """Excel sheet name ≤31 chars; no : \\ / ? * [ ]"""
    s = re.sub(r'[:\\/?*\[\]]', "-", (site or "site").strip()) or "site"
    if len(s) > 31:
        s = s[:31]
    base = s
    n = 2
    while s.lower() in used:
        suffix = f"_{n}"
        s = (base[: 31 - len(suffix)] + suffix) if len(base) + len(suffix) > 31 else base + suffix
        n += 1
    used.add(s.lower())
    return s


def _norm_status(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


def _row2_date_columns(ws: Any, *, max_col: int = 200) -> dict[date, list[int]]:
    """Map calendar date -> list of 1-based column indices (handles duplicate day columns)."""
    by_day: dict[date, list[int]] = defaultdict(list)
    for c in range(1, max_col + 1):
        v = ws.cell(HEADER_DATE_ROW, c).value
        d: date | None = None
        if isinstance(v, datetime):
            d = v.date()
        elif isinstance(v, date):
            d = v
        if d:
            by_day[d].append(c)
    return dict(by_day)


def _dates_in_month(dates: list[date], year: int, month: int) -> list[date]:
    return sorted(d for d in dates if d.year == year and d.month == month)


@dataclass
class RollEmployee:
    name: str
    role: str
    roll_source: str  # "On Roll" | "Off Roll"
    by_date: dict[date, str | None] = field(default_factory=dict)


def _parse_roll_sheet(
    ws: Any,
    roll_source: str,
    *,
    year_month: tuple[int, int] | None,
) -> dict[str, list[RollEmployee]]:
    """site_key -> list of RollEmployee (one row per sheet row)."""
    by_day_cols = _row2_date_columns(ws)
    all_dates = sorted(by_day_cols.keys())
    if year_month:
        y, m = year_month
        grid_dates = _dates_in_month(all_dates, y, m)
    else:
        grid_dates = all_dates

    by_site: dict[str, list[RollEmployee]] = defaultdict(list)

    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if not row:
            continue
        site = row[COL_SITE - 1] if len(row) >= COL_SITE else None
        if site is None or not str(site).strip():
            continue
        site_key = str(site).strip()
        name = str(row[COL_NAME - 1]).strip() if len(row) >= COL_NAME and row[COL_NAME - 1] else ""
        role_cell = row[COL_ROLE - 1] if len(row) >= COL_ROLE else None
        role = str(role_cell).strip() if role_cell is not None else ""

        emp = RollEmployee(name=name, role=role or "", roll_source=roll_source)
        for d in grid_dates:
            cols = by_day_cols.get(d, [])
            vals = []
            for c in cols:
                if c - 1 < len(row):
                    vals.append(row[c - 1])
            emp.by_date[d] = collapse_duplicate_day_cells(vals)
        by_site[site_key].append(emp)

    return dict(by_site)


def _merge_site_employees(
    on_map: dict[str, list[RollEmployee]],
    off_map: dict[str, list[RollEmployee]],
) -> dict[str, list[RollEmployee]]:
    sites = sorted(set(on_map.keys()) | set(off_map.keys()))
    out: dict[str, list[RollEmployee]] = {}
    for site in sites:
        merged: list[RollEmployee] = []
        for emp in sorted(on_map.get(site, []), key=lambda e: e.name.lower()):
            merged.append(emp)
        for emp in sorted(off_map.get(site, []), key=lambda e: e.name.lower()):
            merged.append(emp)
        out[site] = merged
    return out


def _write_footer_summary(
    ws: Worksheet,
    employees: list[RollEmployee],
    dates: list[date],
    *,
    footer_start_row: int,
    ncols: int,
    apply_colors: bool,
) -> None:
    """Block below the calendar: per-employee day counts (blank ≠ present)."""
    title_row = footer_start_row
    ws.cell(title_row, 1).value = "Attendance summary (day counts)"
    if apply_colors:
        for c in range(1, ncols + 1):
            cell = ws.cell(title_row, c)
            cell.fill = _FILL_SUMMARY_BAND
            cell.font = Font(bold=True, size=11)
            cell.alignment = _ALIGN_LEFT if c == 1 else _ALIGN
    else:
        ws.cell(title_row, 1).font = Font(bold=True, size=11)

    metrics: tuple[tuple[str, str], ...] = (
        ("present", "Present days"),
        ("absent", "Absent days"),
        ("week_off", "Week off days"),
        ("leave", "Leave days"),
        ("locum", "Locum days"),
        ("blank", "Blank / unmarked"),
        ("other", "Other / unclassified"),
    )

    r = title_row + 1
    ws.cell(r, 1).value = "Metric"
    for j, emp in enumerate(employees, start=2):
        ws.cell(r, j).value = emp.name if emp.name else ""
    if apply_colors:
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.fill = _FILL_NAME_ROW
            cell.font = Font(bold=True, size=10)
            cell.alignment = _ALIGN_LEFT if c == 1 else _ALIGN

    for key, label in metrics:
        r += 1
        ws.cell(r, 1).value = label
        for j, emp in enumerate(employees, start=2):
            cnt = sum(1 for d in dates if _attendance_bucket(emp.by_date.get(d)) == key)
            ws.cell(r, j).value = int(cnt)
        if apply_colors:
            ws.cell(r, 1).fill = _FILL_DATE_COL
            ws.cell(r, 1).font = _FONT_SUMMARY_LABEL
            ws.cell(r, 1).alignment = _ALIGN_LEFT
            for j in range(2, ncols + 1):
                cell = ws.cell(r, j)
                cell.font = _FONT_GRID
                cell.alignment = _ALIGN


def _write_summary_style_sheet(
    ws: Worksheet,
    site: str,
    employees: list[RollEmployee],
    dates: list[date],
    *,
    apply_colors: bool = True,
) -> None:
    ap = all_period_label()
    ar = all_roll_label()
    ncols = max(1, len(employees) + 1)
    last_data_row = 3 + len(dates)
    footer_start = last_data_row + 2

    ws.cell(1, 1).value = ap
    ws.cell(1, 2).value = site
    ws.cell(1, 5).value = ar
    ws.cell(3, 1).value = "Date"
    for j, emp in enumerate(employees, start=2):
        ws.cell(2, j).value = emp.role if emp.role else None
        ws.cell(3, j).value = emp.name if emp.name else None
    for i, d in enumerate(dates, start=4):
        ws.cell(i, 1).value = d
        for j, emp in enumerate(employees, start=2):
            v = emp.by_date.get(d)
            ws.cell(i, j).value = v

    _write_footer_summary(
        ws,
        employees,
        dates,
        footer_start_row=footer_start,
        ncols=ncols,
        apply_colors=apply_colors,
    )

    if not apply_colors:
        ws.freeze_panes = "A4"
        return

    ws.freeze_panes = "A4"

    # Row 1: filter band (A1: last col)
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = _FILL_FILTER_ROW
        cell.font = _FONT_HEADER
        cell.alignment = _ALIGN
    # Row 2–3 headers
    for c in range(2, ncols + 1):
        ws.cell(2, c).fill = _FILL_ROLE_ROW
        ws.cell(2, c).font = Font(bold=True, size=10)
        ws.cell(2, c).alignment = _ALIGN
        ws.cell(3, c).fill = _FILL_NAME_ROW
        ws.cell(3, c).font = Font(bold=True, size=10)
        ws.cell(3, c).alignment = _ALIGN
    corner = ws.cell(3, 1)
    corner.fill = _FILL_NAME_ROW
    corner.font = Font(bold=True, size=10)
    corner.alignment = _ALIGN
    # Body: date column + status cells (blank → no fill)
    for i in range(4, last_data_row + 1):
        dc = ws.cell(i, 1)
        dc.fill = _FILL_DATE_COL
        dc.font = _FONT_GRID
        dc.alignment = _ALIGN
        for j in range(2, ncols + 1):
            cell = ws.cell(i, j)
            fv = _fill_for_attendance_value(cell.value)
            if fv is not None:
                cell.fill = fv
            cell.font = _FONT_GRID
            cell.alignment = _ALIGN


def export_roll_tabs_to_summary_workbook(
    src_xlsx: Path | str,
    out_xlsx: Path | str,
    *,
    year_month: tuple[int, int] | None = None,
    sites: set[str] | None = None,
    apply_colors: bool = True,
) -> dict[str, Any]:
    """
    Read On Roll + Off Roll from ``src_xlsx``; write ``out_xlsx`` with one sheet per site.

    ``year_month`` — if ``(2026, 2)``, only February 2026 dates appear in the grid (matches
    typical Summary month scope). If ``None``, every date found in row 2 headers is included.

    ``apply_colors`` — attendance-style fills (present/absent/week off/leave/locum). The source
    file typically uses **conditional formatting** (not stored per-cell fills), so we apply a
    matching semantic palette.
    """
    src = Path(src_xlsx).expanduser().resolve()
    out = Path(out_xlsx).expanduser().resolve()
    wb_in = load_workbook(src, read_only=True, data_only=True)
    try:
        if ON_ROLL_SHEET not in wb_in.sheetnames or OFF_ROLL_SHEET not in wb_in.sheetnames:
            raise ValueError(
                f"Expected sheets {ON_ROLL_SHEET!r} and {OFF_ROLL_SHEET!r} in {src}"
            )
        on_map = _parse_roll_sheet(wb_in[ON_ROLL_SHEET], "On Roll", year_month=year_month)
        off_map = _parse_roll_sheet(wb_in[OFF_ROLL_SHEET], "Off Roll", year_month=year_month)
    finally:
        wb_in.close()

    merged = _merge_site_employees(on_map, off_map)
    if sites is not None:
        merged = {k: v for k, v in merged.items() if k in sites}

    # Shared date axis: union of dates any employee has for that site, sorted
    wb_out = Workbook()
    default_ws = wb_out.active
    wb_out.remove(default_ws)
    used_titles: set[str] = set()
    sheet_count = 0

    for site in sorted(merged.keys()):
        emps = merged[site]
        if not emps:
            continue
        date_set: set[date] = set()
        for e in emps:
            date_set.update(e.by_date.keys())
        dates = sorted(date_set)
        if not dates:
            continue
        title = _excel_safe_sheet_title(site, used_titles)
        ws = wb_out.create_sheet(title=title)
        _write_summary_style_sheet(ws, site, emps, dates, apply_colors=apply_colors)
        sheet_count += 1

    wb_out.save(out)
    return {
        "out": str(out),
        "sheet_count": sheet_count,
        "year_month": year_month,
        "apply_colors": apply_colors,
    }


def _status_strings_equal(sa: str, ba: str) -> bool:
    """Loose equality for Summary vs roll export (spacing, week-off casing, Present vs Present ( G ) )."""
    if sa == ba:
        return True

    def week_off_norm(x: str) -> bool:
        z = x.replace(" ", "").replace("-", "")
        return z == "weekoff"

    if week_off_norm(sa) and week_off_norm(ba):
        return True
    if "absent" in sa and "absent" in ba and "present" not in sa and "present" not in ba:
        return True
    if "leave" in sa and "leave" in ba and "present" not in sa and "present" not in ba:
        return True
    if "present" in sa and "present" in ba and "absent" not in sa and "absent" not in ba:
        return True
    if "locum" in sa and "locum" in ba:
        return True
    return False


def validate_site_against_summary_tab(
    src_xlsx: Path | str,
    built_ws: Worksheet,
    *,
    summary_sheet: str = "Summary",
) -> dict[str, Any]:
    """
    Compare ``built_ws`` (our roll-derived grid) to ``Summary`` in the same workbook where
    Summary!B1 matches ``built_ws['B1']``. Returns counts of matching / mismatched body cells.
    """
    src = Path(src_xlsx).expanduser().resolve()
    site_built = built_ws.cell(1, 2).value
    site_built = str(site_built).strip() if site_built else ""

    # read_only=False so max_column / max_row are set (required to scan row 3 names).
    wb = load_workbook(src, read_only=False, data_only=True)
    try:
        if summary_sheet not in wb.sheetnames:
            return {"ok": False, "error": f"No sheet {summary_sheet!r}"}
        ws = wb[summary_sheet]
        b1 = ws.cell(1, 2).value
        b1s = str(b1).strip() if b1 else ""
        if b1s != site_built:
            return {
                "ok": False,
                "error": f"Summary!B1={b1s!r} != built B1={site_built!r} (open correct site sheet)",
                "summary_b1": b1s,
                "built_b1": site_built,
            }

        # Map name -> col (skip junk columns — same heuristic as Summary grid parser)
        name_to_col: dict[str, int] = {}
        bmax = min(built_ws.max_column or 60, 80)
        for c in range(2, bmax + 1):
            nm = built_ws.cell(3, c).value
            if not nm or not str(nm).strip():
                continue
            ns = str(nm).strip()
            if _looks_like_employee_name_cell(ns) and ns not in name_to_col:
                name_to_col[ns] = c

        summary_name_col: dict[str, int] = {}
        smax = min(ws.max_column or 60, 80)
        for c in range(2, smax + 1):
            nm = ws.cell(3, c).value
            if not nm or not str(nm).strip():
                continue
            ns = str(nm).strip()
            # First occurrence only — dynamic-array spill can repeat names in far columns (junk).
            if _looks_like_employee_name_cell(ns) and ns not in summary_name_col:
                summary_name_col[ns] = c

        # Map date -> row in built (col A from row 4)
        built_date_row: dict[date, int] = {}
        for r in range(4, min(built_ws.max_row or 0, 500) + 1):
            v = built_ws.cell(r, 1).value
            d = v.date() if isinstance(v, datetime) else (v if isinstance(v, date) else None)
            if d:
                built_date_row[d] = r

        match = mismatch = missing = skipped_blank_summary = 0
        details: list[dict[str, Any]] = []

        for r in range(4, min((ws.max_row or 0) + 1, 500)):
            dv = ws.cell(r, 1).value
            d = dv.date() if isinstance(dv, datetime) else (dv if isinstance(dv, date) else None)
            if not d:
                continue
            br = built_date_row.get(d)
            if br is None:
                missing += 1
                continue
            for name, scol in summary_name_col.items():
                bcol = name_to_col.get(name)
                if bcol is None:
                    continue
                sval = ws.cell(r, scol).value
                bval = built_ws.cell(br, bcol).value
                sa = _norm_status(str(sval)) if sval is not None and str(sval).strip() else ""
                ba = _norm_status(str(bval)) if bval is not None and str(bval).strip() else ""
                if not sa:
                    skipped_blank_summary += 1
                    continue
                if _status_strings_equal(sa, ba):
                    match += 1
                else:
                    mismatch += 1
                    if len(details) < 30:
                        details.append(
                            {
                                "date": str(d),
                                "name": name,
                                "summary": sval,
                                "built": bval,
                            }
                        )

        return {
            "ok": mismatch == 0,
            "site": site_built,
            "cells_matched": match,
            "cells_mismatch": mismatch,
            "skipped_blank_summary": skipped_blank_summary,
            "dates_missing_in_built": missing,
            "date_axis_aligned": missing == 0,
            "sample_mismatches": details,
        }
    finally:
        wb.close()
