"""
Validate whether a standardized MIS .xlsx (e.g. "Standardized MIS Format - OHC.xlsx") exposes
{{mustache}} placeholders that a future exporter could fill from ``o2c_mis_run`` /
``o2c_mis_summary_row`` (or a fixed cell map).

The canonical key set below is a **compatibility checklist** vs older invoice-oriented templates;
programmatic MIS export today builds workbooks in code (see ``mis_drafts._write_mis_xlsx_from_db``).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.agents.o2c_ohc.mis_template_cells import describe_standardized_mis_coverage

_PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


def _canonical_placeholder_keys() -> frozenset[str]:
    """Keys a classic invoice-style MIS template expected (historical checklist)."""
    base = {
        "invoice_number",
        "client_name",
        "client_short",
        "site_name",
        "billing_period_start",
        "billing_period_end",
        "subtotal",
        "service_charge_total",
        "taxable_amount",
        "gst_rate",
        "gst_amount",
        "total_with_gst",
        "tds_rate",
        "tds_amount",
        "net_payable",
        "client_site_key",
    }
    keys = set(base)
    for i in range(1, 25):
        keys.update(
            {
                f"line_{i}_desc",
                f"line_{i}_qty",
                f"line_{i}_unit",
                f"line_{i}_rate",
                f"line_{i}_sc",
                f"line_{i}_total",
            }
        )
    return frozenset(keys)


@dataclass
class MisTemplateAuditReport:
    template_path: str
    sheet_names: list[str] = field(default_factory=list)
    layout_style: str = "unknown"
    standardized_mis_notes: dict[str, Any] | None = None
    placeholders_found: list[str] = field(default_factory=list)
    placeholders_supported: list[str] = field(default_factory=list)
    placeholders_unsupported: list[str] = field(default_factory=list)
    can_fill_all_placeholders: bool = True
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "template_path": self.template_path,
            "sheet_names": self.sheet_names,
            "layout_style": self.layout_style,
            "standardized_mis_notes": self.standardized_mis_notes,
            "placeholders_found": self.placeholders_found,
            "placeholders_supported": self.placeholders_supported,
            "placeholders_unsupported": self.placeholders_unsupported,
            "can_fill_all_placeholders": self.can_fill_all_placeholders,
            "notes": self.notes,
        }


def audit_mis_template(template_path: Path | str) -> MisTemplateAuditReport:
    """
    Scan all worksheets for {{token}} strings and compare to a canonical invoice-style key set.
    """
    path = Path(template_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)

    supported = _canonical_placeholder_keys()
    found: set[str] = set()
    wb = load_workbook(path, data_only=False, read_only=True)
    names: list[str] = []
    try:
        names = list(wb.sheetnames)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=False):
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and "{{" in v:
                        for m in _PLACEHOLDER.finditer(v):
                            found.add(m.group(1))
    finally:
        wb.close()

    unsupported = sorted(found - supported)
    supported_hits = sorted(found & supported)
    notes: list[str] = [
        "Supported keys mirror a legacy invoice-line layout (header + up to 24 line slots).",
        "If the MIS uses fixed cell addresses without {{placeholders}}, add a named-range map "
        "or extend the MIS exporter with explicit cell coordinates.",
    ]
    if unsupported:
        notes.append(
            "Unsupported placeholders require extra queries (e.g. bank details, HSN, PO) or "
            "static template changes."
        )

    layout_style = "mustache_placeholders" if found else "fixed_cells"
    std_notes: dict[str, Any] | None = None
    if not found and any("OHC Summary" in n for n in names):
        std_notes = describe_standardized_mis_coverage()
        notes.append(
            "No {{placeholders}} detected — template likely uses fixed cells. "
            "See standardized_mis_notes and mis_template_cells.py for DB mapping gaps."
        )

    return MisTemplateAuditReport(
        template_path=str(path),
        sheet_names=names,
        layout_style=layout_style,
        standardized_mis_notes=std_notes,
        placeholders_found=sorted(found),
        placeholders_supported=supported_hits,
        placeholders_unsupported=unsupported,
        can_fill_all_placeholders=len(unsupported) == 0,
        notes=notes,
    )
