"""Optional OpenAI fallback when rule-based sheet tab resolution fails.

Invoked only after :func:`~app.email_automation.engine.excel_reader.read_sheet`
raises ``KeyError``. The model must choose **exactly one** tab name from the
workbook's ``sheetnames`` list (copy-paste); output is validated before use.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.config.settings import settings

log = logging.getLogger(__name__)

_PROMPT_VERSION = "sheet_tab_pick_v2_1"

_ROLE_INVOICE = (
    "Main **invoice / receivable** data for this variant (one row per bill line or similar). "
    "Not a small per-key lookup table."
)
_ROLE_LOOKUP = (
    "Secondary **lookup** sheet (e.g. unaccounted revenue from **party-wise ageing**). "
    "Per-party facts — not the main invoice line grid."
)


class _SheetPick(BaseModel):
    physical_sheet: str | None = Field(
        default=None,
        description="Exact tab title copied from the workbook tab list in the user message, or null",
    )


_SYSTEM = """You match a **configured logical Excel tab name** to at most one **actual** tab title from a fixed list.

The user message includes: intended sheet name, parser error context, tab names already tried, sheet **role** (invoice vs lookup), and a JSON array of **all workbook tab titles** under the heading \"Actual tab names in this workbook\".

Rules:
1. Output JSON only: {"physical_sheet": "<string>" | null}
2. If physical_sheet is a string, it MUST be byte-for-byte identical to **one entry in that JSON tab list** (same spelling, spaces, punctuation, case).
3. Return null if no tab reasonably matches the **intended** sheet (role + configured name). Do not pick an unrelated tab just to return something.
4. If one tab is a clear best match (e.g. renamed or shortened title for the same report), return that exact string.
5. Party-wise ageing / unaccounted: prefer party-wise / H&T style tabs over TDS-only or unbilled-only when the role is lookup and the configured name references party-wise ageing."""

_USER_TMPL = """## Task
Rule-based tab resolution **failed** (exact match, normalizations, and any pack **aliases** were already tried). Pick the **single best** real tab for the intended sheet, or **null** if none fit.

## Configured logical sheet we need
{configured_name!r}

## Sheet role in the pipeline
{role_description}

## Exact tab name(s) already tried without success
{attempted_json}

## Parser / resolver error (for context)
{error_excerpt}

## Actual tab names in this workbook (you MUST copy one of these exactly, or return null)
{names_json}

## Run context
workflow_type={workflow_type}
variant_name={variant_name}
workbook_file={workbook_basename}
prompt_version={version}

Return JSON: {{"physical_sheet": "<exact string copied from the list above>" or null}}
"""


def _role_line(sheet_kind: str) -> str:
    k = (sheet_kind or "").strip().lower()
    if k == "lookup":
        return _ROLE_LOOKUP
    if k == "invoice":
        return _ROLE_INVOICE
    return (
        "Sheet role not specified; infer from configured_sheet_name and variant "
        "(invoice/receivable vs lookup / ageing)."
    )


def _list_sheet_names(xlsx_path: Path | str) -> tuple[str, ...]:
    wb = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    try:
        return tuple(wb.sheetnames)
    finally:
        wb.close()


async def try_resolve_physical_sheet_name(
    xlsx_path: Path | str,
    configured_name: str,
    *,
    variant_name: str,
    workflow_type: str,
    attempted_tab_names: tuple[str, ...] = (),
    sheet_kind: str = "",
    last_error_message: str | None = None,
) -> str | None:
    """Return a physical tab name from the workbook, or ``None`` if unavailable."""

    if not settings.email_automation_sheet_name_ai_fallback_enabled:
        return None
    api_key = (settings.openai_api_key or "").strip()
    if not api_key:
        log.warning(
            "sheet_name_ai_fallback: enabled but openai_api_key empty — skipping"
        )
        return None

    observed = _list_sheet_names(xlsx_path)
    if not observed:
        return None

    err_txt = (last_error_message or "").strip()
    if not err_txt:
        err_txt = (
            "KeyError: configured tab name did not match any sheet after "
            "rule-based resolution (exact name, fingerprints, aliases)."
        )
    if len(err_txt) > 2500:
        err_txt = err_txt[:2497] + "..."

    attempted_json = json.dumps(
        list(attempted_tab_names) if attempted_tab_names else [configured_name],
        ensure_ascii=False,
    )

    user = _USER_TMPL.format(
        configured_name=configured_name,
        role_description=_role_line(sheet_kind),
        attempted_json=attempted_json,
        error_excerpt=err_txt,
        names_json=json.dumps(list(observed), ensure_ascii=False),
        workflow_type=workflow_type,
        variant_name=variant_name,
        workbook_basename=Path(xlsx_path).name,
        version=_PROMPT_VERSION,
    )

    model = (
        settings.email_automation_sheet_name_ai_model or "gpt-5.4-mini"
    ).strip()

    try:
        from openai import AsyncOpenAI

        async with AsyncOpenAI(api_key=api_key, timeout=120.0) as client:
            resp = await client.chat.completions.create(
                model=model,
                temperature=0,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": _SYSTEM},
                    {"role": "user", "content": user},
                ],
            )
    except Exception:
        log.exception(
            "sheet_name_ai_fallback: OpenAI call failed workbook=%s",
            Path(xlsx_path).name,
        )
        return None

    choices = getattr(resp, "choices", None) or []
    if not choices:
        log.warning(
            "sheet_name_ai_fallback: empty choices workbook=%s",
            Path(xlsx_path).name,
        )
        return None
    msg = getattr(choices[0], "message", None)
    raw = ((getattr(msg, "content", None) or "") if msg else "").strip()
    try:
        data: dict[str, Any] = json.loads(raw)
        picked = _SheetPick.model_validate(data).physical_sheet
    except Exception:
        log.warning("sheet_name_ai_fallback: invalid JSON from model: %r", raw[:500])
        return None

    if picked is None or not str(picked).strip():
        log.warning(
            "sheet_name_ai_fallback: model returned null/empty for configured=%r workbook=%s",
            configured_name,
            Path(xlsx_path).name,
        )
        return None

    tab = str(picked)
    if tab not in observed:
        log.warning(
            "sheet_name_ai_fallback: model picked non-existent tab %r — rejecting",
            tab,
        )
        return None

    return tab
