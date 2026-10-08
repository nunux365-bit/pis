"""OHC attendance workbook — Summary sheet → in-memory map client–site → list of row dicts.

Excel "dropdown" filters (All / Onroll–Offroll All) are applied in logic: we keep rows where
configured scope columns and the on-roll column are blank or 'All', then group by the client–site column.
Tune hints via settings (o2c_ohc_*_column_hints).
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config.settings import settings


def _split_hints(s: str) -> list[str]:
    return [x.strip().lower() for x in (s or "").split(",") if x.strip()]


def _norm_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    t = str(v).strip()
    return t


def _header_score(row: list[Any]) -> int:
    return sum(1 for c in row if _norm_cell(c))


def _find_header_row(ws, max_scan: int = 40) -> int:
    best_r, best_score = 1, -1
    for r in range(1, min(max_scan, ws.max_row or max_scan) + 1):
        row = [ws.cell(r, c).value for c in range(1, min(ws.max_column or 26, 80) + 1)]
        sc = _header_score(row)
        if sc > best_score:
            best_score, best_r = sc, r
    return best_r


def _match_col_index(headers: list[str], hints: list[str]) -> int | None:
    for i, h in enumerate(headers):
        hl = h.strip().lower()
        for hint in hints:
            if hint in hl:
                return i
    return None


def _all_scope_col_indices(headers: list[str], hints: list[str]) -> list[int]:
    if not hints:
        return []
    out: list[int] = []
    for i, h in enumerate(headers):
        hl = h.strip().lower()
        if any(hint in hl for hint in hints):
            out.append(i)
    return out


def _row_passes_scope_filters(
    cells: list[Any],
    scope_indices: list[int],
    onroll_idx: int | None,
) -> bool:
    for idx in scope_indices:
        if idx >= len(cells):
            continue
        v = _norm_cell(cells[idx]).lower()
        if v and v != "all":
            return False
    if onroll_idx is not None and onroll_idx < len(cells):
        v = _norm_cell(cells[onroll_idx]).lower()
        if v and v != "all":
            return False
    return True


def load_ohc_summary_by_client_site(
    path: str | Path,
    *,
    sheet_name: str | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """
    Reads the Summary tab; returns map keyed by client–site label (string).
    Each value is a list of row dicts (header → cell value).
    """
    p = Path(path).expanduser().resolve()
    if not p.is_file():
        raise FileNotFoundError(str(p))

    sn = (sheet_name or "Summary").strip()
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        if sn not in wb.sheetnames:
            raise ValueError(f"Sheet {sn!r} not in {wb.sheetnames}")
        ws = wb[sn]
        hr = _find_header_row(ws)
        max_col = min(ws.max_column or 1, 120)
        headers_raw = [ws.cell(hr, c).value for c in range(1, max_col + 1)]
        headers = [_norm_cell(h) or f"col_{i}" for i, h in enumerate(headers_raw)]

        site_hints = _split_hints(settings.o2c_ohc_site_column_hints)
        onroll_hints = _split_hints(settings.o2c_ohc_onroll_column_hints)
        scope_hints = _split_hints(settings.o2c_ohc_scope_filter_column_hints)

        site_idx = _match_col_index(headers, site_hints)
        if site_idx is None:
            raise ValueError(
                f"Could not find client–site column using hints {site_hints!r}. Headers: {headers[:30]}"
            )
        onroll_idx = _match_col_index(headers, onroll_hints) if onroll_hints else None
        scope_indices = _all_scope_col_indices(headers, scope_hints)

        out: dict[str, list[dict[str, Any]]] = {}
        for r in range(hr + 1, (ws.max_row or hr) + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, max_col + 1)]
            if not any(_norm_cell(x) for x in row_vals):
                continue
            if not _row_passes_scope_filters(row_vals, scope_indices, onroll_idx):
                continue
            site_key = _norm_cell(row_vals[site_idx]) if site_idx < len(row_vals) else ""
            if not site_key or site_key.lower() == "all":
                continue
            rec = {headers[i]: row_vals[i] if i < len(row_vals) else None for i in range(len(headers))}
            out.setdefault(site_key, []).append(rec)
        return out
    finally:
        wb.close()


def normalize_site_key(label: str) -> str:
    """Stable key for matching attendance export labels to contract site_alias / codes."""
    s = label.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")
