"""End-to-end tests for the aggregator variants — PAYMENT_REMINDER_WEEKLY.

Jul-2026 requirement: the diagnostics aggregator no longer reads the party-level
``Receivable as on`` sheet. It is now two invoice-level variants:

* ``diagnostics_aggregator`` — e-Diagnostic-Aggregator rows on ``Invoice details-H&T``.
  Filter: Business Unit = e-Diagnostic-Aggregator, Remarks ∈ {Invoice, TDS,
  Partial Invoice is pending, …}, Final Amount Pending ≥ 1. Table has a Remarks
  column. Not Due rows kept only when due by month end (Invoice Date + Credit
  Days), relabelled "Due this Month".
* ``platform_aggregator`` — e-Pharmacy (Platform Aggregator) rows on
  ``Invoice wise Agg.``. Filter: ``BU = LGP``, Company Code Currency Value ≥ 1,
  and a ``Month Bucket = Not Due`` row kept only when ``Projected Ageing of Next
  Month`` (AZ) reads "OverDue" — those render as "Due this Month". No Remarks
  column on this variant.

Both carry the party name (Name-HANA) in the subject. The table has no Ageing
column and the summary shows only a single "Amount to be Paid" total (per AR
Jul-2026); the overdue split is still computed internally for the skip gate.

Exercises the full ``build_variant_plans`` pipeline from a real openpyxl workbook
to final rendered plans — no unittest.mock patches.
"""

from __future__ import annotations

import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from app.email_automation.pipeline import process as _process
from app.email_automation.pipeline.process import build_variant_plans
from app.email_automation.workflow_packs.payment_reminder import (
    DIAGNOSTICS_AGGREGATOR_VARIANT,
    PLATFORM_AGGREGATOR_VARIANT,
    PAYMENT_REMINDER_WEEKLY,
)

# Pin "today" to the sample workbook's receivable date so the Not Due month-end
# sieve is deterministic (month = July 2026, month end = 2026-07-31).
_REF_DATE = datetime.date(2026, 7, 8)


@pytest.fixture(autouse=True)
def _pin_reference_date():
    _process._set_projection_reference_date_for_tests(_REF_DATE)
    yield
    _process._set_projection_reference_date_for_tests(None)


# --------------------------------------------------------------------------- #
# Invoice details-H&T (e-Diagnostic-Aggregator)                               #
# --------------------------------------------------------------------------- #

_HT_HEADER: list[str] = [
    "Co", "Code", "HANA Code", "Name of the party", "Brand Name",
    "KAM 2", "RPT", "Business Owner", "Business Unit", "Segment",
    "Credit Days", "Type", "Doc. No.", "Invoice Date", "Invoice No:-",
    "Original Amount from Aug'24", "Net amount pending", "Final Amount Pending",
    "Due Days", "Ageing type", "Remarks", "Consider",
]


def _ht_row(
    *,
    hana: str,
    name: str = "Test Aggregator Party",
    business_unit: str = "e-Diagnostic-Aggregator",
    invoice_date: datetime.date = datetime.date(2026, 4, 13),
    invoice_no: str = "TM2606000053",
    credit_days: int = 30,
    final_pending: float = 80_500.0,
    ageing_type: str = "1-3 Months",
    remarks: str = "Invoice",
) -> dict[str, Any]:
    return {
        "Co": "TATA 1MGH",
        "Code": "OTI-001",
        "HANA Code": hana,
        "Name of the party": name,
        "Brand Name": None,
        "KAM 2": None,
        "RPT": None,
        "Business Owner": "Owner",
        "Business Unit": business_unit,
        "Segment": "Aggregator",
        "Credit Days": credit_days,
        "Type": None,
        "Doc. No.": None,
        "Invoice Date": invoice_date,
        "Invoice No:-": invoice_no,
        "Original Amount from Aug'24": final_pending,
        "Net amount pending": final_pending,
        "Final Amount Pending": final_pending,
        "Due Days": 30,
        "Ageing type": ageing_type,
        "Remarks": remarks,
        "Consider": "Positive",
    }


_PARTY_WISE_HEADER: list[str] = [
    "Code", "Name of the party", "KAM 2", "RPT", "Business Unit", "Segment",
    "Business Owner", "Credit days", "Not Due", "0-1 Months", "1-3 Months",
    "3-6 Months", "6-9 Months", "9-12 Months", "+1yrs", "Grand Total",
    "Unaccounted receipts", "Amount in SAP", "Remarks",
]


def _party_wise_row(
    *,
    code: str,
    business_unit: str = "e-Diagnostic-Aggregator",
    unaccounted: float = 0.0,
) -> list[object]:
    return [
        code, "Test Aggregator Party", None, None, business_unit, "Aggregator",
        None, 30, 0, 0, 0, 0, 0, 0, 0, unaccounted, unaccounted, None, None,
    ]


# --------------------------------------------------------------------------- #
# Invoice wise Agg. (e-Pharmacy (Platform Aggregator))                        #
# --------------------------------------------------------------------------- #

# Mirrors the real sheet: ``Month Bucket`` (AT) is the displayed ageing bucket,
# ``Projected Ageing of Next Month`` (AZ, formerly ``Remarks``) is the overdue
# flag AR uses to rescue "Not Due" rows, and ``BU`` (BA) scopes the variant.
_AGG_HEADER: list[str] = [
    "Company Code", "Customer", "Company Code Currency Value", "Posting Date",
    "Reference", "Search term", "Customer Account: Name 1", "Month Bucket",
    "Ageing Date", "Projected Ageing of Next Month", "BU",
]


def _agg_row(
    *,
    customer: object,
    value: float,
    posting_date: datetime.date = datetime.date(2026, 5, 1),
    reference: str = "HT2606000126",
    search_term: str = "TRUWORTH.",
    name: str = "TRUWORTH HEALTH TECHNOLOGIES",
    month_bucket_at: str = "1-3 Months",
    projected_ageing_az: str | None = "OverDue",
    bu: str | None = "LGP",
) -> list[object]:
    return [
        "1MGH", customer, value, posting_date, reference, search_term, name,
        month_bucket_at, posting_date, projected_ageing_az, bu,
    ]


# --------------------------------------------------------------------------- #
# Workbook + tracker helpers                                                  #
# --------------------------------------------------------------------------- #


def _write_workbook(
    tmp_path: Path,
    *,
    ht_rows: list[dict[str, Any]] | None = None,
    party_wise_rows: list[list[object]] | None = None,
    agg_rows: list[list[object]] | None = None,
    filename: str = "receivables.xlsx",
) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws_ht = wb.create_sheet(title="Invoice details-H&T")
    ws_ht.append([])  # title rows above the header (header detection)
    ws_ht.append([])
    ws_ht.append(_HT_HEADER)
    for row in (ht_rows or []):
        ws_ht.append([row.get(h) for h in _HT_HEADER])

    ws_party = wb.create_sheet(title="Party wise Ageing-H&T")
    ws_party.append(_PARTY_WISE_HEADER)
    for row in (party_wise_rows or []):
        ws_party.append(row)

    ws_agg = wb.create_sheet(title="Invoice wise Agg.")
    ws_agg.append(_AGG_HEADER)
    for row in (agg_rows or []):
        ws_agg.append(row)

    # Filename matters: the email's "as on <date>" is parsed from it. The default
    # carries no date, so most tests render without the clause.
    path = tmp_path / filename
    wb.save(path)
    return path


def _tracker_rows(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for entry in entries:
        row: dict[str, Any] = {"Code": entry.get("Code", ""), "BU": entry.get("BU", "Lab")}
        for i in range(1, 12):
            row[f"Mail ID {i}"] = entry.get(f"Mail ID {i}")
        rows.append(row)
    return rows


async def _build(variant, xlsx: Path, tracker: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return await build_variant_plans(
        pack=PAYMENT_REMINDER_WEEKLY,
        variant=variant,
        xlsx_path=xlsx,
        tracker_rows=tracker,
        period="2026-W28",
    )


# =========================================================================== #
# diagnostics_aggregator (e-Diagnostic-Aggregator)                            #
# =========================================================================== #


async def test_diag_normal_send_one_plan_per_hana(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            _ht_row(hana="8106", name="TATA CAPITAL LIMITED", final_pending=133900.0),
            _ht_row(hana="8106", name="TATA CAPITAL LIMITED", final_pending=105800.0),
        ],
        party_wise_rows=[_party_wise_row(code="8106", unaccounted=200.0)],
    )
    tracker = _tracker_rows([{"Code": "8106", "Mail ID 1": "client@agg.com"}])

    plans = await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker)

    assert len(plans) == 1
    plan = plans[0]
    assert plan["variant"] == "diagnostics_aggregator"
    assert plan["business_key"] == "8106"
    assert plan["row_count"] == 2
    assert plan["skip_gate"] is None
    assert Decimal(plan["totals"]["_net_pending"]) == Decimal("239700")
    assert "client@agg.com" in plan["resolved_to"]
    assert "client@agg.com" not in plan["resolved_cc"]


async def test_diag_subject_includes_name_hana(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[_ht_row(hana="8106", name="TATA CAPITAL LIMITED")],
    )
    tracker = _tracker_rows([{"Code": "8106", "Mail ID 1": "c@agg.com"}])

    plan = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0]

    assert plan["rendered_subject"] == (
        "Payment Reminder – Overdue Invoices | TATA CAPITAL LIMITED | Tata1mg | 8106"
    )


async def test_diag_body_matches_reference_email(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            # Overdue row
            _ht_row(
                hana="8106", name="TATA CAPITAL LIMITED", final_pending=133900.0,
                ageing_type="1-3 Months", remarks="Invoice",
            ),
            # Not Due row due this month (Invoice 2026-06-17 + 30d = 2026-07-17)
            _ht_row(
                hana="8106", name="TATA CAPITAL LIMITED",
                invoice_date=datetime.date(2026, 6, 17), credit_days=30,
                final_pending=345000.0, ageing_type="Not Due", remarks="Invoice",
            ),
        ],
    )
    tracker = _tracker_rows([{"Code": "8106", "Mail ID 1": "c@agg.com"}])
    body = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0]["rendered_body_html"]

    assert "Dear Team," in body
    # Intro names the Business Unit these invoices belong to. (The "as on
    # <date>" clause that follows is asserted in its own test.)
    assert (
        "the following e-Diagnostic-Aggregator invoices that remain overdue "
        "in our records" in body
    )
    # Table columns incl. Remarks, but NO Ageing column (dropped per AR).
    for col in ("HANA Code", "Name-HANA", "Invoice Date", "Invoice Number",
                "Net Amount Pending", "Remarks"):
        assert f">{col}<" in body
    assert ">Ageing<" not in body
    # Upcoming-invoice clause removed from both aggregator variants per AR.
    assert "upcoming invoices listed above" not in body
    # Summary: single "Amount to be Paid" total; breakdown lines removed.
    # (The "Net Amount Pending" table header legitimately remains — assert on
    # the summary labels specifically, which carry a trailing colon.)
    assert "Amount to be Paid:" in body
    assert "Amount Pending:" not in body
    assert "Due this month" not in body
    assert "Overdue Amount Pending" not in body
    assert "Unaccounted Revenue" not in body
    assert "Current Total Overdue" not in body
    assert "{receivable_as_on_date}" not in body


async def test_subject_keeps_hana_as_trailing_token_for_reply_tracker(tmp_path):
    """The HANA code must stay the subject's trailing token.

    Reply Tracker tooling recovers ``business_key`` by parsing the subject, so
    the ``… | Tata1mg | <HANA>`` suffix is a contract, not cosmetics:

    * ``reconstruct_send_anchors_local.py`` regex — a miss means the thread gets
      no ``sent`` anchor and silently never appears in Reply Tracker.
    * ``temp_sent_replies_bounces._parse_hana_code`` — returns the last ``|``
      segment, so a trailing party name would masquerade as the HANA code.

    Both parsers are re-implemented here (they live in ``scripts/``, which is not
    an importable package) and MUST be kept in lock-step with the originals.
    """

    import re as _re

    # Verbatim from scripts/reconstruct_send_anchors_local.py
    subject_key_re = _re.compile(r"Tata1mg\s*\|\s*([A-Za-z0-9._\-/]+)\s*$")

    # Verbatim from scripts/temp_sent_replies_bounces.py::_parse_hana_code
    def parse_hana_code(subject: str) -> str | None:
        parts = [p.strip() for p in subject.split("|")]
        if len(parts) >= 2 and _re.search(r"tata\s*1mg", parts[-2], _re.I):
            return parts[-1] or None
        return parts[-1] or None

    xlsx = _write_workbook(
        tmp_path,
        # A name with a pipe and the brand token — the nastiest realistic input.
        ht_rows=[_ht_row(hana="8106", name="ACME | Tata1mg Labs Pvt Ltd")],
        agg_rows=[_agg_row(customer=1000017221.0, value=5000.0)],
    )
    tracker = _tracker_rows([
        {"Code": "8106", "Mail ID 1": "c@agg.com"},
        {"Code": "1000017221", "Mail ID 1": "t@agg.com"},
    ])

    for variant, expected_key in (
        (DIAGNOSTICS_AGGREGATOR_VARIANT, "8106"),
        (PLATFORM_AGGREGATOR_VARIANT, "1000017221"),
    ):
        subject = (await _build(variant, xlsx, tracker))[0]["rendered_subject"]
        m = subject_key_re.search(subject)
        assert m is not None, f"anchor regex failed to parse {subject!r}"
        assert m.group(1) == expected_key
        assert parse_hana_code(subject) == expected_key


@pytest.mark.parametrize(
    "filename, expected",
    [
        # Real wire shapes: mixed '-'/'.' separators, and a trailing revision
        # marker that puts the date mid-name rather than at the end.
        ("Receivable-08-07.2026.xlsx", "08-Jul-2026"),
        ("Receivable-31.05.2026-v3.xlsx", "31-May-2026"),
        ("Receivables as on 17.06.2026.xlsx", "17-Jun-2026"),
        ("Receivable_01-12-26.xlsx", "01-Dec-2026"),  # 2-digit year
    ],
)
async def test_intro_carries_as_on_date_from_attachment_filename(
    tmp_path, filename, expected
):
    """The as-on date comes from the receivables attachment filename."""

    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[_ht_row(hana="8106")],
        agg_rows=[_agg_row(customer=1000017221.0, value=5000.0)],
        filename=filename,
    )
    tracker = _tracker_rows([
        {"Code": "8106", "Mail ID 1": "c@agg.com"},
        {"Code": "1000017221", "Mail ID 1": "t@agg.com"},
    ])

    for variant, bu in (
        (DIAGNOSTICS_AGGREGATOR_VARIANT, "e-Diagnostic-Aggregator"),
        (PLATFORM_AGGREGATOR_VARIANT, "e-Pharmacy (Platform Aggregator)"),
    ):
        body = (await _build(variant, xlsx, tracker))[0]["rendered_body_html"]
        assert (
            f"the following {bu} invoices that remain overdue in our records "
            f"as on <b>{expected}</b>:" in body
        )


async def test_intro_omits_date_gracefully_when_filename_has_none(tmp_path):
    """Undated filename → sentence still reads correctly (no dangling 'as on')."""

    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[_ht_row(hana="8106")],
        # default filename "receivables.xlsx" carries no date
    )
    tracker = _tracker_rows([{"Code": "8106", "Mail ID 1": "c@agg.com"}])

    body = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0][
        "rendered_body_html"
    ]
    assert (
        "the following e-Diagnostic-Aggregator invoices that remain overdue "
        "in our records:" in body
    )
    assert "as on" not in body
    assert "{receivable_as_on_clause_html}" not in body


async def test_diag_remarks_allowlist_filter(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            _ht_row(hana="KEEP1", remarks="Invoice"),
            _ht_row(hana="KEEP2", remarks="TDS"),
            _ht_row(hana="KEEP3", remarks="Partial Invoice is pending"),
            _ht_row(hana="DROP1", remarks="Received"),
            _ht_row(hana="DROP2", remarks="Write off"),
            _ht_row(hana="DROP3", remarks="Closed"),
        ],
    )
    tracker = _tracker_rows(
        [{"Code": c, "Mail ID 1": "x@agg.com"} for c in ("KEEP1", "KEEP2", "KEEP3")]
    )

    plans = await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker)
    assert {p["business_key"] for p in plans} == {"KEEP1", "KEEP2", "KEEP3"}


async def test_diag_business_unit_filter(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            _ht_row(hana="OK", business_unit="e-Diagnostic-Aggregator"),
            _ht_row(hana="WRONG1", business_unit="e-Diagnostic"),
            _ht_row(hana="WRONG2", business_unit="e-Pharmacy (Platform Aggregator)"),
            _ht_row(hana="WRONG3", business_unit="Corporate Wellness"),
        ],
    )
    tracker = _tracker_rows([{"Code": "OK", "Mail ID 1": "ok@agg.com"}])

    plans = await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker)
    assert {p["business_key"] for p in plans} == {"OK"}


async def test_diag_final_amount_materiality_filter(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            _ht_row(hana="H1", final_pending=0.0),
            _ht_row(hana="H1", final_pending=0.50),
            _ht_row(hana="H1", final_pending=2500.0),
        ],
    )
    tracker = _tracker_rows([{"Code": "H1", "Mail ID 1": "h1@agg.com"}])

    plans = await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker)
    assert len(plans) == 1
    assert plans[0]["row_count"] == 1
    assert Decimal(plans[0]["totals"]["_net_pending"]) == Decimal("2500")


async def test_diag_not_due_month_end_sieve(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[
            # Due this month: 2026-06-17 + 30d = 2026-07-17 ≤ month end → keep
            _ht_row(
                hana="H1", invoice_date=datetime.date(2026, 6, 17), credit_days=30,
                final_pending=345000.0, ageing_type="Not Due",
            ),
            # Due next month: 2026-06-17 + 60d = 2026-08-16 > month end → drop
            _ht_row(
                hana="H1", invoice_date=datetime.date(2026, 6, 17), credit_days=60,
                final_pending=999999.0, ageing_type="Not Due",
            ),
        ],
    )
    tracker = _tracker_rows([{"Code": "H1", "Mail ID 1": "h1@agg.com"}])

    plan = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0]
    assert plan["row_count"] == 1
    assert Decimal(plan["totals"]["_net_pending"]) == Decimal("345000")
    # Kept Not Due row contributes nothing to overdue exposure.
    assert Decimal(plan["totals"]["_overdue_pending"]) == Decimal("0")
    # The upcoming-invoice clause is removed from aggregator emails per AR,
    # even though the Not Due row is still projected as Due this Month.
    assert "upcoming invoices listed above" not in plan["rendered_body_html"]


async def test_diag_overpayment_skip_via_unaccounted(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        ht_rows=[_ht_row(hana="H1", final_pending=500.0)],
        party_wise_rows=[_party_wise_row(code="H1", unaccounted=600.0)],
    )
    tracker = _tracker_rows([{"Code": "H1", "Mail ID 1": "h1@agg.com"}])

    plan = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0]
    assert plan["skip_gate"]["code"] == "potential_overpayment"


async def test_diag_recipient_split_to_and_cc(tmp_path):
    xlsx = _write_workbook(tmp_path, ht_rows=[_ht_row(hana="MT", final_pending=8000.0)])
    tracker = _tracker_rows([
        {
            "Code": "MT",
            "Mail ID 1": "primary@client.com",
            "Mail ID 2": "cc1@client.com",
            "Mail ID 3": "cc2@client.com",
        }
    ])

    plan = (await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker))[0]
    assert plan["skip_gate"] is None
    assert "primary@client.com" in plan["resolved_to"]
    assert "himanshu.singh@1mg.com" in plan["resolved_to"]
    assert "cc1@client.com" in plan["resolved_cc"]
    assert "cc2@client.com" in plan["resolved_cc"]
    assert "cc1@client.com" not in plan["resolved_to"]


# =========================================================================== #
# platform_aggregator (e-Pharmacy (Platform Aggregator))                      #
# =========================================================================== #


async def test_platform_normal_send_cleans_float_code(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        agg_rows=[
            _agg_row(customer=1000017221.0, value=5159069.98),
            _agg_row(customer=1000017221.0, value=11914939.0),
        ],
    )
    tracker = _tracker_rows([{"Code": "1000017221", "Mail ID 1": "truworth@agg.com"}])

    plans = await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker)
    assert len(plans) == 1
    plan = plans[0]
    assert plan["variant"] == "platform_aggregator"
    # SAP float customer code rendered without the trailing ".0".
    assert plan["business_key"] == "1000017221"
    assert plan["row_count"] == 2
    assert Decimal(plan["totals"]["_net_pending"]) == Decimal("17074008.98")
    assert plan["rendered_subject"] == (
        "Payment Reminder – Overdue Invoices | TRUWORTH HEALTH TECHNOLOGIES | "
        "Tata1mg | 1000017221"
    )


async def test_platform_bu_filter_scopes_to_lgp(tmp_path):
    """Only ``BU = LGP`` rows belong to the platform aggregator."""

    xlsx = _write_workbook(
        tmp_path,
        agg_rows=[
            _agg_row(customer=1000017221.0, value=5000.0, bu="LGP"),
            # Same sheet, other BUs / no BU → must not leak into this variant.
            _agg_row(customer=2000000000.0, value=9000.0, bu=None),
            _agg_row(customer=3000000000.0, value=9000.0, bu="OTHER"),
        ],
    )
    tracker = _tracker_rows([{"Code": "1000017221", "Mail ID 1": "t@agg.com"}])

    plans = await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker)
    assert {p["business_key"] for p in plans} == {"1000017221"}


async def test_platform_not_due_requires_overdue_flag(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        agg_rows=[
            # Month Bucket=Not Due + Projected Ageing(AZ)=OverDue → kept,
            # rendered as "Due this Month".
            _agg_row(customer=1000017221.0, value=4000.0,
                     month_bucket_at="Not Due", projected_ageing_az="OverDue"),
            # Not Due and Projected Ageing(AZ) not OverDue → dropped
            _agg_row(customer=1000017221.0, value=999999.0,
                     month_bucket_at="Not Due", projected_ageing_az=None),
        ],
    )
    tracker = _tracker_rows([{"Code": "1000017221", "Mail ID 1": "t@agg.com"}])

    plan = (await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker))[0]
    assert plan["row_count"] == 1
    assert Decimal(plan["totals"]["_net_pending"]) == Decimal("4000")
    # A Not Due bucket kept via the AZ "OverDue" flag is labelled Due this Month
    # and must NOT count toward overdue exposure (it is due, not yet overdue).
    assert Decimal(plan["totals"]["_overdue_pending"]) == Decimal("0")
    body = plan["rendered_body_html"]
    # Ageing is no longer a table column, and the upcoming-invoice clause is
    # removed from aggregator emails per AR (even though the Not-Due-but-OverDue
    # row is still projected as "Due this Month" internally).
    assert ">Ageing<" not in body
    assert "upcoming invoices listed above" not in body
    assert ">Remarks<" not in body  # no Remarks column on this variant
    # Intro names this variant's Business Unit, not the diagnostics one.
    assert (
        "the following e-Pharmacy (Platform Aggregator) invoices that remain "
        "overdue in our records" in body
    )
    assert "e-Diagnostic-Aggregator" not in body


# =========================================================================== #
# Schema-drift safety — missing columns must not mail wrong rows              #
# =========================================================================== #


def _write_agg_workbook_without(tmp_path: Path, drop_column: str) -> Path:
    """Workbook whose ``Invoice wise Agg.`` sheet is missing *drop_column*.

    The H&T sheet stays intact so we can assert the other variant is unaffected.
    """
    wb = Workbook()
    wb.remove(wb.active)

    ws_ht = wb.create_sheet(title="Invoice details-H&T")
    ws_ht.append([])
    ws_ht.append([])
    ws_ht.append(_HT_HEADER)
    ws_ht.append([_ht_row(hana="8106")[h] for h in _HT_HEADER])

    ws_party = wb.create_sheet(title="Party wise Ageing-H&T")
    ws_party.append(_PARTY_WISE_HEADER)

    keep = [i for i, h in enumerate(_AGG_HEADER) if h != drop_column]
    ws_agg = wb.create_sheet(title="Invoice wise Agg.")
    ws_agg.append([_AGG_HEADER[i] for i in keep])
    row = _agg_row(customer=1000017221.0, value=5000.0)
    ws_agg.append([row[i] for i in keep])

    path = tmp_path / "receivables.xlsx"
    wb.save(path)
    return path


@pytest.mark.parametrize(
    "dropped, expected_key",
    [
        # Filter column. Critically, "Month Bucket" is used with `ne` — an absent
        # column reads as None, so `ne(None, "Not Due")` would flip to TRUE and
        # mail rows that should never qualify. Must fail closed instead.
        ("BU", "bu#0"),
        ("Month Bucket", "month bucket#0"),
        ("Projected Ageing of Next Month", "projected ageing of next month#0"),
        # Projection column — the amount the email is built from.
        ("Company Code Currency Value", "company code currency value#0"),
    ],
)
async def test_missing_column_aborts_variant_and_sends_nothing(
    tmp_path, dropped, expected_key
):
    from app.email_automation.pipeline.process import _RequiredColumnsMissing

    xlsx = _write_agg_workbook_without(tmp_path, dropped)
    tracker = _tracker_rows([{"Code": "1000017221", "Mail ID 1": "t@agg.com"}])

    with pytest.raises(_RequiredColumnsMissing) as exc:
        await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker)

    # The error names the offending column so ops can see what AR renamed.
    assert expected_key in exc.value.missing
    assert exc.value.variant == "platform_aggregator"


async def test_missing_column_in_one_variant_leaves_others_working(tmp_path):
    """A broken sheet must not take the whole workbook down."""

    from app.email_automation.pipeline.process import _RequiredColumnsMissing

    xlsx = _write_agg_workbook_without(tmp_path, "BU")
    tracker = _tracker_rows([
        {"Code": "8106", "Mail ID 1": "c@agg.com"},
        {"Code": "1000017221", "Mail ID 1": "t@agg.com"},
    ])

    with pytest.raises(_RequiredColumnsMissing):
        await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker)

    # Same workbook, healthy sheet → diagnostics still renders and would send.
    plans = await _build(DIAGNOSTICS_AGGREGATOR_VARIANT, xlsx, tracker)
    assert [p["business_key"] for p in plans] == ["8106"]
    assert plans[0]["skip_gate"] is None


async def test_platform_currency_materiality_filter(tmp_path):
    xlsx = _write_workbook(
        tmp_path,
        agg_rows=[
            _agg_row(customer=1000017221.0, value=0.0),
            _agg_row(customer=1000017221.0, value=0.50),
            _agg_row(customer=1000017221.0, value=7500.0),
        ],
    )
    tracker = _tracker_rows([{"Code": "1000017221", "Mail ID 1": "t@agg.com"}])

    plan = (await _build(PLATFORM_AGGREGATOR_VARIANT, xlsx, tracker))[0]
    assert plan["row_count"] == 1
    assert Decimal(plan["totals"]["_net_pending"]) == Decimal("7500")
