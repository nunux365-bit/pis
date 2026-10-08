"""Unit tests for receivables executive dashboard parser."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services import receivable_dashboard as rd


def test_tab_name_rules() -> None:
    assert rd.is_receivable_tab("Receivable as on 15.04.2026")
    assert not rd.is_receivable_tab("Partywise  Unbilled")
    assert rd.is_unbilled_tab("Partywise Unbilled")
    assert rd.is_unbilled_tab("Party-wise Unbilled")
    assert not rd.is_receivable_tab("Summary")
    assert rd.is_tds_aging_tab("TDS Ageing")
    assert rd.is_tds_aging_tab("TDS Aging")
    assert rd.is_tds_aging_tab("TDS-Ageing")
    assert not rd.is_tds_aging_tab("Receivable as on 15.04.2026")
    assert not rd.is_tds_aging_tab("Partywise Unbilled")


def test_normalize_match_bu() -> None:
    a = rd.normalize_bu_unbilled("Speciality-Pharma")
    b = rd.normalize_bu_unbilled("Speciality Pharma")
    assert a == b


def test_normalize_bu_key_same_as_unbilled() -> None:
    """:func:`normalize_bu_key` is the unbilled segment key (hyphens/space, not order-dependent)."""
    for s in (
        "Retail-West",
        "Retail West",
        "Hospitals & Hospital IPD",
        "Hospitals&HospitalIPD",
    ):
        assert rd.normalize_bu_key(s) == rd.normalize_bu_unbilled(s)
    assert rd.normalize_bu_key("A-B C") == "abc"


def test_build_receivable_payload_minimal(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "0-1 Months",
        "1-3 Months",
        "3-6 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        ws.cell(3, c, 100_000)  # 1 lakh INR per KPI cell
    ws.cell(5, 1, "C1")
    ws.cell(5, 2, "Party A")
    ws.cell(5, 3, "BU1")
    ws.cell(5, 4, 50_000)
    ws.cell(5, 5, 0)
    ws.cell(5, 6, 50_000)
    ws.cell(5, 7, 0)
    ws.cell(5, 8, 0)
    ws.cell(5, 9, 1_000_000)  # 10 lakh net
    ws.cell(6, 1, "C2")
    ws.cell(6, 2, "Party B")
    ws.cell(6, 3, "BU1")
    for c in range(4, 9):
        ws.cell(6, c, 0)
    ws.cell(6, 9, 2_000_000)
    u = wb.create_sheet("Partywise Unbilled")
    u.cell(1, 1, "Segment")
    u.cell(1, 2, "Unbilled Amt")
    u.cell(2, 1, "BU1")
    u.cell(2, 2, 100_000)
    path = tmp_path / "recv.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(
        path, {"C1": 1, "C2": 3}
    )
    assert pl is not None
    assert pl["meta"].get("reminder_grid_format") == rd.REMINDER_GRID_FORMAT
    # "all" = sum of party detail (excludes payables), not the workbook control row
    k_all = pl["kpi_lakh"]["all"]
    assert k_all["Not Due"] == 0.0
    assert k_all["0-1 Months"] == pytest.approx(0.5)
    assert k_all.get("Net Due") == pytest.approx(30.0)  # 1M + 2M
    top = pl["top_parties"]["all"]
    # Excel logic: sort by sum(age buckets) then Net Due → C1 (50k ageing) before C2 (0 ageing).
    assert [r["code"] for r in top] == ["C1", "C2"]
    assert all("wtd_age_months" in r and r["wtd_age_months"] >= 0 for r in top)
    by_code = {r["code"]: r["wtd_age_months"] for r in top}
    # C1: weighted on 0–1 mo slice only → midpoint 0.5 mo. C2: no bucket INR → dominant 0–1 midpoint.
    assert by_code["C1"] == pytest.approx(0.5, abs=0.01)
    assert by_code["C2"] == pytest.approx(0.5, abs=0.01)
    assert top[0]["net_due_lakh"] == pytest.approx(0.5)  # SUM(age) / 1e5, not Net Due
    # C2 has no money in ageing columns; Net Due is 20L but “due” = bucket sum = 0.
    assert top[1]["net_due_lakh"] == 0.0
    assert pl["unbilled_lakh"]["all"] == 1.0
    g = pl["collection_grids"]["all"]
    assert g["reminder_bands"] == ["0-1", "2", "3", "4", "5", "6+"]
    assert len(g["values_lakh"]) == 5
    assert len(g["values_lakh"][0]) == 6
    assert g["ageing_buckets"] == list(rd.EXCEL_GRID_AGE_LABELS)
    assert "Not Due" not in (g.get("ageing_buckets") or [])
    # Only C1 has ageing INR; matrix uses bucket sums (50k INR = 0.5 L) in one cell.
    assert sum(v for row in g["values_lakh"] for v in row) == pytest.approx(0.5)
    # Party lines include ₹ / months like the Excel “Parties” matrix.
    lines = g["client_names"]
    assert any("₹" in x for row in lines for cell in row for x in cell)
    assert "BU1" in pl["kpi_lakh"]["by_business_unit"]


def test_kpi_and_grid_use_norm_rollout_replicated_for_each_bu_string(
    tmp_path: Path,
) -> None:
    """``normalize_bu_key``-equivalent BUs share one KPI rollup; it is under every raw spelling.
    Top/collection for each spelling use the same party set (all rows in that norm)."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "0-1 Months",
        "1-3 Months",
        "3-6 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        ws.cell(3, c, 0)
    ws.cell(5, 1, "A")
    ws.cell(5, 2, "P1")
    ws.cell(5, 3, "Alpha & Beta Unit")
    for c in range(4, 9):
        ws.cell(5, c, 0)
    ws.cell(5, 6, 50_000)
    ws.cell(5, 9, 50_000)
    ws.cell(6, 1, "B")
    ws.cell(6, 2, "P2")
    ws.cell(6, 3, "Alpha&BetaUnit")
    for c in range(4, 9):
        ws.cell(6, c, 0)
    ws.cell(6, 6, 30_000)
    ws.cell(6, 9, 30_000)
    path = tmp_path / "twospell.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    k = pl["kpi_lakh"]["by_business_unit"]
    a = "Alpha & Beta Unit"
    b = "Alpha&BetaUnit"
    assert a in k and b in k
    # Same norm: 50k + 30k = 80k in 0-1
    assert k[a]["0-1 Months"] == pytest.approx(0.8)
    assert k[b]["0-1 Months"] == pytest.approx(0.8)
    g = pl["collection_grids"]["by_business_unit"]
    assert sum(
        v for row in g[a]["values_lakh"] for v in row
    ) == sum(v for row in g[b]["values_lakh"] for v in row)


def test_is_excluded_payable_ledger_name() -> None:
    assert rd.is_excluded_payable_ledger_name("ASSAM X - Payable")
    assert rd.is_excluded_payable_ledger_name("ASSAM X -Payable")
    assert rd.is_excluded_payable_ledger_name("Y—Payable")
    assert not rd.is_excluded_payable_ledger_name("Accounts Payable")
    assert not rd.is_excluded_payable_ledger_name("Acme Corp")


def test_kpi_includes_ageing_when_subtotal_cell_blank(tmp_path: Path) -> None:
    """Subtotal row may omit a figure; allowlist is header-driven so detail still rolls up (A, C)."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "0-1 Months",
        "1-3 Months",
        "3-6 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        if c == 6:
            continue
        ws.cell(3, c, 0)
    # Row 3 col 6 = 0-1 Months: blank in KPI row, but party line has 100k in that column
    ws.cell(5, 1, "C1")
    ws.cell(5, 2, "Party A")
    ws.cell(5, 3, "BU1")
    for c in range(4, 9):
        ws.cell(5, c, 0)
    ws.cell(5, 6, 100_000)
    ws.cell(5, 9, 0)
    path = tmp_path / "blank_subtotal_01.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    assert pl["kpi_lakh"]["by_business_unit"]["BU1"].get("0-1 Months") == pytest.approx(1.0)
    g = pl["collection_grids"]["all"]
    assert sum(v for row in g["values_lakh"] for v in row) == pytest.approx(1.0)


def test_duplicate_identical_header_labels_sum_both_columns(tmp_path: Path) -> None:
    """B: two columns with the same title — KPI and party row sum both; warn at runtime."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "0-1 Months",
        "0-1 Months",
        "1-3 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 8):
        ws.cell(3, c, 0)
    ws.cell(5, 1, "X")
    ws.cell(5, 2, "P")
    ws.cell(5, 3, "BU1")
    ws.cell(5, 4, 40_000)
    ws.cell(5, 5, 10_000)
    ws.cell(5, 6, 0)
    ws.cell(5, 7, 0)
    path = tmp_path / "dup_hdr.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    # _sum_kpi_by_bu iterates each column: 40k+10k = 0.5 L
    assert pl["kpi_lakh"]["by_business_unit"]["BU1"].get("0-1 Months") == pytest.approx(0.5)
    p = pl["top_parties"]["all"][0]
    assert p["code"] == "X"
    assert p["net_due_lakh"] == pytest.approx(0.5)


def test_excludes_trailing_payable_ledger_name_from_bu_kpi_and_grid(
    tmp_path: Path,
) -> None:
    """Negative SAP *-Payable* line must not offset receivable in KPI or show on grid."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "0-1 Months",
        "1-3 Months",
        "3-6 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        ws.cell(3, c, 100_000)
    # Receivable line
    ws.cell(5, 1, "R1")
    ws.cell(5, 2, "Main Customer")
    ws.cell(5, 3, "BU1")
    for c in range(4, 9):
        ws.cell(5, c, 0)
    ws.cell(5, 6, 100_000)
    ws.cell(5, 9, 100_000)
    # Contra: would net KPI to 0 if included
    ws.cell(6, 1, "L1")
    ws.cell(6, 2, "Main Customer -Payable")
    ws.cell(6, 3, "BU1")
    for c in range(4, 9):
        ws.cell(6, c, 0)
    ws.cell(6, 6, -100_000)
    ws.cell(6, 9, -100_000)
    path = tmp_path / "with_payable.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    assert pl["meta"]["excluded_payable_ledger_rows"] == 1
    assert pl["kpi_lakh"]["all"]["0-1 Months"] == pytest.approx(1.0)
    assert pl["kpi_lakh"]["by_business_unit"]["BU1"]["0-1 Months"] == pytest.approx(
        1.0
    )
    g = pl["collection_grids"]["all"]
    assert sum(v for row in g["values_lakh"] for v in row) == pytest.approx(1.0)


def test_no_receivable_tab_returns_none(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Random Summary"
    ws.cell(1, 1, "x")
    path = tmp_path / "n.xlsx"
    wb.save(path)
    assert rd.build_receivable_payload(path, {}) is None


def test_unbilled_segment_credits_all_matching_bus(tmp_path: Path) -> None:
    """Same normalized segment must attribute the row to every BU spelling in the sheet."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "0-1 Months",
        "1-3 Months",
        "3-6 Months",
        "Net Due",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        ws.cell(3, c, 0)
    ws.cell(5, 1, "C1")
    ws.cell(5, 2, "Party A")
    ws.cell(5, 3, "Retail-West")
    for c in range(4, 9):
        ws.cell(5, c, 0)
    ws.cell(5, 9, 100_000)
    ws.cell(6, 1, "C2")
    ws.cell(6, 2, "Party B")
    ws.cell(6, 3, "Retail West")
    for c in range(4, 9):
        ws.cell(6, c, 0)
    ws.cell(6, 9, 100_000)
    u = wb.create_sheet("Partywise Unbilled")
    u.cell(1, 1, "Segment")
    u.cell(1, 2, "Unbilled Amt")
    u.cell(2, 1, "Retail West")
    u.cell(2, 2, 300_000)
    path = tmp_path / "recv_unb.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    assert pl["unbilled_lakh"]["all"] == pytest.approx(3.0)
    ub = pl["unbilled_lakh"]["by_business_unit"]
    assert ub.get("Retail-West") == pytest.approx(3.0)
    assert ub.get("Retail West") == pytest.approx(3.0)


def test_moved_header_row_and_reordered_columns(tmp_path: Path) -> None:
    """Header/KPI are not on fixed row 4/3; key columns are not in default order."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "CHW-Receivable"
    h_row, kpi_row, d_row = 7, 6, 8
    headers = [
        "Net Due",
        "Business Unit",
        "0-1 Months",
        "Name of the party",
        "1-3 Months",
        "3-6 Months",
        "Not Due",
        "Code",
        "Receivables",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(h_row, c, h)
    for c in range(1, len(headers) + 1):
        ws.cell(kpi_row, c, 100_000)
    # Party row: Code col 8, Name 4, BU 2, Receivables 9, Not Due 7, age 3,5,6, Net 1
    row_vals = {8: "CX", 4: "X", 2: "BU1", 9: 0, 7: 0, 3: 40_000, 5: 0, 6: 0, 1: 500_000}
    for col, v in row_vals.items():
        ws.cell(d_row, col, v)
    path = tmp_path / "moved.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {"CX": 0})
    assert pl is not None
    assert pl["meta"]["receivable_layout"]["header_excel_row"] == h_row
    assert pl["meta"]["receivable_layout"]["data_start_excel_row"] == d_row
    # Party row Net 500_000 = 5.0 L; control row is not used for kpi "all"
    assert pl["kpi_lakh"]["all"].get("Net Due") == pytest.approx(5.0)
    g = pl["collection_grids"]["all"]
    assert sum(v for row in g["values_lakh"] for v in row) == pytest.approx(0.4)


def test_age_buckets_semantic_order_independent_of_sheet_order(tmp_path: Path) -> None:
    """O…T midpoints follow calendar bucket order, not the physical column order in the file."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2020"
    headers = [
        "Code",
        "Name of the party",
        "Business Unit",
        "Receivables",
        "Not Due",
        "3-6 Months",
        "0-1 Months",
        "1-3 Months",
        "O/s Amount",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(4, c, h)
    for c in range(1, 10):
        ws.cell(3, c, 100_000)
    ws.cell(5, 1, "C1")
    ws.cell(5, 2, "Party A")
    ws.cell(5, 3, "BU1")
    ws.cell(5, 4, 0)
    ws.cell(5, 5, 0)
    ws.cell(5, 6, 0)
    ws.cell(5, 7, 100_000)
    ws.cell(5, 8, 0)
    ws.cell(5, 9, 500_000)
    u = wb.create_sheet("Partywise Unbilled")
    u.cell(1, 1, "Segment")
    u.cell(1, 2, "Unbilled Amt")
    u.cell(2, 1, "BU1")
    u.cell(2, 2, 0)
    path = tmp_path / "permute_age.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {"C1": 0})
    assert pl is not None
    by = {r["code"]: r["wtd_age_months"] for r in pl["top_parties"]["all"]}
    assert by["C1"] == pytest.approx(0.5, abs=0.02)


# ── TDS Ageing helpers ────────────────────────────────────────────────────────

def _make_receivable_wb(parties: list[dict]) -> "Workbook":
    """Build a minimal receivable workbook with optional TDS Ageing sheet.

    Each party dict: code, name, bu, buckets (list of INR values for
    ``0-1 Months`` and ``1-3 Months``), net_due. Optionally ``tds_buckets``.
    """
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2025"
    headers = ["Code", "Name of the party", "Business Unit",
               "0-1 Months", "1-3 Months", "Net Due"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, h)  # KPI row
    for c in range(1, 7):
        ws.cell(3, c, 0)
    for c, h in enumerate(headers, 1):
        ws.cell(4, c, h)  # header row
    for ri, p in enumerate(parties, start=5):
        ws.cell(ri, 1, p["code"])
        ws.cell(ri, 2, p["name"])
        ws.cell(ri, 3, p["bu"])
        ws.cell(ri, 4, p["buckets"][0])
        ws.cell(ri, 5, p["buckets"][1])
        ws.cell(ri, 6, p["net_due"])
    has_tds = any("tds_buckets" in p for p in parties)
    if has_tds:
        tws = wb.create_sheet("TDS Ageing")
        tds_headers = ["Code", "Name of the party", "Business Unit",
                       "0-1 Months", "1-3 Months"]
        for c, h in enumerate(tds_headers, 1):
            tws.cell(1, c, h)
        row = 2
        for p in parties:
            if "tds_buckets" not in p:
                continue
            tws.cell(row, 1, p["code"])
            tws.cell(row, 2, p["name"])
            tws.cell(row, 3, p["bu"])
            tws.cell(row, 4, p["tds_buckets"][0])
            tws.cell(row, 5, p["tds_buckets"][1])
            row += 1
    return wb


def test_tds_subtracted_from_bucket_amounts(tmp_path: Path) -> None:
    """Grid cell and top-party amounts should reflect rec − TDS per bucket."""
    wb = _make_receivable_wb([
        {"code": "A1", "name": "Alpha", "bu": "BU1",
         "buckets": [300_000, 200_000], "net_due": 500_000,
         "tds_buckets": [50_000, 20_000]},
    ])
    path = tmp_path / "tds_basic.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    # Grid total should be (300k−50k) + (200k−20k) = 250k+180k = 430k = 4.3 L
    g = pl["collection_grids"]["all"]
    total = sum(v for row in g["values_lakh"] for v in row)
    assert total == pytest.approx(4.3, abs=0.01)
    # KPI bucket values should also be TDS-adjusted
    kpi = pl["kpi_lakh"]["all"]
    assert kpi.get("0-1 Months", 0) == pytest.approx(2.5, abs=0.01)
    assert kpi.get("1-3 Months", 0) == pytest.approx(1.8, abs=0.01)
    # No clip events expected
    dq = pl["meta"]["data_quality"]
    assert dq["tds_clip_count"] == 0
    assert dq["tds_clip_parties"] == []


def test_tds_multiple_rows_same_code_same_bu(tmp_path: Path) -> None:
    """Two receivable rows + two TDS rows for the same (code, BU): aggregate then subtract."""
    wb = _make_receivable_wb([
        {"code": "A1", "name": "Alpha", "bu": "BU1",
         "buckets": [100_000, 200_000], "net_due": 300_000},
        {"code": "A1", "name": "Alpha", "bu": "BU1",
         "buckets": [50_000, 30_000], "net_due": 80_000,
         "tds_buckets": [20_000, 30_000]},
    ])
    # Also add a TDS row for the first split that is a separate TDS row
    ws_tds = wb["TDS Ageing"]
    # Add second TDS row for same code
    ws_tds.cell(3, 1, "A1")
    ws_tds.cell(3, 2, "Alpha")
    ws_tds.cell(3, 3, "BU1")
    ws_tds.cell(3, 4, 5_000)   # extra TDS for 0-1 bucket
    ws_tds.cell(3, 5, 0)
    path = tmp_path / "tds_multi_row.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    # rec_total 0-1: 150k, TDS 0-1: 25k → net 125k = 1.25 L
    # rec_total 1-3: 230k, TDS 1-3: 30k → net 200k = 2.0 L
    kpi = pl["kpi_lakh"]["all"]
    assert kpi.get("0-1 Months", 0) == pytest.approx(1.25, abs=0.01)
    assert kpi.get("1-3 Months", 0) == pytest.approx(2.0, abs=0.01)


def test_tds_clip_flagged_when_tds_exceeds_receivable(tmp_path: Path) -> None:
    """When TDS > receivable in a bucket, the bucket is clamped to 0 and flagged."""
    wb = _make_receivable_wb([
        {"code": "B1", "name": "Beta", "bu": "BU1",
         "buckets": [10_000, 200_000], "net_due": 200_000,
         "tds_buckets": [50_000, 20_000]},  # TDS 50k > rec 10k in bucket 0
    ])
    path = tmp_path / "tds_clip.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    dq = pl["meta"]["data_quality"]
    assert dq["tds_clip_count"] == 1
    clip = dq["tds_clip_parties"][0]
    assert clip["code"] == "B1"
    assert clip["bucket"] == "0-1 Months"
    assert clip["rec_lakh"] == pytest.approx(0.1, abs=0.001)
    assert clip["tds_lakh"] == pytest.approx(0.5, abs=0.001)
    # Grid should use clamped 0 for bucket 0; bucket 1 = 200k − 20k = 1.8 L
    g = pl["collection_grids"]["all"]
    total = sum(v for row in g["values_lakh"] for v in row)
    assert total == pytest.approx(1.8, abs=0.01)


def test_no_tds_sheet_still_parses(tmp_path: Path) -> None:
    """Workbook without TDS Ageing sheet still produces a valid payload (TDS = 0)."""
    wb = _make_receivable_wb([
        {"code": "C1", "name": "Gamma", "bu": "BU1",
         "buckets": [100_000, 50_000], "net_due": 150_000},
    ])
    path = tmp_path / "no_tds.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    # Amounts unchanged — no TDS sheet means zero subtraction
    kpi = pl["kpi_lakh"]["all"]
    assert kpi.get("0-1 Months", 0) == pytest.approx(1.0, abs=0.01)
    assert kpi.get("1-3 Months", 0) == pytest.approx(0.5, abs=0.01)
    dq = pl["meta"]["data_quality"]
    assert dq["tds_clip_count"] == 0


def test_tds_different_bu_same_code_subtracted_independently(tmp_path: Path) -> None:
    """Same HANA code in two BUs: TDS subtracted per BU, not mixed across BUs."""
    wb = Workbook()
    ws = wb.active
    assert ws
    ws.title = "Receivable as on 01.01.2025"
    headers = ["Code", "Name of the party", "Business Unit",
               "0-1 Months", "1-3 Months", "Net Due"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, 0)
        ws.cell(4, c, h)
    # Same code, different BUs
    ws.cell(5, 1, "X1"); ws.cell(5, 2, "Xray"); ws.cell(5, 3, "Retail")
    ws.cell(5, 4, 200_000); ws.cell(5, 5, 0); ws.cell(5, 6, 200_000)
    ws.cell(6, 1, "X1"); ws.cell(6, 2, "Xray"); ws.cell(6, 3, "Pharma")
    ws.cell(6, 4, 100_000); ws.cell(6, 5, 0); ws.cell(6, 6, 100_000)
    # TDS sheet: same code, different BUs
    tws = wb.create_sheet("TDS Ageing")
    tds_hdrs = ["Code", "Name of the party", "Business Unit", "0-1 Months", "1-3 Months"]
    for c, h in enumerate(tds_hdrs, 1):
        tws.cell(1, c, h)
    tws.cell(2, 1, "X1"); tws.cell(2, 2, "Xray"); tws.cell(2, 3, "Retail")
    tws.cell(2, 4, 30_000); tws.cell(2, 5, 0)
    tws.cell(3, 1, "X1"); tws.cell(3, 2, "Xray"); tws.cell(3, 3, "Pharma")
    tws.cell(3, 4, 10_000); tws.cell(3, 5, 0)
    path = tmp_path / "tds_diff_bu.xlsx"
    wb.save(path)
    pl = rd.build_receivable_payload(path, {})
    assert pl is not None
    kpi_bu = pl["kpi_lakh"]["by_business_unit"]
    # Retail: 200k − 30k = 170k = 1.7 L
    assert kpi_bu["Retail"].get("0-1 Months", 0) == pytest.approx(1.7, abs=0.01)
    # Pharma: 100k − 10k = 90k = 0.9 L
    assert kpi_bu["Pharma"].get("0-1 Months", 0) == pytest.approx(0.9, abs=0.01)


def test_kpi_key_is_due_ageing_bucket_columns_only() -> None:
    """Due-ageing = month/yrs buckets only — not finance summary columns named *overdue*."""
    assert rd.kpi_key_is_due_ageing("0-1 Months")
    assert rd.kpi_key_is_due_ageing("+1Yrs")
    assert rd.kpi_key_is_due_ageing("6-12 Months")
    assert not rd.kpi_key_is_due_ageing("Net Due")
    assert not rd.kpi_key_is_due_ageing("Not Due")
    assert not rd.kpi_key_is_due_ageing("Receivables")
    assert not rd.kpi_key_is_due_ageing("Final Overdue as on 30th June")
    assert not rd.kpi_key_is_due_ageing("Final Overdues")
    assert not rd.kpi_key_is_due_ageing("Overdue before adjustment")
    assert not rd.kpi_key_is_due_ageing("Overdue")
